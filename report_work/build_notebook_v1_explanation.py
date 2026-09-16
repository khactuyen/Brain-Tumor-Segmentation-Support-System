from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"D:\SIC_Capstone 2026")
OUTPUT = ROOT / "doc" / "SIC_Capstone_v1_Giai_thich_chi_tiet.docx"
ASSETS = ROOT / "report_work" / "notebook_v1_explained_assets"
LABEL_VISUAL = ROOT / "report_work" / "brats_label_visual_for_report.png"

NAVY = "1E3A8A"
BLUE = "1E40AF"
LIGHT_BLUE = "EAF2FF"
PALE = "F8FAFC"
GRAY = "475569"
LIGHT_GRAY = "E2E8F0"
WHITE = "FFFFFF"
GREEN = "166534"
AMBER = "92400E"
RED = "991B1B"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_paragraph_border(paragraph, color: str, side: str = "left", size: str = "18") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    border = OxmlElement(f"w:{side}")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), size)
    border.set(qn("w:space"), "8")
    border.set(qn("w:color"), color)
    p_bdr.append(border)


def set_repeat_header_footer(section, header_text: str) -> None:
    header = section.header
    hp = header.paragraphs[0]
    hp.text = header_text
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in hp.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor.from_string(GRAY)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("Page ")
    add_field(run, "PAGE")
    run = fp.add_run(" / ")
    add_field(run, "NUMPAGES")
    for run in fp.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor.from_string(GRAY)


def add_field(run, instruction: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)


def add_toc(doc: Document) -> None:
    p = doc.add_paragraph()
    run = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    # Giữ mục lục gọn ở hai cấp; 24 tiêu đề code cell đã được trình bày
    # tuần tự trong Mục 5 nên không cần làm mục lục tràn sang một trang riêng.
    instr.text = ' TOC \\o "1-2" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Mục lục sẽ được cập nhật khi mở bằng Microsoft Word."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, placeholder, end):
        run._r.append(node)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(1.8)
    set_repeat_header_footer(section, "SIC Capstone 2026  |  Giải thích notebook v1")

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12

    for style_name, size, color in (
        ("Heading 1", 17, NAVY),
        ("Heading 2", 14, BLUE),
        ("Heading 3", 12, NAVY),
    ):
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    caption.font.name = "Times New Roman"
    caption.font.size = Pt(9.5)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(GRAY)
    caption.paragraph_format.space_after = Pt(7)

    code = styles.add_style("Code Block", 1)
    code.font.name = "Consolas"
    code.font.size = Pt(8.3)
    code.paragraph_format.left_indent = Cm(0.25)
    code.paragraph_format.right_indent = Cm(0.15)
    code.paragraph_format.space_before = Pt(3)
    code.paragraph_format.space_after = Pt(5)
    code.paragraph_format.line_spacing = 1.0

    settings = doc.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def add_heading_1(doc: Document, text: str, page_break: bool = False) -> None:
    p = doc.add_paragraph(text, style="Heading 1")
    if page_break:
        p.paragraph_format.page_break_before = True
    set_paragraph_shading(p, LIGHT_BLUE)
    set_paragraph_border(p, BLUE, side="bottom", size="12")


def add_body(doc: Document, text: str, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)


def add_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)


def add_numbered(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)


def add_code(doc: Document, code: str) -> None:
    p = doc.add_paragraph(style="Code Block")
    set_paragraph_shading(p, "F1F5F9")
    set_paragraph_border(p, "94A3B8", side="left", size="16")
    p.add_run(code.rstrip())


def add_callout(doc: Document, title: str, text: str, kind: str = "note") -> None:
    color = {"note": BLUE, "good": GREEN, "warning": AMBER, "risk": RED}[kind]
    fill = {"note": "EFF6FF", "good": "F0FDF4", "warning": "FFFBEB", "risk": "FEF2F2"}[kind]
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=100, bottom=100, start=120, end=120)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{title}: ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(color)
    p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string(WHITE)
        if widths:
            cell.width = Inches(widths[i])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cell = cells[i]
            if row_index % 2:
                set_cell_shading(cell, PALE)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(value))
            run.font.name = "Times New Roman"
            run.font.size = Pt(9.2)
            if widths:
                cell.width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_picture(doc: Document, path: Path, caption: str, width: float = 6.25) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    cp = doc.add_paragraph(caption, style="Caption")
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.keep_together = True


def add_cell_explanation(
    doc: Document,
    number: int,
    title: str,
    purpose: str,
    code: str,
    explanation: list[str],
    output: str,
    question: str,
    answer: str,
) -> None:
    doc.add_paragraph(f"Code cell {number}: {title}", style="Heading 3")
    add_body(doc, purpose, bold_prefix="Mục tiêu. ")
    add_code(doc, code)
    add_bullets(doc, explanation)
    add_callout(doc, "Đầu ra cần nhìn thấy", output, "good")
    add_callout(doc, "Câu hỏi bảo vệ", f"{question} — {answer}", "note")


def build_document() -> None:
    doc = Document()
    configure_document(doc)

    # Cover
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(75)
    r = p.add_run("SIC CAPSTONE 2026")
    r.font.name = "Arial"
    r.font.size = Pt(17)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    r = p.add_run("GIẢI THÍCH CHI TIẾT\nNOTEBOOK V1")
    r.font.name = "Arial"
    r.font.size = Pt(30)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run("SIC_Capstone_v1.ipynb\nBrain Tumor Segmentation với 3D U-Net")
    r.font.name = "Times New Roman"
    r.font.size = Pt(16)
    r.font.italic = True
    r.font.color.rgb = RGBColor.from_string(GRAY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    p.add_run(
        "Tài liệu học và chuẩn bị bảo vệ\n"
        "Giải thích luồng dữ liệu, từng code cell, kết quả và các điểm cần audit"
    ).font.size = Pt(13)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(75)
    r = p.add_run("Cập nhật: 15/09/2026")
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor.from_string(GRAY)

    add_callout(
        doc,
        "Phạm vi",
        "Tài liệu giải thích đúng notebook v1 đang có trong workspace. Kết quả là đầu ra nghiên cứu, chưa được thẩm định lâm sàng.",
        "warning",
    )
    doc.add_page_break()

    add_heading_1(doc, "MỤC LỤC")
    add_toc(doc)
    doc.add_page_break()

    # 1
    add_heading_1(doc, "1. Cách đọc tài liệu này")
    add_body(
        doc,
        "Tài liệu đi theo đúng thứ tự thực thi của notebook v1. Mỗi code cell được giải thích theo bốn câu hỏi: cell nhận gì, xử lý gì, tạo ra gì và nếu sai thì hậu quả gì. Phần cuối tập trung vào những câu hỏi khó thường xuất hiện khi bảo vệ đồ án.",
    )
    add_numbered(doc, [
        "Đọc mục 2–4 để nắm bài toán, dataset và toàn bộ pipeline.",
        "Đọc mục 5 theo từng code cell và đối chiếu trực tiếp với notebook.",
        "Học mục 6–9 để giải thích tensor, training, metric và kết quả.",
        "Dùng mục 10–11 để chuẩn bị câu hỏi phản biện và thuật ngữ.",
    ])
    add_callout(
        doc,
        "Câu trả lời cốt lõi",
        "V1 huấn luyện 3D U-Net trên patch MRI 3D bốn kênh, chọn checkpoint bằng full validation và chỉ dùng held-out test để báo cáo cuối.",
        "good",
    )

    # 2
    add_heading_1(doc, "2. Tổng quan notebook v1", page_break=True)
    add_body(
        doc,
        "Notebook v1 xây dựng một pipeline phân đoạn u não 3D hoàn chỉnh: đọc ảnh NIfTI BraTS, kiểm tra dữ liệu, chuẩn hóa và cache NPZ, chia train/validation/test, huấn luyện 3D U-Net, chọn checkpoint, đánh giá trên held-out test, trực quan hóa ba mặt phẳng và xuất báo cáo PDF nghiên cứu.",
    )
    add_table(doc, ["Thành phần", "Giá trị quan sát trong notebook v1"], [
        ["Dataset hợp lệ", "1.251 ca BraTS"],
        ["Split", "875 train / 188 validation / 188 held-out test"],
        ["Đầu vào", "4 modality: T1n, T1c, T2w, T2-FLAIR"],
        ["Mô hình thực sự train", "3D U-Net (MODELS_TO_TRAIN = ['unet3d'])"],
        ["Số tham số", "4.811.129"],
        ["Patch", "96 × 96 × 96 voxel"],
        ["Checkpoint được chọn", "Epoch 40, dựa trên full validation"],
        ["Kết quả test", "Mean Dice 0,7858; Mean IoU 0,6934; HD95 31,117 mm"],
    ], widths=[1.65, 4.55])
    doc.add_paragraph("Pipeline ở mức khái niệm", style="Heading 2")
    add_table(doc, ["Bước", "Dòng chảy"], [
        ["1", "NIfTI bốn modality + segmentation mask"],
        ["2", "Kiểm tra shape, affine, spacing, intensity và label"],
        ["3", "Clip percentile → Z-score → remap label → NPZ"],
        ["4", "Chia case 70% / 15% / 15%"],
        ["5", "Crop patch có ưu tiên vùng u + augmentation 3D"],
        ["6", "3D U-Net → 4 logit voxel"],
        ["7", "DiceCE loss → AdamW → cập nhật trọng số"],
        ["8", "Full validation → lưu checkpoint tốt nhất"],
        ["9", "Sliding-window + TTA + hậu xử lý → held-out test"],
        ["10", "NIfTI prediction + hình ba mặt phẳng + PDF"],
    ], widths=[0.6, 5.6])

    # 3
    add_heading_1(doc, "3. Dataset, modality và nhãn", page_break=True)
    add_body(
        doc,
        "Mỗi ca là một volume MRI 3D của cùng bệnh nhân. Bốn modality phải đồng nhất về shape và affine để mỗi voxel ở bốn kênh mô tả cùng một vị trí giải phẫu. Notebook ghép chúng theo thứ tự cố định thành tensor C×H×W×D.",
    )
    add_table(doc, ["Modality", "Vai trò trực quan", "Tên hậu tố"], [
        ["T1 native", "Giải phẫu nền", "t1n"],
        ["T1 contrast", "Làm nổi vùng tăng tương phản", "t1c"],
        ["T2 weighted", "Nhạy với thành phần giàu nước", "t2w"],
        ["T2-FLAIR", "Làm rõ phù và tổn thương lan tỏa", "t2f"],
    ], widths=[1.2, 3.8, 1.2])
    doc.add_paragraph("Nhãn gốc và nhãn nội bộ", style="Heading 2")
    add_table(doc, ["Raw", "Class mô hình", "Ý nghĩa"], [
        ["0", "0", "Background"],
        ["1", "1", "NCR/NET – lõi hoại tử hoặc không tăng tương phản"],
        ["2", "2", "ED – vùng phù nề"],
        ["4", "3", "ET – vùng tăng tương phản"],
    ], widths=[0.8, 1.2, 4.2])
    add_body(
        doc,
        "Notebook cũng chấp nhận raw label 3 và ánh xạ nó về class 3 để tương thích một số cache/biến thể dữ liệu. Với bộ raw theo quy ước 0, 1, 2, 4, ET được đổi từ 4 thành 3 để đầu ra mô hình có bốn class liên tiếp 0–3.",
    )
    add_table(doc, ["Vùng đánh giá", "Công thức trên class nội bộ", "Ý nghĩa"], [
        ["WT", "label > 0", "Toàn bộ u: 1 + 2 + 3"],
        ["TC", "label ∈ {1, 3}", "Lõi u: NCR/NET + ET"],
        ["ET", "label = 3", "Vùng tăng tương phản"],
    ], widths=[0.8, 2.25, 3.15])
    if LABEL_VISUAL.exists():
        add_picture(
            doc,
            LABEL_VISUAL,
            "Hình 1. Vị trí các nhãn trên một lát cắt T2-FLAIR và quan hệ ET ⊂ TC ⊂ WT.",
            width=6.25,
        )
    add_callout(
        doc,
        "Điểm phải nhớ",
        "Raw label 4 và model class 3 đều chỉ cùng một mô ET. Đây là đổi chỉ số, không phải tạo thêm loại mô.",
        "note",
    )

    # 4
    add_heading_1(doc, "4. Tiền xử lý và kiểm soát dữ liệu", page_break=True)
    add_body(
        doc,
        "Mỗi modality được kiểm tra độc lập nhưng dùng modality đầu làm mốc shape, affine và spacing. Chỉ các voxel có intensity lớn hơn 0 được xem là vùng não để tính percentile, mean và standard deviation.",
    )
    add_numbered(doc, [
        "Đọc NIfTI thành float32 và xác nhận volume có đúng ba chiều.",
        "Kiểm tra mọi giá trị hữu hạn; loại NaN hoặc infinity.",
        "Kiểm tra shape và affine của bốn modality trùng nhau.",
        "Tạo brain mask bằng điều kiện intensity > 0.",
        "Clip intensity tại percentile 0,5% và 99,5% để giảm ảnh hưởng outlier.",
        "Chuẩn hóa Z-score riêng từng modality trong vùng não: (x − μ) / σ.",
        "Đặt voxel ngoài brain mask về 0.",
        "Đọc segmentation, kiểm tra shape/affine, kiểm tra label và remap 4 → 3.",
        "Lưu image, label, affine, spacing, case ID và preprocess version vào NPZ nén.",
    ])
    add_code(doc, "brain_mask = array > 0\nlow, high = np.percentile(values, [0.5, 99.5])\narray = np.clip(array, low, high)\narray = (array - mean_value) / std_value\narray[~brain_mask] = 0.0")
    add_callout(
        doc,
        "Tại sao Z-score theo từng modality",
        "T1, T1c, T2 và FLAIR có thang intensity khác nhau. Chuẩn hóa riêng giúp gradient ổn định và tránh một kênh chi phối chỉ vì có trị số lớn hơn.",
        "note",
    )
    add_callout(
        doc,
        "Giới hạn cache",
        "PREPROCESS_VERSION được lưu trong NPZ/checkpoint, nhưng cache_is_current hiện chỉ kiểm tra file tồn tại và đọc được; code chưa từ chối cache có preprocess_version cũ.",
        "risk",
    )
    if (ASSETS / "v1_eda_raw_nifti.png").exists():
        doc.add_page_break()
        add_picture(
            doc,
            ASSETS / "v1_eda_raw_nifti.png",
            "Hình 2. EDA của năm ca đầu tiên theo thứ tự thư mục; mỗi hàng gồm bốn modality và ground-truth overlay.",
            width=6.15,
        )
        add_body(
            doc,
            "EDA chọn lát axial có tổng diện tích u lớn nhất, nên hình dễ quan sát nhưng không đại diện cho phân bố lát cắt ngẫu nhiên. Năm ca cũng được lấy bằng sorted(... )[:5], không phải mẫu ngẫu nhiên hoặc mẫu phân tầng.",
        )

    # 5 cell-by-cell
    add_heading_1(doc, "5. Giải thích theo từng code cell", page_break=True)
    add_body(
        doc,
        "Notebook có 24 code cell. Số dưới đây là thứ tự code cell, không phải execution_count hiển thị bên trái Jupyter. Hai cell định nghĩa hàm cache chưa được chạy riêng trong trạng thái lưu, nhưng các hàm đó phải tồn tại trước khi cell chuẩn bị dataset được thực thi.",
    )

    cell_specs = [
        (1, "Cài thư viện", "Mục tiêu. Chuẩn bị MONAI, nibabel, PyTorch và các thư viện đánh giá/đồ họa.",
         '%pip install -q "monai>=1.3,<2.0" "nibabel>=5.2,<6" ...',
         ["Version range giảm nguy cơ API thay đổi lớn.", "Lệnh %pip cài vào đúng kernel notebook đang hoạt động."],
         "Cell hoàn tất cài đặt, không có lỗi import.",
         "Vì sao phải khóa phiên bản?", "Để một API thay đổi sau này không làm notebook chạy khác hoặc lỗi."),
        (2, "Import", "Mục tiêu. Nạp module xử lý file, mảng, deep learning, metric, visualization và report.",
         "import numpy as np\nimport torch\nimport nibabel as nib\nimport monai.transforms as mt\nfrom monai.networks.nets import UNet, SwinUNETR",
         ["nibabel đọc/ghi NIfTI.", "MONAI cung cấp transform, model và sliding-window inference.", "SciPy được dùng cho surface distance và connected components."],
         "Không có output; các tên module có mặt trong namespace.",
         "MONAI thay thế PyTorch hay không?", "Không. MONAI xây trên PyTorch và bổ sung thành phần chuyên cho ảnh y khoa."),
        (3, "Mount Google Drive", "Mục tiêu. Cho notebook dùng dữ liệu và checkpoint lưu bền trên Google Drive khi chạy Colab.",
         'if not Path("/content/drive/MyDrive").exists():\n    drive.mount("/content/drive")',
         ["ImportError cho biết đang chạy ngoài Colab.", "Code kiểm tra MyDrive sau mount thay vì chỉ tin lệnh mount."],
         "Google Drive is ready.",
         "Tại sao không lưu checkpoint chỉ trong /content?", "Runtime Colab là tạm thời; mất kết nối có thể xóa dữ liệu local."),
        (4, "Seed và thiết bị", "Mục tiêu. Đặt seed cho Python/NumPy/PyTorch/MONAI và chọn CUDA nếu có.",
         "SEED = 42\nset_determinism(seed=SEED)\ntorch.backends.cudnn.deterministic = True\ndevice = torch.device('cuda' if torch.cuda.is_available() else 'cpu')",
         ["Seed giúp lặp lại split và phần lớn phép augmentation.", "Deterministic có thể giảm tốc độ so với benchmark mode."],
         "Device cuda; GPU Tesla T4; VRAM 14,56 GB.",
         "Seed 42 có bảo đảm giống tuyệt đối không?", "Không hoàn toàn; còn phụ thuộc GPU kernel, version thư viện, worker và phần cứng."),
        (5, "Đường dẫn và cấu hình", "Mục tiêu. Đặt toàn bộ đường dẫn, mapping nhãn và hyperparameter ở một nơi.",
         "PATCH_SIZE = (96, 96, 96)\nBATCH_SIZE = 1\nNUM_SAMPLES = 2\nSTEPS_PER_EPOCH = 150\nMODELS_TO_TRAIN = ['unet3d']",
         ["V1 chỉ train 3D U-Net.", "VAL_SPLIT và TEST_SPLIT đều là 0,15.", "NPZ được đồng bộ từ Drive sang SSD Colab để đọc nhanh."],
         "In ba đường dẫn raw data, Drive NPZ và runtime NPZ.",
         "Batch size 1 có nghĩa mỗi epoch chỉ có một mẫu?", "Không. Mỗi step có một case đầu vào; num_samples=2 tạo hai patch, và có tối đa 150 step mỗi epoch."),
        (6, "EDA từ raw NIfTI", "Mục tiêu. Kiểm tra trực quan bốn modality và nhãn trước preprocessing.",
         "z_sum = np.sum(label_mapped > 0, axis=(0, 1))\nz_idx = int(np.argmax(z_sum))",
         ["Chọn lát có diện tích u lớn nhất.", "Overlay dùng đỏ cho NCR, xanh cho ED và cam cho ET.", "Chỉ lấy năm thư mục đầu sau khi sắp xếp."],
         "Một figure 5×5 gồm năm ca và năm cột.",
         "EDA này có chứng minh toàn dataset sạch không?", "Không. Nó là kiểm tra định tính một số ca; vẫn cần audit tự động toàn bộ dataset."),
        (7, "Đọc cache và preprocess một ca", "Mục tiêu. Chuẩn hóa schema cache và tạo NPZ an toàn từ raw NIfTI.",
         "image = np.stack(channels, axis=0)\nlabel[raw_seg == 4] = 3\nnp.savez_compressed(...)",
         ["Hỗ trợ nhiều tên key như image/img/data/mri.", "Chuyển channels-last thành channels-first khi cần.", "Ghi file tạm rồi replace để tránh NPZ dở dang."],
         "Một file NPZ có image, label, seg, affine, spacing, version và case ID.",
         "Tại sao label phải dùng nearest-neighbor khi resample?", "Nội suy tuyến tính tạo label thập phân không còn là class hợp lệ."),
        (8, "Chuẩn bị toàn dataset", "Mục tiêu. Tái dùng cache, preprocess ca còn thiếu và đồng bộ sang SSD local.",
         "cases_to_build = [cid for cid in raw_cases if cid not in reusable_cases]\nshutil.copy2(source, destination)",
         ["Không preprocess lại 1.251 ca nếu cache đã có.", "So kích thước file để quyết định đồng bộ.", "Trả về danh sách case ID từ thư mục đích."],
         "Loaded 1251 preprocessed cases.",
         "Vì sao dùng SSD local nhanh hơn Drive?", "Drive mount có độ trễ I/O lớn; train đọc NPZ lặp lại nhiều lần."),
        (9, "Chạy chuẩn bị và QC một mẫu", "Mục tiêu. Xác nhận cache có shape, dtype, label, spacing và affine hợp lệ.",
         "Image: (4, 240, 240, 155) float32\nLabel: (1, 240, 240, 155) labels=[0, 1, 2, 3]",
         ["Bốn kênh ảnh ở trục đầu.", "Label có một kênh và đã remap về 0–3.", "Spacing quan sát là 1×1×1 mm."],
         "Valid cases: 1251; affine determinant 1,0000.",
         "Determinant affine bằng 1 nói được gì?", "Nó phù hợp với scale/orientation đang quan sát, nhưng một số duy nhất không thay thế kiểm tra đầy đủ ma trận affine."),
        (10, "Chia train/validation/test", "Mục tiêu. Tạo ba tập case rời nhau với seed cố định.",
         "train_val_cases, test_cases = train_test_split(..., test_size=0.15)\nrelative_val_size = 0.15 / 0.85",
         ["Tách test trước, rồi tách validation khỏi phần còn lại.", "relative_val_size bảo đảm validation cuối cùng vẫn bằng 15% tổng.", "Ba assert kiểm tra không trùng case."],
         "875 train; 188 validation; 188 held-out test.",
         "Tại sao không gọi test_size=0.15 lần hai?", "Vì lần hai chỉ còn 85% dữ liệu; phải dùng 0,15/0,85 để validation bằng 15% tổng."),
        (11, "Lớp BraTSDataset3D", "Mục tiêu. Biến danh sách case thành giao diện Dataset mà PyTorch/MONAI hiểu.",
         "def __getitem__(self, index):\n    cached = load_validated_cache(npz_path)\n    return self.transforms(sample) if self.transforms else sample",
         ["Mỗi sample giữ image, label, case_id, affine và spacing.", "Label chuyển int64 để dùng cho loss phân loại voxel.", "Có fallback đọc cache từ Drive."],
         "Dataset trả một dictionary thống nhất cho DataLoader.",
         "Vì sao giữ affine và spacing trong sample?", "Affine dùng để lưu output đúng không gian; spacing dùng tính thể tích và HD95 theo mm."),
        (12, "Transform train và evaluation", "Mục tiêu. Tạo patch giàu vùng u, augmentation cho train và evaluation xác định.",
         "RandCropByPosNegLabeld(pos=3.0, neg=1.0, num_samples=2)\nRandAffined(mode=['bilinear', 'nearest'])",
         ["Tỷ lệ pos:neg 3:1 giảm số patch chỉ có nền.", "Ảnh affine bằng bilinear; label bằng nearest.", "Evaluation chỉ EnsureTyped, không augmentation."],
         "Transform train có crop, flip, rotate, affine, noise, scale và shift intensity.",
         "Tại sao augmentation ảnh và mask phải đồng bộ?", "Nếu biến hình học khác nhau, label không còn nằm đúng vị trí tổn thương."),
        (13, "Dataset và DataLoader", "Mục tiêu. Tạo loader riêng cho train, fast validation, full validation và test.",
         "train_loader = MonaiDataLoader(..., batch_size=1, shuffle=True)\nval_loader_full = MonaiDataLoader(..., shuffle=False)",
         ["Train được shuffle; evaluation giữ thứ tự.", "pin_memory bật khi có CUDA.", "Fast validation có 20 ca; full validation và test mỗi tập 188 ca."],
         "Train loader có 875 step khả dụng nhưng training chỉ lấy tối đa 150.",
         "Một epoch ở đây có quét hết train set không?", "Không. Vòng lặp dừng ở 150/875 step; đây là epoch theo ngân sách step của notebook."),
        (14, "Metric, sliding window, TTA và hậu xử lý", "Mục tiêu. Tạo WT/TC/ET, tính metric và dự đoán full volume trong giới hạn VRAM.",
         "REGION_BUILDERS = {'wt': array > 0, 'tc': isin([1,3]), 'et': array == 3}\npred = torch.argmax(probs, dim=1, keepdim=True)",
         ["Sliding window dùng ROI 96³, overlap 0,5 và Gaussian blending.", "TTA lật một trục rồi trung bình xác suất.", "Xóa connected component nhỏ hơn 50 voxel theo từng class."],
         "Trả prediction shape [1,1,H,W,D] và bảng metric theo từng ca.",
         "Metric đang đo model nguyên bản hay pipeline hoàn chỉnh?", "Pipeline hoàn chỉnh, vì evaluation dùng TTA và hậu xử lý trước khi tính metric."),
        (15, "Model factory", "Mục tiêu. Khởi tạo 3D U-Net theo một contract bốn kênh vào và bốn class ra.",
         "UNet(spatial_dims=3, in_channels=4, out_channels=4,\n     channels=(16,32,64,128,256), num_res_units=2,\n     norm='instance', dropout=0.2)",
         ["InstanceNorm phù hợp hơn BatchNorm khi batch rất nhỏ.", "Skip connection giữ chi tiết không gian.", "Residual units giúp gradient đi qua mạng sâu ổn định hơn."],
         "unet3d: 4,811,129 parameters.",
         "Vì sao out_channels=4 khi chỉ báo cáo WT/TC/ET?", "Model dự đoán bốn class độc quyền 0–3; WT/TC/ET được gộp sau argmax để đánh giá."),
        (16, "DiceCE loss và tiện ích checkpoint", "Mục tiêu. Kết hợp overlap toàn vùng với lỗi phân loại từng voxel.",
         "DiceCELoss(to_onehot_y=True, softmax=True,\n           include_background=False, lambda_dice=1.0, lambda_ce=1.0)",
         ["Dice thành phần giảm ảnh hưởng mất cân bằng vùng u.", "Cross Entropy tạo tín hiệu cục bộ cho từng voxel.", "Background bị loại khỏi phần Dice để tránh lớp nền áp đảo."],
         "Khởi tạo criterion, MODEL_RUNS, GradScaler và EarlyStopping.",
         "Tại sao vẫn cần CE khi đã có Dice?", "CE ổn định gradient ở mức voxel, nhất là đầu training khi overlap còn thấp."),
        (17, "Training, resume và chọn checkpoint", "Mục tiêu. Train tới epoch 50 từ checkpoint cũ, validate định kỳ và chọn mô hình bằng full validation.",
         "MODEL_RUNS['unet3d'] = train_one_model('unet3d', total_epochs=50, resume=True)",
         ["Resume dùng learning rate 2e-5.", "Mixed precision dùng autocast + GradScaler.", "Gradient norm bị clip ở 1,0.", "Fast val mỗi 2 epoch; full val mỗi 10 epoch.", "Checkpoint chỉ lưu khi full Dice tốt hơn."],
         "Selected model from validation only: unet3d, epoch 40.",
         "Resume có tiếp tục optimizer hoàn toàn không?", "Không. Code lưu optimizer_state nhưng khi resume chỉ load model và scaler; AdamW được khởi tạo mới với LR 2e-5."),
        (18, "Vẽ learning curves", "Mục tiêu. Đọc history từ checkpoint và quan sát loss cùng Dice theo epoch.",
         "history = checkpoint.get('history', {})\naxes[0].plot(history['train_loss'])\naxes[1].plot(history['full_val_dice'])",
         ["Loss giảm thể hiện tối ưu hóa đang tiến triển.", "Full validation đáng tin hơn fast subset 20 ca.", "Không chọn checkpoint bằng train loss."],
         "Một figure gồm Training Loss và Validation Dice.",
         "Vì sao fast Dice có các điểm tụt mạnh?", "Subset chỉ 20 ca nên phương sai lớn; model cũng thay đổi theo epoch. Nó chỉ là tín hiệu theo dõi."),
        (19, "In bảng kết quả cuối", "Mục tiêu. Báo cáo metric validation và held-out test của checkpoint được chọn.",
         "for split_name in ('val', 'test'):\n    summary = run[f'{split_name}_summary']",
         ["Best epoch giống nhau cho hai dòng vì cùng một checkpoint.", "Validation dùng chọn model; test chỉ báo cáo cuối.", "HD95 có đơn vị mm vì có spacing."],
         "Val Mean Dice 0,8052; test Mean Dice 0,7858.",
         "Có được chọn epoch dựa trên test không?", "Không. Làm vậy gây test leakage và làm điểm test mất ý nghĩa ước lượng tổng quát hóa."),
        (20, "Hàm trực quan ba mặt phẳng", "Mục tiêu. So sánh FLAIR, ground truth và prediction trên axial/coronal/sagittal.",
         "z_idx = argmax(sum(mask > 0, axis=(0,1)))\ny_idx = argmax(sum(mask > 0, axis=(0,2)))\nx_idx = argmax(sum(mask > 0, axis=(1,2)))",
         ["Mỗi mặt phẳng chọn lát có diện tích mask lớn nhất.", "Nếu có ground truth, lát được chọn theo ground truth; nếu không thì theo prediction.", "Overlay dùng cùng bảng màu class 1–3."],
         "Hàm được định nghĩa, chưa tạo hình cho tới cell tiếp theo.",
         "Vì sao hình đẹp chưa chứng minh model tốt?", "Đây là một ca và các lát thuận lợi nhất; phải dùng metric toàn tập và xem cả ca lỗi."),
        (21, "Dự đoán một held-out test case", "Mục tiêu. Chạy checkpoint tốt nhất trên một ca test, tính metric, lưu NIfTI và vẽ kết quả.",
         "sample_batch = next(iter(test_loader))\nsample_prediction = predict_volume(best_model, sample_image)\nnib.save(Nifti1Image(sample_prediction, sample_affine), prediction_path)",
         ["Ca là phần tử đầu của test loader không shuffle.", "Giữ affine gốc khi lưu.", "Metric dùng spacing của ca."],
         "Case BraTS-GLI-01021-000; một file prediction NIfTI và figure 3×3.",
         "Output đã tương thích raw BraTS chưa?", "Chưa hoàn toàn: class ET=3 đang được lưu trực tiếp; nếu yêu cầu raw convention phải đổi 3 trở lại 4."),
        (22, "Lưu hình segmentation", "Mục tiêu. Tạo PNG ba mặt phẳng để nhúng vào PDF report.",
         "save_segmentation_figure(image_4d, prediction, output_file, ground_truth)",
         ["Dùng FLAIR làm ảnh nền.", "Có legend class 1,2,3.", "Đóng figure sau khi lưu để tránh chiếm bộ nhớ."],
         "File PNG segmentation được tạo trong reports_dir.",
         "Tại sao report dùng FLAIR làm nền?", "FLAIR thường làm phù và vùng tổn thương lan tỏa rõ; các modality khác vẫn cần cho model."),
        (23, "Định nghĩa PDF research report", "Mục tiêu. Đóng gói case ID, model, spacing, volume, hình và metric thành báo cáo nghiên cứu.",
         "class ResearchSegmentationReport(FPDF): ...\ndef export_research_report(...): ...",
         ["Header ghi rõ research use only.", "Volume tính theo cm³.", "Metric được gọi là reference-based, không phải confidence score.", "Limitations nằm ở trang riêng."],
         "Hai hàm/class sẵn sàng để xuất PDF.",
         "Dice của một ca có phải độ tin cậy khi không có ground truth không?", "Không. Dice cần ground truth; khi triển khai thật không thể tự tính Dice cho ca mới."),
        (24, "Xuất report mẫu", "Mục tiêu. Gọi hàm lưu hình và xuất PDF cho ca demo đã suy luận.",
         "sample_volumes = {region: sample_metrics[f'volume_cm3_{region}'] ...}\nexport_research_report(...) ",
         ["Lấy volume WT/TC/ET từ sample_metrics.", "Nhúng PNG segmentation vào report.", "Đưa metric chỉ vì ca test có ground truth."],
         "Research report của BraTS-GLI-01021-000 được lưu thành công.",
         "Report này có dùng để chẩn đoán được không?", "Không. Nó là output nghiên cứu, chưa có clinical validation hoặc chứng nhận thiết bị y tế."),
    ]

    for spec in cell_specs:
        add_cell_explanation(doc, *spec)

    # 6 Tensor flow
    add_heading_1(doc, "6. Dòng chảy tensor qua mô hình", page_break=True)
    add_table(doc, ["Giai đoạn", "Shape điển hình", "Giải thích"], [
        ["NPZ image", "(4, 240, 240, 155)", "4 modality, volume đầy đủ"],
        ["NPZ label", "(1, 240, 240, 155)", "Một class index tại mỗi voxel"],
        ["Một patch image", "(4, 96, 96, 96)", "Crop 3D đưa vào train"],
        ["Batch hiệu dụng", "xấp xỉ (2, 4, 96, 96, 96)", "batch_size=1 nhưng num_samples=2"],
        ["Logit model", "(B, 4, 96, 96, 96)", "Bốn score chưa chuẩn hóa tại mỗi voxel"],
        ["Softmax", "(B, 4, 96, 96, 96)", "Xác suất bốn class, tổng bằng 1"],
        ["Argmax", "(B, 1, 96, 96, 96)", "Chọn class có xác suất cao nhất"],
        ["Full-volume output", "(1, 1, 240, 240, 155)", "Ghép sliding windows"],
    ], widths=[1.35, 1.85, 3.0])
    add_callout(
        doc,
        "Phân biệt logit và mask",
        "Logit là số thực cho từng class và còn gradient; mask là class rời rạc sau argmax, không dùng để backpropagate.",
        "note",
    )

    # 7 training
    add_heading_1(doc, "7. Training, validation và checkpoint", page_break=True)
    add_body(
        doc,
        "Một training step thực hiện: lấy patch → chuyển GPU → zero gradient → forward trong autocast → DiceCE loss → backward qua GradScaler → unscale → clip gradient → optimizer step → cập nhật scaler. Scheduler chỉ đổi learning rate sau mỗi epoch.",
    )
    add_table(doc, ["Khái niệm", "Trong notebook v1", "Vai trò"], [
        ["Optimizer", "AdamW", "Cập nhật trọng số và tách weight decay khỏi gradient update"],
        ["Resume LR", "2e-5", "Fine-tune nhẹ hơn sau checkpoint"],
        ["Scheduler", "CosineAnnealingLR đến 1e-6", "Giảm LR mượt theo cosine"],
        ["Mixed precision", "autocast + GradScaler", "Giảm VRAM/tăng tốc CUDA"],
        ["Gradient clipping", "max_norm=1,0", "Giảm nguy cơ gradient bùng nổ"],
        ["Fast validation", "20 ca, mỗi 2 epoch", "Theo dõi nhanh"],
        ["Full validation", "188 ca, mỗi 10 epoch", "Lưu checkpoint và early stopping"],
        ["Held-out test", "188 ca, sau khi chọn checkpoint", "Báo cáo cuối"],
    ], widths=[1.3, 1.75, 3.15])
    add_callout(
        doc,
        "Điểm tinh tế về early stopping",
        "patience=8 tăng bộ đếm chỉ khi full validation chạy. Vì full validation cách 10 epoch, đây là 8 lần kiểm tra chứ không phải 8 epoch; với lần chạy tới epoch 50, cơ chế này gần như không có đủ tám lần để dừng.",
        "warning",
    )
    if (ASSETS / "v1_training_curves.png").exists():
        add_picture(
            doc,
            ASSETS / "v1_training_curves.png",
            "Hình 3. Train loss giảm và validation Dice tăng; checkpoint tốt nhất được chọn tại epoch 40.",
            width=6.25,
        )

    # 8 metrics
    add_heading_1(doc, "8. Dice, IoU, Precision, Recall, HD95 và volume", page_break=True)
    add_table(doc, ["Metric", "Công thức/ý nghĩa", "Chiều tốt"], [
        ["Dice", "2|P∩G| / (|P|+|G|): độ chồng lấp", "Càng gần 1 càng tốt"],
        ["IoU", "|P∩G| / |P∪G|: độ chồng lấp nghiêm ngặt hơn", "Càng gần 1 càng tốt"],
        ["Precision", "TP/(TP+FP): phần dự đoán dương tính là đúng", "Càng cao càng tốt"],
        ["Recall", "TP/(TP+FN): phần tổn thương thật được tìm thấy", "Càng cao càng tốt"],
        ["HD95", "Phân vị 95% khoảng cách hai bề mặt, có spacing", "Càng thấp càng tốt"],
        ["Volume", "Số voxel dự đoán × thể tích một voxel / 1000", "Không có hướng tốt cố định"],
    ], widths=[1.0, 4.25, 1.25])
    add_callout(
        doc,
        "Quy tắc mask rỗng",
        "Nếu prediction và ground truth cùng rỗng, notebook gán Dice/IoU/Precision/Recall = 1 và HD95 = 0. Quy tắc này hợp lý theo nghĩa không có lỗi, nhưng có thể làm macro mean cao nếu nhiều ca không có ET.",
        "warning",
    )
    add_callout(
        doc,
        "Một bên rỗng",
        "Dice, IoU, precision hoặc recall phản ánh thất bại; HD95 được gán bằng độ dài đường chéo vật lý của volume để tạo hình phạt lớn thay vì NaN.",
        "note",
    )

    # 9 results
    add_heading_1(doc, "9. Kết quả và cách diễn giải", page_break=True)
    add_table(doc, ["Split", "Mean Dice", "WT", "TC", "ET", "Mean IoU", "HD95 (mm)"], [
        ["Validation", "0,8052", "0,8308", "0,8125", "0,7724", "0,7127", "29,683"],
        ["Held-out test", "0,7858", "0,8158", "0,7938", "0,7476", "0,6934", "31,117"],
    ], widths=[1.15, 0.85, 0.65, 0.65, 0.65, 0.85, 0.95])
    add_bullets(doc, [
        "Test Mean Dice thấp hơn validation 0,0194; đây là khoảng cách tổng quát hóa quan sát được.",
        "WT cao nhất trên test (0,8158), TC tiếp theo (0,7938), ET thấp nhất (0,7476).",
        "HD95 test cao hơn validation 1,434 mm, cho thấy biên hoặc outlier trên dữ liệu chưa dùng để chọn model khó hơn.",
        "Các con số mô tả đúng pipeline có TTA và lọc component nhỏ; không phải model thô trước hậu xử lý.",
        "Kết quả chỉ là internal held-out test trên cùng hệ sinh thái BraTS; chưa chứng minh khả năng dùng trên bệnh viện khác.",
    ])
    if (ASSETS / "v1_prediction_three_planes.png").exists():
        add_picture(
            doc,
            ASSETS / "v1_prediction_three_planes.png",
            "Hình 4. Ca test minh họa BraTS-GLI-01021-000 trên axial, coronal và sagittal.",
            width=6.2,
        )
    add_callout(
        doc,
        "Không được suy rộng quá mức",
        "Một hình dự đoán đẹp không thay thế đánh giá 188 ca. Dice không phải phần trăm chẩn đoán đúng và report không phải kết luận lâm sàng.",
        "risk",
    )

    # 10 audit
    add_heading_1(doc, "10. Những điểm cần audit hoặc cải thiện", page_break=True)
    add_table(doc, ["Mức", "Điểm cần kiểm tra", "Hướng xử lý"], [
        ["Cao", "Prediction class 3 được lưu thẳng vào NIfTI", "Đổi class 3 → raw label 4 trước khi xuất nếu cần chuẩn BraTS"],
        ["Cao", "Cache version được lưu nhưng chưa enforce", "So preprocess_version trước khi tái dùng NPZ"],
        ["Vừa", "Resume không load optimizer_state", "Gọi rõ là fine-tune hoặc khôi phục optimizer/scheduler đầy đủ"],
        ["Vừa", "Early stopping đếm full-val event, không phải epoch", "Đổi tên patience_checks hoặc điều chỉnh interval/patience"],
        ["Vừa", "EDA lấy năm ca đầu và lát u lớn nhất", "Bổ sung random/stratified sampling và thống kê toàn cohort"],
        ["Vừa", "Split chưa stratify theo burden/ET presence", "Audit phân bố WT/TC/ET giữa ba split"],
        ["Vừa", "Ngưỡng component là 50 voxel", "Ghi đơn vị vật lý hoặc kiểm chứng ngưỡng trên validation"],
        ["Thấp", "TTA flip phụ thuộc hướng dữ liệu", "Chuẩn hóa orientation và kiểm tra affine trước khi gọi là trái–phải"],
    ], widths=[0.7, 2.9, 2.9])
    add_callout(
        doc,
        "Cách trình bày khi bảo vệ",
        "Nêu thẳng đây là các giới hạn kỹ thuật đã nhận diện. Khả năng chỉ ra giới hạn và đề xuất kiểm chứng thường thuyết phục hơn việc khẳng định pipeline hoàn hảo.",
        "good",
    )

    # 11 viva
    add_heading_1(doc, "11. Bộ câu hỏi bảo vệ khó và câu trả lời ngắn", page_break=True)
    qa = [
        ("Tại sao dùng 3D U-Net thay vì U-Net 2D?", "Volume 3D chứa tính liên tục qua các lát. Convolution 3D học ngữ cảnh theo cả ba trục, đổi lại tốn VRAM hơn."),
        ("Vì sao cần bốn modality?", "Mỗi modality nhấn một loại tương phản mô; ghép bốn kênh cung cấp thông tin bổ sung cho cùng voxel."),
        ("Raw label 4 và class 3 khác nhau thế nào?", "Chỉ khác mã hóa. Cả hai đều là ET; remap giúp class liên tục 0–3."),
        ("Tại sao Z-score chỉ tính trong brain mask?", "Nền 0 rất lớn sẽ kéo mean/std và làm chuẩn hóa mô não sai lệch."),
        ("Clip percentile để làm gì?", "Giảm ảnh hưởng của vài intensity cực trị trước khi tính mean/std."),
        ("Tại sao cache NPZ?", "Giảm chi phí đọc NIfTI và preprocessing lặp lại; đổi lại phải quản lý version cache."),
        ("Held-out test khác validation thế nào?", "Validation dùng chọn checkpoint; test chỉ mở sau khi đã khóa model để ước lượng tổng quát hóa."),
        ("Tại sao chia ở mức case?", "Để lát/patch của cùng bệnh nhân không rơi vào nhiều split và gây leakage."),
        ("Patch 96³ có nhược điểm gì?", "Giảm VRAM nhưng giới hạn ngữ cảnh; sliding window và overlap dùng để tái tạo full volume."),
        ("NUM_SAMPLES=2 khác BATCH_SIZE=1 thế nào?", "Một case tạo hai patch; DataLoader nhận một case mỗi lần nhưng có thể collate thành hai patch hiệu dụng."),
        ("Tại sao crop pos:neg = 3:1?", "U chiếm ít voxel; ưu tiên patch chứa u giúp model nhận đủ tín hiệu lớp hiếm."),
        ("Tại sao label dùng nearest trong affine transform?", "Label là class rời rạc; bilinear sẽ tạo giá trị trung gian không hợp lệ."),
        ("Tại sao InstanceNorm?", "Batch rất nhỏ làm thống kê BatchNorm thiếu ổn định; InstanceNorm chuẩn hóa riêng từng sample/channel."),
        ("DiceCE kết hợp lợi ích gì?", "Dice tập trung overlap/lớp hiếm; CE tạo gradient cục bộ ổn định ở từng voxel."),
        ("Vì sao loại background khỏi Dice loss?", "Nền áp đảo số voxel và có thể làm loss đẹp dù vùng u kém."),
        ("Fast validation có dùng chọn checkpoint không?", "Không. Nó chỉ theo dõi nhanh; checkpoint dựa trên full validation 188 ca."),
        ("Tại sao epoch 40 tốt hơn epoch 50?", "Checkpoint được chọn theo full validation; epoch cuối có thể overfit hoặc không cải thiện."),
        ("AdamW khác Adam ở đâu?", "AdamW tách weight decay khỏi cập nhật gradient thích nghi, cho regularization rõ hơn."),
        ("Mixed precision có đổi kết quả không?", "Có thể tạo sai khác số học nhỏ; mục tiêu chính là giảm VRAM và tăng tốc, GradScaler tránh underflow."),
        ("Gradient clipping làm gì?", "Co gradient khi norm vượt 1,0 để tránh bước cập nhật quá lớn."),
        ("Sliding-window overlap 0,5 có tác dụng gì?", "Mỗi voxel được nhìn trong nhiều cửa sổ; Gaussian blending giảm artefact tại biên patch."),
        ("TTA có phải train lại model không?", "Không. Nó dự đoán ảnh gốc và ảnh lật, đảo kết quả rồi trung bình xác suất."),
        ("Dice cao nhưng HD95 cao nghĩa là gì?", "Phần lớn vùng chồng lấp tốt nhưng vẫn có đoạn biên hoặc cụm dự đoán xa."),
        ("Tại sao ET thường khó hơn WT?", "ET nhỏ, biến thiên và mất cân bằng hơn; sai vài voxel tạo tỷ lệ lỗi lớn."),
        ("Empty-empty Dice=1 có nguy cơ gì?", "Nếu nhiều ca không có ET, macro mean có thể cao dù model chưa chứng minh phát hiện ET dương tính."),
        ("Kết quả test thấp hơn validation có bất thường không?", "Không nhất thiết. Test chưa dùng chọn model thường khó hơn; chênh lệch nhỏ là generalization gap quan sát được."),
        ("Số tham số lớn hơn có chắc tốt hơn không?", "Không. Hiệu quả phụ thuộc dữ liệu, protocol, regularization và đánh giá công bằng."),
        ("NIfTI prediction hiện có điểm gì cần sửa?", "Nếu theo raw BraTS, phải reverse-map class 3 thành label 4 trước khi lưu."),
        ("Tài liệu PDF có thể đưa cho bác sĩ chẩn đoán không?", "Chỉ dùng nghiên cứu/minh họa; chưa clinical validation và không phải thiết bị y tế."),
        ("Bằng chứng mạnh nhất của v1 là gì?", "Có held-out test 188 ca, checkpoint chọn bằng validation và metric gồm overlap lẫn khoảng cách biên."),
    ]
    add_table(doc, ["Câu hỏi", "Câu trả lời đề xuất"], [[q, a] for q, a in qa], widths=[2.5, 4.0])

    # 12 glossary
    add_heading_1(doc, "12. Thuật ngữ cần thuộc", page_break=True)
    glossary = [
        ("Affine", "Ma trận ánh xạ voxel sang tọa độ không gian thực của ảnh."),
        ("Augmentation", "Biến đổi dữ liệu train để tăng tính đa dạng mà vẫn giữ nhãn đúng."),
        ("Batch", "Nhóm sample xử lý trong một optimizer step."),
        ("Checkpoint", "Tệp lưu trọng số và trạng thái cần thiết của một thời điểm training."),
        ("Cohort", "Nhóm ca bệnh tham gia một thí nghiệm."),
        ("Ground truth", "Mask tham chiếu dùng để train hoặc đánh giá."),
        ("Held-out test", "Tập giữ riêng, không dùng điều chỉnh model."),
        ("Logit", "Score thô của model trước softmax."),
        ("Macro mean", "Trung bình cho các class/vùng với trọng số ngang nhau."),
        ("Modality", "Một kiểu chuỗi xung MRI tạo tương phản mô khác nhau."),
        ("NIfTI", "Định dạng ảnh y khoa giữ volume và metadata không gian."),
        ("NPZ", "Tệp nén chứa nhiều mảng NumPy và metadata."),
        ("Patch", "Khối con 3D cắt từ volume để giảm bộ nhớ."),
        ("Protocol", "Tập quy tắc về data, split, preprocessing, checkpoint và metric."),
        ("Spacing", "Kích thước vật lý của một voxel theo ba trục, thường tính bằng mm."),
        ("Voxel", "Phần tử thể tích 3D, tương tự pixel trong ảnh 2D."),
    ]
    add_table(doc, ["Thuật ngữ", "Giải thích"], [[a, b] for a, b in glossary], widths=[1.45, 5.05])

    doc.add_paragraph("Checklist trước khi trình bày", style="Heading 2")
    add_bullets(doc, [
        "Nói được input shape, output shape và mapping nhãn.",
        "Vẽ được pipeline từ NIfTI đến held-out test.",
        "Phân biệt validation với test và fast validation với full validation.",
        "Giải thích được Dice, HD95 và empty-mask policy.",
        "Nêu đúng checkpoint epoch 40 và kết quả test.",
        "Chủ động nêu hai điểm audit: cache version và reverse label mapping khi xuất NIfTI.",
    ])

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()

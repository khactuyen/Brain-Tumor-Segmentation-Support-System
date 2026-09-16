from pathlib import Path
from docx import Document

path = Path(r"D:\SIC_Capstone 2026\doc\SIC_Capstone_v1_Giai_thich_chi_tiet.docx")
doc = Document(path)
text = "\n".join(p.text for p in doc.paragraphs)
missing_cells = [i for i in range(1, 25) if f"Code cell {i}" not in text]
headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
print(f"file={path}")
print(f"size={path.stat().st_size}")
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} images={len(doc.inline_shapes)}")
print(f"headings={len(headings)}")
print(f"missing_cells={missing_cells}")
print(f"has_internal_source_reference={'Nguồn: output cell' in text}")
print(f"has_24_cell_heading={'24 code cell' in text}")
assert not missing_cells
assert "Nguồn: output cell" not in text
assert len(doc.inline_shapes) >= 4
assert path.stat().st_size > 100_000

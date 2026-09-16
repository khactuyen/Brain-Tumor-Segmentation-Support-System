# -*- coding: utf-8 -*-
from __future__ import annotations

import hashlib
import os
import re
import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(r"D:\SIC_Capstone 2026")
INPUT = ROOT / "doc" / "SIC_AI_Capstone Project_Final Report_Brain Tumor Segmentation_Revised.docx"
OUTPUT = ROOT / "report_work" / "SIC_AI_Final_Report_v1_v2_working.docx"
BUILDER = ROOT / "report_work" / "build_report.py"
OLD_GENERATED = ROOT / "report_work" / "generated"
NEW_GENERATED = ROOT / "report_work" / "v1_v2_diagrams"


NUMBER_REPLACEMENTS = {
    "0,7613": "0,8308",
    "0,7717": "0,8125",
    "0,7334": "0,7724",
    "0,7555": "0,8052",
    "0,6490": "0,7127",
    "43,087": "29,683",
    "0,7489": "0,8158",
    "0,7534": "0,7938",
    "0,7061": "0,7476",
    "0,7361": "0,7858",
    "0,6310": "0,6934",
    "43,058": "31,117",
}


def shift_versions(text: str) -> str:
    """Rename the two compared notebooks without cascading v2 into v1."""
    text = text.replace("SIC_Capstone_v2.ipynb", "__BASELINE_NOTEBOOK__")
    text = text.replace("SIC_Capstone_v3.ipynb", "__ADVANCED_NOTEBOOK__")
    text = re.sub(r"\bV2\b", "__BASELINE_VERSION_UPPER__", text)
    text = re.sub(r"\bv2\b", "__BASELINE_VERSION_LOWER__", text)
    text = re.sub(r"\bV3\b", "V2", text)
    text = re.sub(r"\bv3\b", "v2", text)
    text = text.replace("__BASELINE_VERSION_UPPER__", "V1")
    text = text.replace("__BASELINE_VERSION_LOWER__", "v1")
    text = text.replace("__BASELINE_NOTEBOOK__", "SIC_Capstone_v1.ipynb")
    text = text.replace("__ADVANCED_NOTEBOOK__", "SIC_Capstone_v2.ipynb")
    return text


def update_text(text: str) -> str:
    updated = shift_versions(text)
    updated = updated.replace(
        "Quy trình.pdf",
        "SIC_AI_Chapter 11. Starting an AI Project.pdf",
    )
    for old, new in NUMBER_REPLACEMENTS.items():
        updated = updated.replace(old, new)
    updated = updated.replace("best epoch 30", "best epoch 40")
    updated = updated.replace("Best epoch 30", "Best epoch 40")
    updated = updated.replace("HD95 khoảng 43 mm", "HD95 khoảng 31 mm")
    updated = updated.replace(
        "TC đạt Dice test cao nhất 0,7938, tiếp theo WT 0,8158 và ET 0,7476.",
        "WT đạt Dice test cao nhất 0,8158, tiếp theo TC 0,7938 và ET 0,7476.",
    )
    updated = updated.replace(
        "Danh mục hình và bảng chính",
        "Danh mục hình và bảng",
    )
    return updated


def update_all_text_nodes(doc: Document) -> tuple[int, dict[str, int]]:
    changed_nodes = 0
    counts: dict[str, int] = {}
    roots = [doc.element.body]
    for section in doc.sections:
        roots.extend([section.header._element, section.footer._element])

    seen_roots: set[int] = set()
    for root in roots:
        if id(root) in seen_roots:
            continue
        seen_roots.add(id(root))
        for node in root.iter(qn("w:t")):
            original = node.text or ""
            updated = update_text(original)
            if updated != original:
                node.text = updated
                changed_nodes += 1
                counts[original] = counts.get(original, 0) + 1
    return changed_nodes, counts


def keep_sic_phase_table_together(doc: Document) -> bool:
    for table in doc.tables:
        if not table.rows:
            continue
        header = [" ".join(cell.text.split()) for cell in table.rows[0].cells]
        if header != ["Pha", "Đầu ra", "Tiêu chí hoàn tất"]:
            continue
        for row_index, row in enumerate(table.rows):
            tr_pr = row._tr.get_or_add_trPr()
            if tr_pr.find(qn("w:cantSplit")) is None:
                tr_pr.append(OxmlElement("w:cantSplit"))
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = row_index < len(table.rows) - 1
        return True
    return False


def build_updated_diagrams() -> None:
    source = BUILDER.read_text(encoding="utf-8")
    source = shift_versions(source)
    source = source.replace(
        "unet = [0.7489, 0.7534, 0.7061, 0.7361]",
        "unet = [0.8158, 0.7938, 0.7476, 0.7858]",
    )
    namespace = {"__name__": "report_diagram_update", "__file__": str(BUILDER)}
    exec(compile(source, str(BUILDER), "exec"), namespace)
    if NEW_GENERATED.exists():
        shutil.rmtree(NEW_GENERATED)
    namespace["GENERATED"] = NEW_GENERATED
    namespace["create_diagrams"]()


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def replace_embedded_diagrams(docx_path: Path) -> dict[str, str]:
    diagram_names = [
        "split_comparison.png",
        "metric_comparison.png",
        "unet_architecture.png",
        "swin_architecture.png",
    ]
    old_hash_to_name = {
        sha256_bytes((OLD_GENERATED / name).read_bytes()): name
        for name in diagram_names
    }

    with zipfile.ZipFile(docx_path, "r") as source_zip:
        replacements: dict[str, bytes] = {}
        replaced_names: dict[str, str] = {}
        for member in source_zip.namelist():
            if not member.startswith("word/media/"):
                continue
            digest = sha256_bytes(source_zip.read(member))
            diagram_name = old_hash_to_name.get(digest)
            if diagram_name:
                replacements[member] = (NEW_GENERATED / diagram_name).read_bytes()
                replaced_names[diagram_name] = member

        missing = sorted(set(diagram_names) - set(replaced_names))
        if missing:
            raise RuntimeError(f"Could not locate embedded diagrams: {missing}")

        temp_path = docx_path.with_suffix(".media-update.docx")
        with zipfile.ZipFile(temp_path, "w") as target_zip:
            for item in source_zip.infolist():
                payload = replacements.get(item.filename, source_zip.read(item.filename))
                target_zip.writestr(item, payload)
    os.replace(temp_path, docx_path)
    return replaced_names


def audit(docx_path: Path) -> None:
    doc = Document(docx_path)
    text_parts = []
    for node in doc.element.body.iter(qn("w:t")):
        text_parts.append(node.text or "")
    full_text = "".join(text_parts)
    forbidden = [
        "SIC_Capstone_v3.ipynb",
        "Quy trình.pdf",
        "0,7361",
        "0,7489",
        "0,7534",
        "0,7061",
        "43,058",
    ]
    leftovers = [token for token in forbidden if token in full_text]
    version_leftovers = re.findall(r"\b[vV]3\b", full_text)
    if leftovers or version_leftovers:
        raise RuntimeError(
            f"Audit failed. Old tokens: {leftovers}; v3 count: {len(version_leftovers)}"
        )
    required = [
        "SIC_Capstone_v1.ipynb",
        "SIC_Capstone_v2.ipynb",
        "0,7858",
        "0,8052",
        "31,117",
        "SIC_AI_Chapter 11. Starting an AI Project.pdf",
    ]
    missing = [token for token in required if token not in full_text]
    if missing:
        raise RuntimeError(f"Audit failed. Missing updated tokens: {missing}")


def main() -> None:
    build_updated_diagrams()
    doc = Document(INPUT)
    changed_nodes, replacement_counts = update_all_text_nodes(doc)
    if changed_nodes < 60:
        raise RuntimeError(f"Unexpectedly few text updates: {changed_nodes}")
    if not keep_sic_phase_table_together(doc):
        raise RuntimeError("SIC phase table was not found")

    doc.core_properties.subject = "So sánh notebook v1 3D U-Net và notebook v2 Swin UNETR trên dữ liệu BraTS"
    doc.core_properties.comments = (
        "Báo cáo sử dụng kết quả có sẵn trong SIC_Capstone_v1.ipynb và "
        "SIC_Capstone_v2.ipynb; không huấn luyện lại mô hình."
    )
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    media = replace_embedded_diagrams(OUTPUT)
    audit(OUTPUT)
    print(f"Saved: {OUTPUT}")
    print(f"Size: {OUTPUT.stat().st_size} bytes")
    print(f"Changed text nodes: {changed_nodes}")
    print(f"Updated diagrams: {media}")
    print(f"Unique original text nodes changed: {len(replacement_counts)}")


if __name__ == "__main__":
    main()

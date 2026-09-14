# -*- coding: utf-8 -*-
from __future__ import annotations

import math
import shutil
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"D:\SIC_Capstone 2026")
TEMPLATE = ROOT / "doc" / "SIC_AI_Capstone Project_Final Report.docx"
OUTPUT = ROOT / "doc" / "SIC_AI_Capstone Project_Final Report_Brain Tumor Segmentation_Revised.docx"
WORK = ROOT / "report_work"
ASSETS = WORK / "assets"
GENERATED = WORK / "generated"

BLUE = "#1428A0"
LIGHT_BLUE = "#EAF0FF"
MID_BLUE = "#4B67C7"
DARK = "#171717"
GREY = "#666666"
LIGHT_GREY = "#F2F4F7"
WHITE = "#FFFFFF"
RED = "#B3261E"
GREEN = "#26734D"


def font(size: int, bold: bool = False):
    candidates = [
        Path(r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf"),
        Path(r"C:\Windows\Fonts\calibrib.ttf" if bold else r"C:\Windows\Fonts\calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def rounded_box(draw, box, fill, outline=BLUE, radius=24, width=4):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def arrow(draw, start, end, color=BLUE, width=6):
    draw.line([start, end], fill=color, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    head = 18
    left = (end[0] - head * math.cos(angle - math.pi / 6), end[1] - head * math.sin(angle - math.pi / 6))
    right = (end[0] - head * math.cos(angle + math.pi / 6), end[1] - head * math.sin(angle + math.pi / 6))
    draw.polygon([end, left, right], fill=color)


def centered_multiline(draw, box, text, fnt, fill=DARK, spacing=8):
    lines = text.split("\n")
    heights = []
    widths = []
    for line in lines:
        bb = draw.textbbox((0, 0), line, font=fnt)
        widths.append(bb[2] - bb[0])
        heights.append(bb[3] - bb[1])
    total = sum(heights) + spacing * (len(lines) - 1)
    y = box[1] + (box[3] - box[1] - total) / 2
    for line, w, h in zip(lines, widths, heights):
        x = box[0] + (box[2] - box[0] - w) / 2
        draw.text((x, y), line, font=fnt, fill=fill)
        y += h + spacing


def create_diagrams():
    GENERATED.mkdir(parents=True, exist_ok=True)
    title_f = font(42, True)
    box_f = font(30, True)
    small_f = font(24, False)

    # Design Thinking / problem-follow-up process
    img = Image.new("RGB", (1800, 850), "white")
    d = ImageDraw.Draw(img)
    d.text((70, 45), "Quy trình xác định và theo dõi vấn đề", font=title_f, fill=DARK)
    labels = [
        "Discover\nThu thập bằng chứng",
        "Define\nAi – Cái gì – Ở đâu",
        "Analyze\nNguyên nhân và tác động",
        "Ideate\nGiải pháp khả thi",
        "Prototype & Test\nĐo lường – cải tiến",
    ]
    colors = ["#EAF0FF", "#DDE6FF", "#CFDAFF", "#C2CEFF", "#B4C2FF"]
    xs = [70, 415, 760, 1105, 1450]
    for i, (x, label, c) in enumerate(zip(xs, labels, colors)):
        box = (x, 240, x + 280, 500)
        rounded_box(d, box, c)
        centered_multiline(d, box, label, box_f, fill=BLUE)
        if i < len(labels) - 1:
            arrow(d, (x + 285, 370), (xs[i + 1] - 12, 370))
    d.text((70, 625), "Vòng phản hồi: kết quả kiểm thử quay lại bước Define/Analyze nếu chưa đạt tiêu chí.", font=small_f, fill=GREY)
    arrow(d, (1590, 555), (225, 555), color=MID_BLUE, width=5)
    img.save(GENERATED / "problem_followup.png", quality=95)

    # End-to-end AI workflow
    img = Image.new("RGB", (1800, 980), "white")
    d = ImageDraw.Draw(img)
    d.text((70, 45), "Luồng xử lý đầu cuối của hệ thống", font=title_f, fill=DARK)
    rows = [
        ("4 chuỗi MRI NIfTI", "t1n, t1c, t2w, t2f"),
        ("Tiền xử lý", "clip cường độ – chuẩn hóa – ghép kênh"),
        ("Suy luận 3D", "3D U-Net hoặc Swin UNETR – sliding window"),
        ("Hậu xử lý", "nhãn WT, TC, ET – thống kê thể tích"),
        ("Giao diện và báo cáo", "3 mặt phẳng – mask NIfTI – PDF minh họa"),
    ]
    y = 165
    for i, (a, b) in enumerate(rows):
        box = (250, y, 1550, y + 125)
        rounded_box(d, box, LIGHT_BLUE if i % 2 == 0 else LIGHT_GREY, width=3)
        d.text((300, y + 25), a, font=box_f, fill=BLUE)
        d.text((820, y + 30), b, font=small_f, fill=DARK)
        if i < len(rows) - 1:
            arrow(d, (900, y + 130), (900, y + 170), width=5)
        y += 165
    img.save(GENERATED / "end_to_end_workflow.png", quality=95)

    # System architecture
    img = Image.new("RGB", (1800, 950), "white")
    d = ImageDraw.Draw(img)
    d.text((70, 45), "Kiến trúc logic của bản mẫu", font=title_f, fill=DARK)
    boxes = {
        "ui": (80, 250, 460, 520),
        "pre": (560, 160, 980, 355),
        "model": (1080, 160, 1630, 355),
        "post": (560, 485, 980, 680),
        "export": (1080, 485, 1630, 680),
    }
    texts = {
        "ui": "Gradio UI\nTải 4 ảnh MRI\nChọn mô hình",
        "pre": "Preprocessing\nChuẩn hóa và tensor 4 kênh",
        "model": "Inference engine\n3D U-Net / Swin UNETR",
        "post": "Postprocessing\nMask và thể tích vùng u",
        "export": "Output\nViewer – NIfTI – PDF",
    }
    for key, box in boxes.items():
        rounded_box(d, box, LIGHT_BLUE if key != "ui" else "#DCE6FF")
        centered_multiline(d, box, texts[key], box_f, fill=BLUE)
    arrow(d, (460, 330), (560, 255))
    arrow(d, (980, 255), (1080, 255))
    arrow(d, (1355, 355), (870, 485))
    arrow(d, (980, 585), (1080, 585))
    d.text((80, 795), "Ranh giới an toàn: kết quả là hỗ trợ nghiên cứu; người dùng phải kiểm tra chất lượng đầu vào và không sử dụng như chẩn đoán tự động.", font=small_f, fill=RED)
    img.save(GENERATED / "system_architecture.png", quality=95)

    # Split comparison
    img = Image.new("RGB", (1800, 900), "white")
    d = ImageDraw.Draw(img)
    d.text((70, 45), "So sánh phân chia dữ liệu giữa notebook", font=title_f, fill=DARK)
    datasets = [("v2 – 3D U-Net", [875, 188, 188]), ("v3 – Swin UNETR", [812, 204, 0])]
    colors = [BLUE, "#5D78D5", "#B9C4E9"]
    labels = ["Train", "Validation", "Held-out test"]
    maxv = 1251
    y = 230
    for name, vals in datasets:
        d.text((90, y - 65), name, font=box_f, fill=DARK)
        x = 90
        usable = 1550
        total = sum(vals)
        for val, c, lab in zip(vals, colors, labels):
            w = usable * val / maxv
            if val > 0:
                d.rectangle((x, y, x + w, y + 110), fill=c)
                centered_multiline(d, (x, y, x + w, y + 110), f"{lab}\n{val}", small_f, fill=WHITE if c != "#B9C4E9" else DARK)
            x += w
        d.text((90, y + 140), f"Tổng: {total} ca", font=small_f, fill=GREY)
        y += 300
    d.text((90, 800), "Lưu ý: v3 không có tập kiểm thử độc lập; so sánh chéo phải được xem là mô tả, không phải đối đầu công bằng.", font=small_f, fill=RED)
    img.save(GENERATED / "split_comparison.png", quality=95)

    # As-reported metric comparison
    img = Image.new("RGB", (1800, 1050), "white")
    d = ImageDraw.Draw(img)
    d.text((70, 45), "Chỉ số được báo cáo – khác giao thức đánh giá", font=title_f, fill=DARK)
    cats = ["WT", "TC", "ET", "Mean Dice"]
    unet = [0.7489, 0.7534, 0.7061, 0.7361]
    swin = [0.8498, 0.7461, 1.0000, 0.8653]
    base_y = 880
    chart_top = 160
    chart_h = 650
    d.line((130, base_y, 1670, base_y), fill=DARK, width=3)
    d.line((130, chart_top, 130, base_y), fill=DARK, width=3)
    for tick in range(0, 11, 2):
        val = tick / 10
        yy = base_y - chart_h * val
        d.line((120, yy, 1670, yy), fill="#D8DCE6", width=2)
        d.text((55, yy - 15), f"{val:.1f}", font=small_f, fill=GREY)
    group_w = 350
    bar_w = 95
    for i, cat in enumerate(cats):
        gx = 220 + i * group_w
        for j, (val, c) in enumerate([(unet[i], BLUE), (swin[i], "#7691F0")]):
            x = gx + j * 120
            y = base_y - chart_h * val
            d.rectangle((x, y, x + bar_w, base_y), fill=c)
            d.text((x - 5, y - 38), f"{val:.3f}", font=small_f, fill=DARK)
        d.text((gx + 55, base_y + 30), cat, font=box_f, fill=DARK)
    d.rectangle((1050, 85, 1100, 125), fill=BLUE)
    d.text((1120, 85), "3D U-Net – held-out test v2", font=small_f, fill=DARK)
    d.rectangle((1050, 135, 1100, 175), fill="#7691F0")
    d.text((1120, 135), "Swin UNETR – full validation v3", font=small_f, fill=DARK)
    d.text((130, 990), "ET của v3 = 1,000 vì nhãn và dự đoán ET đều rỗng ở toàn bộ 204 ca; không được diễn giải là phân đoạn ET hoàn hảo.", font=small_f, fill=RED)
    img.save(GENERATED / "metric_comparison.png", quality=95)

    # Resource comparison
    img = Image.new("RGB", (1800, 930), "white")
    d = ImageDraw.Draw(img)
    d.text((70, 45), "Quy mô mô hình và tệp triển khai", font=title_f, fill=DARK)
    items = [
        ("3D U-Net", 4.811, 55.1, BLUE),
        ("Swin UNETR", 15.706, 67.2, "#7691F0"),
    ]
    x0 = 430
    for i, (name, params, size, c) in enumerate(items):
        y = 230 + i * 300
        d.text((80, y + 40), name, font=box_f, fill=DARK)
        w = 900 * params / 16.0
        d.rectangle((x0, y, x0 + w, y + 90), fill=c)
        d.text((x0 + w + 25, y + 20), f"{params:.3f} triệu tham số", font=small_f, fill=DARK)
        w2 = 900 * size / 70.0
        d.rectangle((x0, y + 130, x0 + w2, y + 220), fill="#B8C5EF")
        d.text((x0 + w2 + 25, y + 150), f"xấp xỉ {size:.1f} MiB", font=small_f, fill=DARK)
    d.text((430, 825), "Thanh đậm: tham số học được   |   Thanh nhạt: kích thước checkpoint/tệp triển khai quan sát trong dự án", font=small_f, fill=GREY)
    img.save(GENERATED / "resource_comparison.png", quality=95)

    # 3D U-Net architecture used by notebook v2
    img = Image.new("RGB", (2000, 1120), "white")
    d = ImageDraw.Draw(img)
    d.text((70, 45), "Kiến trúc 3D U-Net trong notebook v2", font=title_f, fill=DARK)
    levels = [
        ("Input", "4 × 96³", 190),
        ("Enc 1", "16 × 48³", 410),
        ("Enc 2", "32 × 24³", 630),
        ("Enc 3", "64 × 12³", 850),
        ("Enc 4", "128 × 6³", 1070),
        ("Bottleneck", "256 × 6³", 1290),
    ]
    for idx, (name, shape, x) in enumerate(levels):
        y = 220 + idx * 75
        box = (x, y, x + 205, y + 125)
        rounded_box(d, box, LIGHT_BLUE if idx < 5 else "#DCE6FF", width=3)
        centered_multiline(d, box, f"{name}\n{shape}", small_f, fill=BLUE)
        if idx < len(levels) - 1:
            arrow(d, (x + 205, y + 62), (levels[idx + 1][2] - 15, y + 137), width=4)
    decoder = [
        ("Dec 4", "128 × 12³", 1510, 600),
        ("Dec 3", "64 × 24³", 1600, 455),
        ("Dec 2", "32 × 48³", 1690, 310),
        ("Output", "4 × 96³", 1780, 165),
    ]
    for idx, (name, shape, x, y) in enumerate(decoder):
        box = (x, y, x + 180, y + 120)
        rounded_box(d, box, "#F0F3FF", width=3)
        centered_multiline(d, box, f"{name}\n{shape}", small_f, fill=BLUE)
        if idx < len(decoder) - 1:
            nx, ny = decoder[idx + 1][2], decoder[idx + 1][3]
            arrow(d, (x + 90, y), (nx + 90, ny + 120), width=4)
    skip_pairs = [((520, 430), (1760, 370)), ((740, 580), (1670, 515)), ((960, 730), (1580, 660))]
    for start, end in skip_pairs:
        d.line([start, (start[0], 930), (end[0], 930), end], fill="#7691F0", width=4)
    d.text((90, 1010), "Mỗi tầng sử dụng residual unit, InstanceNorm và dropout 0,2; skip connection truyền đặc trưng encoder sang decoder.", font=small_f, fill=GREY)
    img.save(GENERATED / "unet_architecture.png", quality=95)

    # Swin UNETR architecture used by notebook v3
    img = Image.new("RGB", (2000, 1120), "white")
    d = ImageDraw.Draw(img)
    d.text((70, 45), "Kiến trúc Swin UNETR trong notebook v3", font=title_f, fill=DARK)
    stages = [
        ("Input", "4 × 96³", 90, 230),
        ("Patch embedding", "24 × 48³", 350, 230),
        ("Swin stage 1", "24 × 48³", 650, 230),
        ("Patch merge", "48 × 24³", 950, 230),
        ("Swin stage 2", "48 × 24³", 1250, 230),
        ("Deep stages", "96–384 kênh", 1550, 230),
    ]
    for idx, (name, shape, x, y) in enumerate(stages):
        box = (x, y, x + 235, y + 150)
        rounded_box(d, box, LIGHT_BLUE if "Swin" not in name else "#DCE6FF", width=3)
        centered_multiline(d, box, f"{name}\n{shape}", small_f, fill=BLUE)
        if idx < len(stages) - 1:
            arrow(d, (x + 235, y + 75), (stages[idx + 1][2] - 15, y + 75), width=4)
    decs = [
        ("UNETR decoder", "Upsample + conv", 1510, 650),
        ("Skip fusion", "Ghép đa tỷ lệ", 1060, 650),
        ("Segmentation head", "4 × 96³ logits", 600, 650),
    ]
    for idx, (name, shape, x, y) in enumerate(decs):
        box = (x, y, x + 300, y + 150)
        rounded_box(d, box, "#F0F3FF", width=3)
        centered_multiline(d, box, f"{name}\n{shape}", small_f, fill=BLUE)
        if idx < len(decs) - 1:
            nx = decs[idx + 1][2]
            arrow(d, (x, y + 75), (nx + 300 + 15, y + 75), width=4)
    for x in (760, 1360, 1670):
        d.line([(x, 380), (x, 590), (1220, 590), (1220, 650)], fill="#7691F0", width=4)
    d.text((90, 960), "W-MSA học quan hệ trong cửa sổ; SW-MSA dịch cửa sổ để trao đổi thông tin giữa các vùng; decoder phục hồi độ phân giải.", font=small_f, fill=GREY)
    d.text((90, 1010), "feature_size = 24; gradient checkpointing được bật để giảm bộ nhớ activation khi huấn luyện.", font=small_f, fill=GREY)
    img.save(GENERATED / "swin_architecture.png", quality=95)


def remove_body_content(doc: Document):
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill.lstrip("#"))


def set_cell_border(cell, color="B6BDCC", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color.lstrip("#"))


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run_font(run, name="Times New Roman", size=10.5, bold=None, color=None, italic=None):
    # The report uses one consistent typeface throughout.  Keep the name
    # argument for call-site compatibility, but normalize every run to the
    # user-requested Times New Roman family.
    name = "Times New Roman"
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))


def set_cell_text(cell, text, bold=False, color=DARK, size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run(str(text))
    set_run_font(r, "Times New Roman", size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_border(cell)


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(10.5)
    pf = normal.paragraph_format
    pf.line_spacing = 1.15
    pf.space_after = Pt(6)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, size, color, before, after in [
        ("Heading 1", 18, WHITE, 0, 0),
        ("Heading 2", 14, BLUE, 8, 6),
        ("Heading 3", 11.5, DARK, 7, 4),
    ]:
        st = styles[name]
        st.font.name = "Times New Roman"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    if "Figure Caption" not in [s.name for s in styles]:
        styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    cap = styles["Figure Caption"]
    cap.font.name = "Times New Roman"
    cap._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    cap.font.size = Pt(9)
    cap.font.italic = True
    cap.font.color.rgb = RGBColor.from_string(GREY.lstrip("#"))
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(3)
    cap.paragraph_format.space_after = Pt(8)
    cap.paragraph_format.keep_with_next = False

    if "Report Bullet" not in [s.name for s in styles]:
        styles.add_style("Report Bullet", WD_STYLE_TYPE.PARAGRAPH)
    bullet = styles["Report Bullet"]
    bullet.base_style = styles["Normal"]
    bullet.paragraph_format.left_indent = Cm(0.65)
    bullet.paragraph_format.first_line_indent = Cm(-0.35)
    bullet.paragraph_format.space_after = Pt(3)

    if "Source Note" not in [s.name for s in styles]:
        styles.add_style("Source Note", WD_STYLE_TYPE.PARAGRAPH)
    src = styles["Source Note"]
    src.font.name = "Times New Roman"
    src._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    src.font.size = Pt(8.5)
    src.font.color.rgb = RGBColor.from_string(GREY.lstrip("#"))
    src.paragraph_format.space_before = Pt(1)
    src.paragraph_format.space_after = Pt(6)


def configure_page(doc: Document):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3.0)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)
    section.different_first_page_header_footer = True


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "Cập nhật trường trong Microsoft Word"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def configure_footer(doc: Document):
    for section in doc.sections:
        footer = section.footer
        for p in footer.paragraphs:
            p._element.getparent().remove(p._element)
        p = footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run("Page ")
        set_run_font(r, "Arial", 8.5, color=GREY)
        add_field(p, "PAGE")
        r = p.add_run(" / ")
        set_run_font(r, "Arial", 8.5, color=GREY)
        add_field(p, "NUMPAGES")


def force_times_new_roman(doc: Document):
    """Normalize all paragraph/run and style font declarations in the report."""
    font_name = "Times New Roman"

    # Template styles can carry their own font declarations, so update every
    # paragraph style rather than only Normal and the three report headings.
    for style in doc.styles:
        try:
            style.font.name = font_name
            r_pr = style._element.get_or_add_rPr()
            r_fonts = r_pr.rFonts
            if r_fonts is None:
                r_fonts = OxmlElement("w:rFonts")
                r_pr.insert(0, r_fonts)
            for key in ("ascii", "hAnsi", "eastAsia", "cs"):
                r_fonts.set(qn(f"w:{key}"), font_name)
        except Exception:
            # Some latent/built-in styles do not expose a writable rPr.
            continue

    def normalize_paragraph(paragraph):
        for run in paragraph.runs:
            run.font.name = font_name
            r_pr = run._element.get_or_add_rPr()
            r_fonts = r_pr.rFonts
            if r_fonts is None:
                r_fonts = OxmlElement("w:rFonts")
                r_pr.insert(0, r_fonts)
            for key in ("ascii", "hAnsi", "eastAsia", "cs"):
                r_fonts.set(qn(f"w:{key}"), font_name)

    def normalize_table(table):
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    normalize_paragraph(paragraph)
                for nested in cell.tables:
                    normalize_table(nested)

    for paragraph in doc.paragraphs:
        normalize_paragraph(paragraph)
    for table in doc.tables:
        normalize_table(table)

    # Preserve the template's header/footer content and fields, but ensure
    # their visible text follows the same typeface.
    for section in doc.sections:
        containers = [
            section.header, section.first_page_header,
            section.footer, section.first_page_footer,
        ]
        for container in containers:
            for paragraph in container.paragraphs:
                normalize_paragraph(paragraph)
            for table in container.tables:
                normalize_table(table)


def normalize_package_font_names(path: Path):
    """Catch fonts inside template content controls not exposed by python-docx."""
    replacements = {
        b"SamsungOne 400": b"Times New Roman",
        b"SamsungOne-700": b"Times New Roman",
        b"Samsung Sharp Sans": b"Times New Roman",
        b"Arial": b"Times New Roman",
    }
    tmp = path.with_name(path.stem + ".fontfix.docx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename.startswith("word/") and info.filename.endswith(".xml") and info.filename != "word/theme/theme1.xml":
                for old, new in replacements.items():
                    data = data.replace(old, new)
            target.writestr(info, data)
    tmp.replace(path)


def shade_paragraph(paragraph, fill=BLUE):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill.lstrip("#"))
    p_pr.append(shd)
    spacing = paragraph.paragraph_format
    # Align the chapter band exactly with the document text area.  Equal
    # zero indents prevent the blue rectangle from appearing shifted relative
    # to the tables and body paragraphs at either page edge.
    spacing.left_indent = Cm(0)
    spacing.right_indent = Cm(0)
    spacing.first_line_indent = Cm(0)
    spacing.space_before = Pt(0)
    spacing.space_after = Pt(10)
    spacing.line_spacing = 1.25


def add_chapter(doc, title):
    p = doc.add_paragraph(style="Heading 1")
    p.add_run(title)
    shade_paragraph(p)
    return p


def add_heading(doc, title, level=2):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(title)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_run_font(r, bold=True)
        r = p.add_run(text[len(bold_lead):])
        set_run_font(r)
    else:
        r = p.add_run(text)
        set_run_font(r)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.widow_control = True
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="Report Bullet")
        r = p.add_run("•  " + item)
        set_run_font(r)


def add_numbered(doc, items):
    for i, item in enumerate(items, 1):
        p = doc.add_paragraph(style="Report Bullet")
        r = p.add_run(f"{i}.  {item}")
        set_run_font(r)


def add_table(doc, headers, rows, widths=None, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    prevent_row_split(hdr)
    for i, text in enumerate(headers):
        if widths:
            hdr.cells[i].width = Cm(widths[i])
        set_cell_shading(hdr.cells[i], BLUE)
        set_cell_text(hdr.cells[i], text, bold=True, color=WHITE, size=font_size, align=WD_ALIGN_PARAGRAPH.CENTER)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for i, value in enumerate(row):
            if widths:
                cells[i].width = Cm(widths[i])
            if ridx % 2 == 1:
                set_cell_shading(cells[i], LIGHT_GREY)
            set_cell_text(cells[i], value, size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_picture(doc, path, caption, max_width=6.2, max_height=6.0):
    path = Path(path)
    with Image.open(path) as im:
        w, h = im.size
    ratio = w / h
    width = max_width
    height = width / ratio
    if height > max_height:
        height = max_height
        width = height * ratio
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    r = p.add_run()
    r.add_picture(str(path), width=Inches(width), height=Inches(height))
    cp = doc.add_paragraph(style="Figure Caption")
    cp.add_run(caption)
    return p


def add_source(doc, text):
    p = doc.add_paragraph(style="Source Note")
    p.add_run(text)


def next_page(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def add_topic_page(doc, title, paragraphs, bullets=None, table=None, figure=None, caption=None, source=None, level=2):
    add_heading(doc, title, level=level)
    for paragraph in paragraphs:
        add_body(doc, paragraph)
    if bullets:
        add_bullets(doc, bullets)
    if table:
        add_table(doc, *table)
    if figure:
        add_picture(doc, figure, caption or Path(figure).stem)
    if source:
        add_source(doc, source)


def cover_page(doc):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("AI Course")
    set_run_font(r, "SamsungOne-700", 18, bold=True, color=DARK)
    p.paragraph_format.space_after = Pt(20)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("Capstone Project\nFinal Report")
    set_run_font(r, "Samsung Sharp Sans", 34, bold=True, color=BLUE)
    p.paragraph_format.space_after = Pt(14)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("HỆ THỐNG HỖ TRỢ PHÂN ĐOẠN KHỐI U NÃO\nTRÊN ẢNH MRI 3D")
    set_run_font(r, "SamsungOne-700", 20, bold=True, color=DARK)
    p.paragraph_format.space_after = Pt(22)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("So sánh 3D U-Net và Swin UNETR")
    set_run_font(r, "Arial", 15, bold=True, color=MID_BLUE)
    p.paragraph_format.space_after = Pt(54)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("Báo cáo dự án cuối khóa – Samsung Innovation Campus 2026")
    set_run_font(r, "Arial", 11, color=GREY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("Ngày hoàn thiện: 14/09/2026")
    set_run_font(r, "Arial", 10, color=GREY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(65)
    r = p.add_run("Chỉ phục vụ nghiên cứu và đào tạo. Không sử dụng làm căn cứ duy nhất cho chẩn đoán hoặc điều trị.")
    set_run_font(r, "Arial", 9, italic=True, color=RED)


def add_architecture_depth(doc):
    add_heading(doc, "2.2.4.1 Từ U-Net gốc đến 3D U-Net", 3)
    add_body(doc, "U-Net gốc do Ronneberger và cộng sự đề xuất cho phân đoạn ảnh y sinh 2D [25]. Kiến trúc có hai nhánh đối xứng. Nhánh co rút liên tiếp áp dụng convolution và downsampling để tăng mức trừu tượng; nhánh mở rộng upsample đặc trưng để khôi phục bản đồ phân đoạn. Điểm cốt lõi là skip connection: feature map có độ phân giải cao ở encoder được ghép với decoder cùng mức, giúp decoder phục hồi biên và cấu trúc nhỏ thay vì chỉ dựa vào biểu diễn bottleneck.")
    add_body(doc, "3D U-Net thay convolution, pooling và upsampling 2D bằng phép toán 3D để mỗi kernel quan sát đồng thời chiều sâu, chiều cao và chiều rộng [5]. Bối cảnh theo trục lát cắt được đưa trực tiếp vào feature map. Đổi lại, số phần tử activation tăng nhanh theo D×H×W; đây là lý do notebook dùng patch 96³ và batch nhỏ. Bản gốc 3D U-Net được thiết kế để học từ chú thích thể tích thưa; dự án này sử dụng biến thể MONAI cho nhãn voxel đầy đủ của BraTS.")
    add_table(doc, ["Thành phần", "U-Net 2D gốc", "3D U-Net", "Cấu hình v2"], [
        ("Đầu vào", "Ảnh H×W", "Thể tích D×H×W", "4×96×96×96"),
        ("Kernel", "k_h×k_w", "k_d×k_h×k_w", "3D convolution"),
        ("Skip connection", "Ghép feature 2D", "Ghép feature 3D", "Encoder–decoder MONAI"),
        ("Chuẩn hóa", "Theo implementation", "Theo implementation", "InstanceNorm"),
        ("Regularization", "Augmentation", "Augmentation/normalization", "Dropout 0,2"),
    ], [3.5, 4.0, 4.1, 4.0], font_size=8.7)

    add_heading(doc, "2.2.4.2 Phép tích chập ba chiều", 3)
    add_body(doc, "Tại một vị trí (i,j,k), đầu ra của kênh c_o là tổng có trọng số trên mọi kênh đầu vào và mọi vị trí kernel. Kernel W có dạng C_out×C_in×K_d×K_h×K_w. Trọng số được dùng lại trên toàn thể tích, tạo tính tương đương tịnh tiến và inductive bias cục bộ của CNN.")
    add_picture(doc, GENERATED / "eq_conv3d.png", "Công thức 1. Phép tích chập 3D cho một kênh đầu ra.", max_width=6.0, max_height=1.25)
    add_body(doc, "Với stride S, padding P, dilation δ và kernel K, kích thước mỗi chiều đầu ra là floor((L+2P−δ(K−1)−1)/S+1). Stride 2 làm mỗi chiều giảm gần một nửa, nên số voxel feature giảm khoảng tám lần. Encoder đổi phần tiết kiệm không gian này lấy số kênh lớn hơn để lưu biểu diễn ngữ nghĩa.")
    add_table(doc, ["Tham số", "Ý nghĩa", "Tác động khi tăng"], [
        ("K", "Kích thước kernel", "Receptive field lớn hơn, FLOPs tăng"),
        ("S", "Stride", "Giảm nhanh độ phân giải"),
        ("P", "Padding", "Kiểm soát kích thước và biên"),
        ("C_in/C_out", "Số kênh", "Năng lực biểu diễn và số tham số tăng"),
        ("δ", "Dilation", "Mở rộng receptive field không tăng kernel"),
    ], [3.2, 6.0, 6.4])

    add_heading(doc, "2.2.4.3 Residual unit InstanceNorm và dropout", 3)
    add_body(doc, "Residual unit học phần dư F(x;W) và cộng lại đầu vào x. Đường tắt giúp gradient truyền qua mạng sâu dễ hơn; khi kích thước hoặc số kênh thay đổi, projection có thể được dùng để làm hai tensor tương thích. Trong v2, mỗi tầng có hai residual unit, vì vậy feature được tinh chỉnh nhiều lần trước khi chuyển mức.")
    add_picture(doc, GENERATED / "eq_residual_norm.png", "Công thức 2. Residual connection và Instance Normalization.", max_width=5.6, max_height=1.7)
    add_body(doc, "InstanceNorm tính trung bình μ và phương sai σ² theo từng mẫu và từng kênh trên các voxel không gian. Cách này ổn định hơn BatchNorm khi batch size bằng 1 như v2. Hai tham số học γ và β khôi phục khả năng co giãn/dịch đặc trưng. Dropout 0,2 đặt ngẫu nhiên một phần activation về 0 ở train time; ở inference dropout tắt và toàn bộ activation được sử dụng.")
    add_body(doc, "Ba cơ chế giải quyết ba vấn đề khác nhau: residual hỗ trợ tối ưu mạng sâu, InstanceNorm giảm biến thiên thang đo trong batch nhỏ, còn dropout hạn chế phụ thuộc quá mức vào một tập activation. Chúng không thay thế augmentation hoặc validation; mức dropout quá lớn có thể làm mất tín hiệu của vùng ET nhỏ.")

    add_heading(doc, "2.2.4.4 Luồng tensor của 3D U-Net v2", 3)
    add_picture(doc, GENERATED / "unet_architecture.png", "Hình 3. Kiến trúc khái niệm 3D U-Net theo cấu hình kênh của notebook v2.", max_width=6.2, max_height=4.2)
    add_body(doc, "Tensor đầu vào có bốn kênh. Chuỗi kênh 16–32–64–128–256 cho thấy số kênh tăng gấp đôi khi độ phân giải giảm. Decoder upsample, kết hợp feature encoder và dự đoán bốn logit nội bộ tương ứng background cùng ba lớp u sau mapping. Softmax chuyển logit thành xác suất; argmax tạo nhãn rời rạc.")
    add_table(doc, ["Mức", "Không gian khái niệm", "Số kênh", "Chức năng"], [
        ("Input", "96³", "4", "Bốn modality MRI"),
        ("Encoder 1", "≈48³", "16", "Biên và texture cục bộ"),
        ("Encoder 2", "≈24³", "32", "Cấu trúc trung gian"),
        ("Encoder 3", "≈12³", "64", "Ngữ cảnh tổn thương"),
        ("Encoder 4", "≈6³", "128", "Đặc trưng ngữ nghĩa sâu"),
        ("Bottleneck", "Mức thấp nhất", "256", "Ngữ cảnh cô đọng"),
        ("Decoder", "Khôi phục đến 96³", "128→16", "Kết hợp skip và định vị"),
        ("Head", "96³", "4", "Logit cho voxel"),
    ], [2.7, 4.0, 3.0, 5.9])
    add_source(doc, "Các kích thước không gian là sơ đồ khái niệm suy từ ROI 96³ và stride; shape chính xác phụ thuộc implementation MONAI và padding của từng block.")

    add_heading(doc, "2.2.4.5 Decoder và skip connection", 3)
    add_body(doc, "Ở decoder, phép upsampling tăng độ phân giải feature map. Feature đã upsample được nối theo chiều kênh với feature encoder cùng mức. Nếu chỉ dùng bottleneck, vị trí biên nhỏ đã bị mất qua nhiều lần downsampling; skip connection cung cấp lại chi tiết hình học. Convolution sau phép nối học cách chọn thông tin encoder nào hữu ích và loại bỏ nhiễu.")
    add_body(doc, "Skip connection cũng tạo đường truyền gradient ngắn từ loss tới encoder. Tuy nhiên nó có thể mang theo texture không liên quan; decoder cần đủ năng lực lọc. Với MRI, alignment không gian giữa các modality và mask đặc biệt quan trọng vì skip connection bảo toàn sai lệch vị trí nếu preprocessing sai.")
    add_table(doc, ["Thao tác", "Thay đổi tensor", "Mục đích"], [
        ("Upsample", "Tăng D,H,W", "Khôi phục độ phân giải"),
        ("Concatenate", "Tăng số kênh", "Đưa chi tiết encoder vào decoder"),
        ("Convolution", "Trộn và lọc kênh", "Hợp nhất ngữ nghĩa và vị trí"),
        ("Segmentation head", "C→4 logits", "Dự đoán lớp cho từng voxel"),
    ], [3.5, 5.1, 7.0])

    add_heading(doc, "2.2.4.6 Nguồn gốc Swin Transformer và Swin UNETR", 3)
    add_body(doc, "Vision Transformer chuẩn biến ảnh thành chuỗi token và dùng self-attention toàn cục. Chi phí attention tăng bình phương theo số token, nên áp dụng trực tiếp cho thể tích 3D độ phân giải cao rất tốn bộ nhớ. Swin Transformer giải quyết bằng cửa sổ cục bộ, kiến trúc phân cấp và shifted window [26].")
    add_body(doc, "Swin UNETR lấy Swin Transformer làm encoder cho ảnh y sinh 3D, sau đó dùng decoder kiểu U-Net để khôi phục segmentation [6]. Feature ở nhiều stage được đưa qua skip connection đến decoder. Kiến trúc giữ khả năng học ngữ cảnh của attention nhưng vẫn tạo đầu ra voxel-level.")
    add_table(doc, ["Kiến trúc", "Đơn vị cơ bản", "Không gian đặc trưng", "Vai trò"], [
        ("ViT", "Global self-attention", "Thường một mức token", "Ngữ cảnh toàn cục"),
        ("Swin Transformer", "W-MSA và SW-MSA", "Phân cấp", "Attention hiệu quả hơn"),
        ("UNETR", "Transformer encoder", "Skip sang CNN decoder", "Segmentation 3D"),
        ("Swin UNETR", "Swin encoder", "Đa tỷ lệ", "Kết hợp shifted window và U-Net"),
    ], [3.2, 4.7, 3.6, 4.1])

    add_heading(doc, "2.2.4.7 Patch embedding và biểu diễn phân cấp", 3)
    add_body(doc, "Patch embedding chia volume thành các khối nhỏ không chồng lấp và chiếu mỗi patch thành vector feature. Với patch size 2³, ROI 96³ trở thành lưới 48³ token ở stage đầu. feature_size=24 nghĩa là mỗi token ban đầu có 24 kênh biểu diễn. Patch merging sau đó gom token lân cận, giảm mỗi chiều không gian và tăng số kênh.")
    add_body(doc, "Cấu trúc phân cấp tương tự pyramid của CNN: stage nông giữ vị trí, stage sâu giữ ngữ nghĩa. Sự khác biệt là phép trộn token được thực hiện bằng attention trong cửa sổ và MLP thay vì chỉ convolution. Decoder sử dụng feature từ nhiều stage để phục hồi biên.")
    add_table(doc, ["Stage khái niệm", "Độ phân giải", "Kênh theo feature size 24", "Thông tin ưu thế"], [
        ("Patch embed", "48³", "24", "Chi tiết cục bộ"),
        ("Stage 2", "24³", "48", "Mẫu mô trung gian"),
        ("Stage 3", "12³", "96", "Ngữ cảnh vùng u"),
        ("Stage 4", "6³", "192", "Ngữ nghĩa sâu"),
        ("Deep feature", "3³ hoặc mức sâu nhất", "384", "Bối cảnh cô đọng"),
    ], [3.6, 3.3, 4.5, 4.2])
    add_source(doc, "Bảng mô tả quy luật phân cấp chuẩn của Swin UNETR; cần dùng forward hook nếu muốn xác nhận shape chính xác của checkpoint cụ thể.")

    add_heading(doc, "2.2.4.8 Scaled dot product self attention", 3)
    add_body(doc, "Với ma trận token X, ba phép chiếu tuyến tính tạo Query Q, Key K và Value V. Tích QKᵀ đo mức liên quan giữa các token. Chia cho √d_k giữ độ lớn logit ổn định khi chiều key tăng; Softmax tạo trọng số tổng bằng 1. Ma trận B là relative position bias giúp mô hình biết quan hệ vị trí bên trong cửa sổ.")
    add_picture(doc, GENERATED / "eq_attention.png", "Công thức 3. Scaled dot product attention có relative position bias.", max_width=5.8, max_height=1.6)
    add_body(doc, "Multi-head attention chia feature thành nhiều head để học các kiểu quan hệ khác nhau, sau đó ghép lại. Một head có thể ưu tiên tương phản T1c quanh ET, head khác ưu tiên vùng FLAIR lan tỏa. Đây là diễn giải chức năng, không phải bằng chứng attention map của checkpoint; muốn khẳng định cần trực quan hóa hoặc attribution riêng.")

    add_heading(doc, "2.2.4.9 Window attention và shifted window", 3)
    add_body(doc, "W-MSA chỉ tính attention giữa token trong cùng cửa sổ. Nếu toàn volume có N token và mỗi cửa sổ có M³ token, chi phí attention xấp xỉ O(N·M³·d) thay vì O(N²·d) của attention toàn cục. Lợi ích càng lớn khi N tăng mạnh trong dữ liệu 3D.")
    add_body(doc, "Giới hạn của cửa sổ cố định là token ở hai cửa sổ kề nhau không trao đổi trực tiếp. Block tiếp theo dịch cửa sổ một nửa kích thước, tạo SW-MSA và đưa các token trước đó thuộc cửa sổ khác vào cùng nhóm. Hai block xen kẽ tạo đường truyền thông tin qua toàn feature map mà không trả chi phí attention toàn cục.")
    add_picture(doc, GENERATED / "eq_swin_block.png", "Công thức 4. Một cặp block W-MSA và SW-MSA với residual connection.", max_width=6.0, max_height=1.5)
    add_body(doc, "LayerNorm được áp dụng trước attention và MLP theo pre-norm. Residual connection bao quanh cả attention lẫn MLP. Cấu trúc này giúp tối ưu Transformer sâu; stochastic depth hoặc dropout có thể được dùng tùy cấu hình.")

    add_heading(doc, "2.2.4.10 Luồng tensor Swin UNETR v3", 3)
    add_picture(doc, GENERATED / "swin_architecture.png", "Hình 4. Kiến trúc khái niệm Swin UNETR theo feature size 24 của notebook v3.", max_width=6.2, max_height=4.1)
    add_body(doc, "Encoder tạo feature ở nhiều độ phân giải. Decoder UNETR upsample feature sâu, ghép feature stage nông qua skip connection và dùng convolution 3D để tinh chỉnh. Segmentation head cuối cùng tạo bốn logit voxel. Vì ROI và output contract giống U-Net, hai mô hình có thể dùng chung phần lớn preprocessing và postprocessing nếu label mapping được khóa.")
    add_body(doc, "Gradient checkpointing trong v3 không thay đổi forward output. Trong train, một số activation không được lưu; backward sẽ tính lại chúng để tiết kiệm VRAM. Đánh đổi là thời gian compute tăng. Cấu hình này giải thích vì sao batch 2 có thể khả thi với mô hình 15,7 triệu tham số nhưng không cung cấp benchmark tốc độ.")

    add_heading(doc, "2.2.4.11 Kiến trúc gốc và cấu hình dự án", 3)
    add_body(doc, "Tên kiến trúc không đủ để tái lập. Bản gốc công bố một họ thiết kế, còn implementation cụ thể phụ thuộc image size, patch size, feature size, số block, head, normalization và decoder. Báo cáo tách rõ kiến trúc gốc khỏi cấu hình trong notebook để tránh suy ra tham số không xuất hiện trong artifact.")
    add_table(doc, ["Thuộc tính", "Công trình gốc", "Điều chỉnh trong dự án", "Lý do"], [
        ("3D U-Net input", "Thể tích 3D", "4 modality, patch 96³", "Giới hạn VRAM và BraTS"),
        ("3D U-Net channels", "Tùy thiết kế", "16–32–64–128–256", "Baseline gọn"),
        ("Swin UNETR feature", "Cấu hình theo paper", "feature_size 24", "Cân bằng bộ nhớ/năng lực"),
        ("Checkpointing", "Tùy implementation", "Bật ở v3", "Giảm activation memory"),
        ("Output", "Segmentation classes", "4 lớp nội bộ", "Background + nhãn BraTS"),
    ], [3.3, 4.2, 4.7, 3.4], font_size=8.5)


def add_optimization_depth(doc):
    add_heading(doc, "2.2.5.1 Logit softmax và xác suất voxel", 3)
    add_body(doc, "Mô hình trả về logit z_i,c cho mỗi voxel i và lớp c. Softmax chuyển logit thành p_i,c=exp(z_i,c)/Σ_k exp(z_i,k). Các xác suất trên một voxel tổng bằng 1. Trong train, loss dùng xác suất mềm để tạo gradient; trong inference, argmax chọn lớp có xác suất lớn nhất.")
    add_body(doc, "Dùng softmax phù hợp khi mỗi voxel thuộc đúng một lớp nội bộ. WT, TC và ET là vùng tổng hợp được xây từ nhãn cơ sở sau dự đoán, không phải ba kênh độc lập chồng lấp trong formulation này. Nếu đổi sang sigmoid multi-label, loss và postprocessing phải được thiết kế lại.")
    add_table(doc, ["Đại lượng", "Ký hiệu", "Miền giá trị", "Vai trò"], [
        ("Logit", "z_i,c", "Số thực", "Điểm chưa chuẩn hóa"),
        ("Probability", "p_i,c", "[0,1]", "Đầu vào Dice/CE"),
        ("One-hot target", "g_i,c", "{0,1}", "Nhãn thật"),
        ("Predicted label", "argmax_c p_i,c", "0…C−1", "Mask rời rạc"),
    ], [3.6, 3.0, 3.5, 5.5])

    add_heading(doc, "2.2.5.2 Dice Loss", 3)
    add_body(doc, "Dice Loss tối ưu trực tiếp mức chồng lấp. Tử số đo phần giao mềm giữa xác suất và nhãn; mẫu số chuẩn hóa theo quy mô hai mask. Thành phần ε tránh chia cho 0 và ổn định số học. Loss được tính theo lớp rồi lấy trung bình trên C lớp được chọn.")
    add_picture(doc, GENERATED / "eq_dicece.png", "Công thức 5. Dice Loss, Cross Entropy và hàm kết hợp DiceCE.", max_width=6.1, max_height=2.0)
    add_body(doc, "include_background=False loại kênh nền khỏi trung bình Dice. Nếu giữ nền, số voxel nền rất lớn có thể làm loss có vẻ thấp dù vùng u sai. Dice Loss vẫn có hạn chế: gradient phụ thuộc tổng thể mask và cách xử lý lớp rỗng; một batch không có ET cung cấp ít tín hiệu học cho ET.")
    add_table(doc, ["Biến", "Ý nghĩa"], [
        ("p_i,c", "Xác suất mô hình tại voxel i, lớp c"),
        ("g_i,c", "Nhãn one-hot tại voxel i, lớp c"),
        ("C", "Số lớp đưa vào trung bình; không gồm background trong v2"),
        ("ε", "Hằng số làm trơn và tránh mẫu số bằng 0"),
        ("λ_D", "Trọng số thành phần Dice"),
    ], [4.2, 11.4])

    add_heading(doc, "2.2.5.3 Cross Entropy", 3)
    add_body(doc, "Cross Entropy phạt âm log xác suất của lớp đúng. Nếu mô hình gán p nhỏ cho lớp thật, −log(p) lớn và gradient mạnh. CE hoạt động ở từng voxel nên cung cấp tín hiệu cục bộ ổn định ngay cả khi overlap toàn vùng còn thấp ở đầu huấn luyện.")
    add_body(doc, "Nhược điểm là lớp có nhiều voxel đóng góp nhiều hơn. Background và WT lớn có thể chi phối gradient nếu không dùng class weight hoặc sampling phù hợp. DiceCE kết hợp hai góc nhìn: CE tối ưu phân loại voxel, Dice tối ưu overlap theo vùng. Trong MONAI, chi tiết squared prediction, smoothing và reduction phải được ghi cùng version để tái lập chính xác.")
    add_table(doc, ["Tình huống", "CE", "Dice", "Ý nghĩa kết hợp"], [
        ("Đầu train", "Gradient voxel rõ", "Overlap có thể rất thấp", "CE hỗ trợ khởi động"),
        ("Lớp nhỏ", "Dễ bị lớp lớn áp đảo", "Chuẩn hóa theo mask", "Dice cân bằng vùng"),
        ("Biên sai", "Phạt từng voxel", "Phạt giao/hợp", "Hai tín hiệu bổ sung"),
        ("Mask rỗng", "Vẫn phạt false positive", "Phụ thuộc smooth", "Cần policy rõ"),
    ], [3.6, 4.0, 4.0, 4.0])

    add_heading(doc, "2.2.5.4 Ví dụ số của Dice và IoU", 3)
    add_body(doc, "Giả sử một vùng có 7.200 voxel true positive, 800 false positive và 1.000 false negative. Dice=2×7.200/(2×7.200+800+1.000)=0,8889. IoU=7.200/(7.200+800+1.000)=0,8000. Cả hai mô tả cùng confusion set nhưng Dice luôn lớn hơn IoU khi giá trị nằm giữa 0 và 1.")
    add_body(doc, "Ví dụ cho thấy một con số overlap không nói lỗi nằm ở đâu. Hai mask có cùng Dice có thể khác hẳn về biên: một mask lệch nhẹ toàn vùng, mask khác có một cụm false positive ở xa. HD95 được dùng để bổ sung thông tin khoảng cách hình học.")
    add_table(doc, ["Thành phần", "Số voxel", "Diễn giải"], [
        ("TP", "7.200", "Dự đoán u và ground truth là u"),
        ("FP", "800", "Dự đoán thừa"),
        ("FN", "1.000", "Bỏ sót"),
        ("Dice", "0,8889", "Overlap theo trung bình điều hòa"),
        ("IoU", "0,8000", "Tỷ lệ giao trên hợp"),
    ], [4.0, 3.2, 8.4])

    add_heading(doc, "2.2.5.5 Lan truyền ngược", 3)
    add_body(doc, "Backpropagation dùng quy tắc chuỗi để tính gradient của loss theo từng tham số. Gradient đi từ segmentation head qua decoder, skip connection và encoder. Autograd lưu graph của forward; vì activation 3D lớn, bộ nhớ train thường bị chi phối bởi activation hơn là riêng số tham số.")
    add_picture(doc, GENERATED / "eq_backprop.png", "Công thức 6. Gradient của loss qua chuỗi lớp theo quy tắc dây chuyền.", max_width=5.9, max_height=1.15)
    add_body(doc, "Mỗi bước tối ưu theo thứ tự: zero gradient, forward, tính DiceCE, backward, optimizer.step và scheduler.step theo thiết kế. Mixed precision có thể giảm bộ nhớ nhưng cần gradient scaler để tránh underflow. Notebook phải ghi rõ nếu dùng AMP; báo cáo không mặc định một tính năng không thấy trong artifact.")
    add_numbered(doc, [
        "Lấy batch patch ảnh và nhãn.",
        "Tạo logit bằng forward pass.",
        "Tính DiceCE trên các lớp mục tiêu.",
        "Tính gradient bằng backward pass.",
        "AdamW cập nhật tham số và áp weight decay tách rời.",
        "Cập nhật learning rate theo scheduler khi đến đúng đơn vị epoch/step.",
    ])

    add_heading(doc, "2.2.5.6 AdamW và weight decay tách rời", 3)
    add_body(doc, "Adam duy trì trung bình động bậc nhất m_t của gradient và bậc hai v_t của bình phương gradient. Bias correction tạo m̂_t và v̂_t ở các bước đầu. Mỗi tham số nhận learning rate hiệu dụng được chuẩn hóa bởi √v̂_t+ε, giúp tối ưu thích nghi theo độ lớn gradient [28].")
    add_picture(doc, GENERATED / "eq_adamw.png", "Công thức 7. Cập nhật AdamW với weight decay tách rời.", max_width=5.8, max_height=2.0)
    add_body(doc, "AdamW tách weight decay khỏi gradient loss [27]. Thành phần (1−η_tλ)θ_t co trọng số trực tiếp, thay vì thêm L2 penalty vào gradient rồi để Adam chuẩn hóa. Trong notebook, λ=10⁻⁴. Weight decay quá thấp ít regularize; quá cao làm underfit. Bias và tham số normalization thường có thể được loại khỏi decay trong thiết kế chi tiết, nhưng phải kiểm tra optimizer parameter groups thực tế.")
    add_table(doc, ["Ký hiệu", "Giá trị thường dùng", "Trong dự án"], [
        ("η_t", "Learning rate tại bước t", "2e-4; v2 resume 2e-5"),
        ("β₁", "Moment bậc nhất", "Theo mặc định AdamW nếu notebook không đổi"),
        ("β₂", "Moment bậc hai", "Theo mặc định AdamW nếu notebook không đổi"),
        ("λ", "Weight decay", "1e-4"),
        ("ε", "Ổn định mẫu số", "Theo implementation"),
    ], [3.2, 6.2, 6.2])

    add_heading(doc, "2.2.5.7 Learning rate và cosine annealing", 3)
    add_body(doc, "Learning rate quyết định độ dài bước cập nhật. Giá trị lớn giúp tiến nhanh nhưng có thể vượt vùng tối ưu; giá trị nhỏ ổn định nhưng chậm. v2 dùng 2e-4 ở giai đoạn đầu và 2e-5 khi resume. Nếu resume cả optimizer state và scheduler, learning rate phải tương ứng với epoch; nếu chỉ nạp model weights, run có hành vi khác.")
    add_picture(doc, GENERATED / "eq_cosine.png", "Công thức 8. Cosine annealing learning rate.", max_width=5.5, max_height=1.1)
    add_body(doc, "v3 dùng cosine annealing: η_t giảm trơn từ η_max về η_min theo nửa chu kỳ cosine. Đầu run cập nhật lớn hơn để khám phá, cuối run cập nhật nhỏ hơn để tinh chỉnh. T_max phải khớp số epoch hoặc bước scheduler. Gọi scheduler theo batch thay vì theo epoch sẽ làm lịch giảm nhanh sai thiết kế.")
    add_table(doc, ["Cách điều khiển LR", "Ưu điểm", "Rủi ro"], [
        ("Constant", "Đơn giản", "Khó cân bằng đầu/cuối run"),
        ("Step decay", "Dễ giải thích", "Thay đổi đột ngột"),
        ("Cosine", "Giảm trơn", "Nhạy T_max và đơn vị step"),
        ("Reduce on plateau", "Theo metric", "Nhạy nhiễu validation"),
    ], [4.0, 5.5, 6.1])

    add_heading(doc, "2.2.5.8 Early stopping và checkpoint selection", 3)
    add_body(doc, "Early stopping theo dõi metric validation; nếu không cải thiện trong patience epoch thì dừng. v3 đặt patience 5. Cơ chế này giảm lãng phí compute và overfitting, nhưng chỉ đáng tin khi validation đủ lớn và đại diện. Fast validation hai ca quá nhiễu để dùng như tiêu chí cuối.")
    add_body(doc, "Checkpoint best phải gắn với metric, tập và epoch. Nếu metric trung bình chứa ET empty–empty, tiêu chí chọn có thể ưu tiên sai. Một thiết kế tốt dùng full validation định kỳ hoặc primary metric không bị thành phần rỗng chi phối, đồng thời giữ last checkpoint để khôi phục run.")
    add_table(doc, ["Artifact", "Nội dung bắt buộc"], [
        ("best checkpoint", "Weights, epoch, primary metric, config hash"),
        ("last checkpoint", "Weights, optimizer, scheduler, scaler, RNG state"),
        ("training log", "Loss, per-class metric, LR, thời gian, lỗi"),
        ("evaluation report", "Case list, checkpoint hash, metric implementation"),
    ], [5.0, 10.6])


def add_inference_metric_depth(doc):
    add_heading(doc, "2.2.7.1 Ghép cửa sổ suy luận", 3)
    add_body(doc, "Với overlap 0,5, mỗi voxel có thể xuất hiện trong nhiều ROI. Mô hình tạo xác suất p_r,c(v) cho voxel v trong cửa sổ r. Bộ ghép tính trung bình có trọng số; trọng số thường lớn ở tâm ROI và nhỏ ở biên để giảm artefact. Sau khi chuẩn hóa tổng trọng số, argmax tạo nhãn cuối.")
    add_picture(doc, GENERATED / "eq_sliding_window.png", "Công thức 9. Ghép xác suất các ROI trong sliding-window inference.", max_width=5.8, max_height=1.45)
    add_body(doc, "Overlap tăng làm số ROI và latency tăng. Với stride bằng một nửa ROI, mỗi voxel nội vùng có thể được dự đoán nhiều lần, đổi compute lấy tính trơn ở biên. Benchmark phải dùng cùng overlap, sw_batch_size và blending mode khi so model.")

    add_heading(doc, "2.2.7.2 Hậu xử lý và khôi phục nhãn BraTS", 3)
    add_body(doc, "Logit được softmax và argmax thành lớp nội bộ 0–3. Khi xuất NIfTI, lớp ET nội bộ 3 phải được map lại raw label 4. WT, TC và ET được tạo bằng logic tập hợp. Nếu sai mapping, ảnh overlay vẫn có màu nhưng ý nghĩa vùng bị sai; vì vậy unit test cần kiểm tra từng nhãn bằng fixture nhỏ.")
    add_table(doc, ["Đầu ra nội bộ", "Raw label xuất", "Vùng tổng hợp"], [
        ("0", "0", "Background"),
        ("1", "1", "Thuộc TC và WT"),
        ("2", "2", "Thuộc WT"),
        ("3", "4", "Thuộc ET, TC và WT"),
    ], [4.2, 4.2, 7.2])
    add_body(doc, "Thể tích vùng bằng số voxel nhân tích spacing theo ba chiều rồi đổi đơn vị. Nếu spacing tính bằng mm, volume_mm3=N_voxel·s_x·s_y·s_z và volume_cm3=volume_mm3/1000. Không được tính thể tích chỉ từ số voxel khi dữ liệu không đẳng hướng.")

    add_heading(doc, "2.2.8.1 Quan hệ Dice IoU và HD95", 3)
    add_picture(doc, GENERATED / "eq_metrics.png", "Công thức 10. Dice, IoU và HD95 cho hai mask P và G.", max_width=6.0, max_height=1.5)
    add_body(doc, "Dice và IoU liên hệ bởi IoU=Dice/(2−Dice) khi tính trên cùng hai mask. Vì vậy chúng không phải hai bằng chứng độc lập, nhưng IoU cho thang nghiêm ngặt hơn. HD95 dùng bề mặt ∂P và ∂G: tính khoảng cách gần nhất hai chiều rồi lấy phân vị 95, giảm ảnh hưởng của một vài outlier cực đoan so với Hausdorff tối đa.")
    add_body(doc, "HD95 phải dùng spacing vật lý. Nếu đo trên chỉ số voxel, khoảng cách theo trục có spacing lớn sẽ bị đánh giá sai. Khi một trong hai mask rỗng, khoảng cách bề mặt không xác định tự nhiên; protocol cần gán penalty, NaN hoặc loại có điều kiện và luôn báo số ca.")
    add_table(doc, ["Tình huống", "Dice/IoU", "HD95", "Thông điệp"], [
        ("Overlap tốt, biên gần", "Cao", "Thấp", "Kết quả nhất quán"),
        ("Overlap khá, cụm FP xa", "Có thể vẫn khá", "Cao", "Có outlier hình học"),
        ("Mask nhỏ lệch vài voxel", "Giảm mạnh", "Tùy spacing", "Vùng nhỏ nhạy sai số"),
        ("Cả hai rỗng", "Theo policy", "Không xác định", "Không đo năng lực phát hiện"),
    ], [4.3, 3.5, 3.5, 4.3])

    add_heading(doc, "2.2.8.2 Macro micro và conditional metric", 3)
    add_body(doc, "Macro Dice lấy trung bình theo case, nên mỗi bệnh nhân có trọng số ngang nhau bất kể kích thước u. Micro Dice cộng TP/FP/FN toàn dataset trước khi tính, nên ca u lớn có ảnh hưởng mạnh. Hai cách trả lời câu hỏi khác nhau và nên được báo song song nếu cần đánh giá toàn diện.")
    add_body(doc, "Conditional ET Dice chỉ tính trên case có ground-truth ET > 0, đo năng lực phân đoạn khi mục tiêu hiện diện. Đối với case ET âm tính, báo false-positive rate hoặc predicted ET volume. Cặp chỉ số này tách khả năng tìm ET khỏi khả năng không báo giả và giải quyết trực tiếp vấn đề v3.")
    add_table(doc, ["Aggregation", "Cách tính", "Ưu điểm", "Rủi ro"], [
        ("Macro", "Mean metric theo case", "Mỗi case ngang nhau", "Nhạy empty policy"),
        ("Micro", "Gộp voxel rồi tính", "Ổn định thống kê", "Ca lớn chi phối"),
        ("Conditional", "Chỉ case có lớp", "Đo năng lực lớp hiện diện", "Cần báo cỡ mẫu"),
        ("Stratified", "Theo size/site/ET", "Phát hiện failure mode", "Cần đủ case mỗi nhóm"),
    ], [3.2, 4.5, 4.2, 3.7])

    add_heading(doc, "2.2.8.3 Độ phức tạp tính toán và bộ nhớ", 3)
    add_body(doc, "Số tham số đo dung lượng weights nhưng không dự đoán đầy đủ VRAM. Train còn lưu activation, gradient, moment m và v của AdamW. Với FP32, riêng weights của P tham số cần khoảng 4P byte; gradient thêm 4P; hai moment thêm 8P. Chưa tính activation, U-Net 4,811 triệu tham số cần tối thiểu khoảng 77 MB cho weights+gradient+moments, còn Swin 15,706 triệu cần khoảng 251 MB. Activation 3D thường lớn hơn nhiều.")
    add_body(doc, "FLOPs của convolution tỷ lệ với D_outH_outW_out·C_inC_out·K_dK_hK_w. Attention cửa sổ tỷ lệ với số token nhân số token mỗi cửa sổ và chiều feature. Gradient checkpointing giảm activation được lưu nhưng tăng forward recomputation. Mixed precision giảm byte mỗi tensor nhưng cần kiểm tra ổn định loss.")
    add_table(doc, ["Thành phần bộ nhớ", "FP32 xấp xỉ", "Ghi chú"], [
        ("Weights", "4P byte", "Có ở train và inference"),
        ("Gradients", "4P byte", "Chỉ train"),
        ("Adam first moment", "4P byte", "Chỉ train"),
        ("Adam second moment", "4P byte", "Chỉ train"),
        ("Activations", "Phụ thuộc batch/shape/layer", "Thường chi phối segmentation 3D"),
    ], [5.2, 4.4, 6.0])

    add_heading(doc, "2.2.8.4 Từ công thức đến protocol đánh giá", 3)
    add_body(doc, "Một metric chỉ có nghĩa khi kèm năm thông tin: case list, preprocessing, mapping vùng, empty-mask policy và aggregation. Thiếu một trong các mục này, cùng tên Dice có thể cho giá trị khác. Đây là lý do báo cáo không so trực tiếp mean v2 và v3 như cùng benchmark.")
    add_body(doc, "Protocol chuẩn đề xuất lưu prediction ở không gian gốc hoặc không gian evaluation đã định, kiểm tra spacing, tạo WT/TC/ET bằng một hàm dùng chung, tính per-case metric và xuất CSV. Script tổng hợp đọc CSV để tạo mean, median, quantile, confidence interval và đồ thị outlier. Test set được khóa cho đến khi model/hyperparameter hoàn tất.")
    add_numbered(doc, [
        "Đóng băng manifest test và checksum.",
        "Chạy hai checkpoint với cùng preprocessing và sliding window.",
        "Khôi phục raw label và tạo ba vùng bằng cùng code.",
        "Tính per-case Dice, IoU, HD95 và volume error.",
        "Tách ET-present và ET-absent trước khi tổng hợp.",
        "Bootstrap theo case và báo khoảng tin cậy chênh lệch.",
    ])


def build_report():
    create_diagrams()
    shutil.copy2(TEMPLATE, OUTPUT)
    doc = Document(OUTPUT)
    remove_body_content(doc)
    configure_page(doc)
    configure_styles(doc)
    doc.core_properties.title = "Capstone Project Final Report – 3D Brain Tumor Segmentation"
    doc.core_properties.subject = "So sánh 3D U-Net và Swin UNETR trên dữ liệu BraTS"
    doc.core_properties.author = "AI Healthcare Team"
    doc.core_properties.comments = "Báo cáo được xây dựng từ mẫu Samsung Innovation Campus, sử dụng kết quả có sẵn trong notebook v2 và v3; không huấn luyện lại mô hình."

    cover_page(doc)

    next_page(doc)
    add_chapter(doc, "PROJECT INFORMATION")
    add_heading(doc, "Thông tin dự án", 2)
    add_table(
        doc,
        ["Hạng mục", "Thông tin"],
        [
            ("Project Title", "Hệ thống hỗ trợ phân đoạn khối u não trên ảnh MRI 3D"),
            ("English Title", "AI-assisted 3D Brain Tumor Segmentation on Multi-modal MRI"),
            ("Team Name", "AI Healthcare Team"),
            ("Course", "Samsung Innovation Campus – AI Course"),
            ("Report Date", "14/09/2026"),
            ("Repository", r"D:\SIC_Capstone 2026"),
        ],
        [4.6, 11.0],
    )
    add_heading(doc, "Thành viên", 2)
    add_table(
        doc,
        ["STT", "Họ và tên", "Mã học viên", "Vai trò chính"],
        [(str(i), "[BỔ SUNG]", "[BỔ SUNG]", role) for i, role in enumerate([
            "Điều phối dự án và tài liệu", "Dữ liệu và tiền xử lý", "Huấn luyện 3D U-Net",
            "Huấn luyện Swin UNETR", "Ứng dụng, kiểm thử và trình bày"], 1)],
        [1.2, 5.0, 3.2, 6.2],
    )
    add_body(doc, "Các trường thành viên được để ở trạng thái [BỔ SUNG] vì kho dự án chưa cung cấp thông tin định danh. Nội dung kỹ thuật không giả định cá nhân chịu trách nhiệm cho phần việc cụ thể.")

    next_page(doc)
    add_chapter(doc, "EXECUTIVE SUMMARY")
    add_heading(doc, "Tóm tắt điều hành", 2)
    for p in [
        "Dự án xây dựng một bản mẫu hỗ trợ phân đoạn khối u não trên ảnh cộng hưởng từ đa chuỗi. Hệ thống nhận bốn ảnh NIfTI gồm T1 native, T1 có tiêm tương phản, T2 và T2-FLAIR; thực hiện chuẩn hóa, suy luận thể tích 3D, hiển thị mặt cắt và xuất mask cùng báo cáo minh họa. Mục tiêu là hỗ trợ quy trình nghiên cứu, giảm thao tác lặp lại và tạo đầu ra định lượng có thể kiểm tra lại.",
        "Hai notebook thể hiện hai nhánh thực nghiệm khác nhau. Notebook v2 huấn luyện 3D U-Net trên 1.251 ca hợp lệ và báo cáo trên tập kiểm thử độc lập 188 ca: Mean Dice 0,7361; Dice WT 0,7489; TC 0,7534; ET 0,7061; Mean IoU 0,6310; HD95 43,058 mm. Notebook v3 huấn luyện Swin UNETR trên 1.016 ca, đánh giá toàn bộ 204 ca validation: Mean Dice 0,8653; WT 0,8498; TC 0,7461; ET 1,0000. Tuy nhiên ET bằng 1,0 vì cả nhãn thật và dự đoán ET đều rỗng ở toàn bộ 204 ca, nên chỉ số trung bình bị thổi phồng.",
        "Do khác tập dữ liệu, cách chia tập và protocol đánh giá, báo cáo không kết luận Swin UNETR vượt trội trực tiếp. 3D U-Net có bằng chứng tổng quát hóa đáng tin cậy hơn nhờ held-out test; Swin UNETR cho tín hiệu tích cực ở WT nhưng cần được chạy lại trên cùng split và trên các ca có ET. Khuyến nghị ưu tiên thiết kế đối chứng công bằng, kiểm tra logic nhãn ET và bổ sung đánh giá chuyên gia trước mọi tuyên bố sử dụng lâm sàng.",
    ]:
        add_body(doc, p)
    add_bullets(doc, [
        "Phạm vi hoàn thành: dữ liệu, huấn luyện, đánh giá định lượng, trực quan hóa 3D và giao diện Gradio.",
        "Phạm vi chưa hoàn thành: thử nghiệm đa trung tâm, đánh giá bác sĩ, hiệu chuẩn thiết bị và chứng nhận y tế.",
        "Nguyên tắc diễn giải: mọi con số gắn với đúng tập đánh giá; không dùng best fast-validation như kết quả cuối cùng.",
    ])

    next_page(doc)
    add_chapter(doc, "ABSTRACT")
    add_heading(doc, "Tóm tắt", 2)
    add_body(doc, "Phân đoạn thể tích khối u não trên MRI là một bài toán thị giác máy tính y sinh có yêu cầu cao về không gian ba chiều, đa phương thức và độ tin cậy của nhãn. Báo cáo trình bày quá trình xác định vấn đề theo Quy trình SIC, thiết kế hệ thống, xử lý dữ liệu BraTS, huấn luyện hai kiến trúc 3D U-Net và Swin UNETR, đánh giá kết quả và tích hợp một ứng dụng minh họa. Phương pháp được tổ chức theo CRISP-DM và vòng Discover–Define–Analyze–Ideate–Prototype–Test nhằm liên kết nhu cầu người dùng với bằng chứng kỹ thuật.")
    add_body(doc, "Kết quả cho thấy 3D U-Net đạt Mean Dice 0,7361 trên held-out test 188 ca. Swin UNETR đạt Mean Dice 0,8653 trên full validation 204 ca, nhưng thành phần ET bằng 1,0 do tất cả trường hợp ET là empty–empty. Vì hai mô hình không dùng cùng cohort và protocol, so sánh được trình bày dưới dạng as-reported, kèm phân tích mối đe dọa đến tính hợp lệ. Báo cáo đề xuất một benchmark thống nhất, stratify theo hiện diện ET, báo cáo macro và conditional Dice, và xác nhận đầu ra bằng chuyên gia. Bản mẫu có giá trị đào tạo và nghiên cứu nhưng chưa đủ bằng chứng cho sử dụng lâm sàng.")
    add_heading(doc, "Từ khóa", 2)
    add_body(doc, "Brain tumor segmentation; BraTS; MRI; 3D U-Net; Swin UNETR; MONAI; Dice; HD95; CRISP-DM; explainable visualization.")

    next_page(doc)
    add_chapter(doc, "TABLE OF CONTENTS")
    p = doc.add_paragraph()
    add_field(p, 'TOC \\o "1-3" \\h \\z \\u')
    add_body(doc, "Mục lục được tạo bằng trường tự động của Microsoft Word. Số trang sẽ được cập nhật trong bước kết xuất cuối cùng.")

    next_page(doc)
    add_chapter(doc, "LISTS AND ABBREVIATIONS")
    add_heading(doc, "Danh mục chữ viết tắt", 2)
    add_table(doc, ["Viết tắt", "Diễn giải"], [
        ("AI", "Artificial Intelligence – Trí tuệ nhân tạo"),
        ("MRI", "Magnetic Resonance Imaging – Cộng hưởng từ"),
        ("BraTS", "Brain Tumor Segmentation Challenge"),
        ("WT", "Whole Tumor – Toàn bộ vùng u"),
        ("TC", "Tumor Core – Lõi u"),
        ("ET", "Enhancing Tumor – Vùng u tăng cường"),
        ("CNN", "Convolutional Neural Network"),
        ("ViT", "Vision Transformer"),
        ("IoU", "Intersection over Union"),
        ("HD95", "95th percentile Hausdorff Distance"),
        ("NIfTI", "Neuroimaging Informatics Technology Initiative format"),
        ("ROI", "Region of Interest"),
    ], [3.0, 12.6])
    add_heading(doc, "Danh mục hình và bảng chính", 2)
    add_bullets(doc, [
        "Hình 1–2: quy trình xác định vấn đề và so sánh phân chia dữ liệu.",
        "Hình 3–7: kiến trúc 3D U-Net, Swin UNETR, pipeline và hệ thống.",
        "Hình 8–19: EDA, đường học, dự đoán và so sánh hai mô hình.",
        "Bảng 1–18: dữ liệu, cấu hình, kết quả, rủi ro, kiểm thử và truy vết yêu cầu.",
    ])

    # CHAPTER 1
    next_page(doc)
    add_chapter(doc, "1 Introduction")
    add_heading(doc, "1.1 Background Information", 2)
    for p in [
        "Glioma là nhóm u não nguyên phát có hình thái, kích thước và mức độ xâm lấn đa dạng. MRI đa chuỗi cho phép quan sát các đặc trưng mô khác nhau mà không dùng bức xạ ion hóa. Trong benchmark BraTS, bốn chuỗi MRI được đồng đăng ký và chuẩn hóa để phục vụ phân đoạn các vùng khối u có ý nghĩa đánh giá định lượng [1]–[4].",
        "Phân đoạn thủ công thể tích 3D đòi hỏi người đọc lần lượt xem nhiều lát cắt, duy trì tính nhất quán giữa các mặt phẳng và xử lý vùng biên mơ hồ. Học sâu có thể tạo mask sơ bộ và thống kê thể tích, nhưng đầu ra chỉ hữu ích khi pipeline kiểm soát đúng dữ liệu đầu vào, quy ước nhãn và sai số. Vì vậy dự án được đặt trong vai trò hỗ trợ nghiên cứu, không thay thế bác sĩ.",
    ]:
        add_body(doc, p)
    add_source(doc, "Nguồn nền tảng: BraTS/CBICA [1], Menze et al. [2], Bakas et al. [3], BraTS 2023 [4].")

    add_topic_page(doc, "1.1.1 Bối cảnh bài toán y sinh", [
        "Ảnh MRI não là dữ liệu thể tích với voxel có quan hệ không gian theo ba chiều. Một vùng u có thể xuất hiện nhỏ trên một lát cắt nhưng kéo dài qua nhiều lát khác; mô hình 2D riêng lẻ dễ bỏ lỡ tính liên tục này. Dự án do đó lựa chọn mô hình 3D, chấp nhận chi phí bộ nhớ lớn hơn để học ngữ cảnh thể tích.",
        "Các vùng BraTS được đánh giá theo các cấu trúc tổng hợp WT, TC và ET. Chúng không đơn giản là ba lớp độc lập: ET nằm trong TC, và TC nằm trong WT theo cách quy đổi nhãn. Điều này tạo yêu cầu hậu xử lý và diễn giải nhất quán; một chỉ số Dice tổng hợp không đủ để mô tả lỗi lâm sàng tiềm ẩn ở từng vùng.",
    ], table=(["Vùng", "Thành phần khái niệm", "Ý nghĩa trong báo cáo"], [
        ("WT", "Toàn bộ mô u và phù liên quan", "Phản ánh phạm vi tổn thương tổng thể"),
        ("TC", "Lõi u, gồm vùng hoại tử/không tăng cường và ET", "Phản ánh cấu trúc trung tâm của u"),
        ("ET", "Phần tăng cường sau tiêm tương phản", "Nhạy với lỗi nhãn rỗng và kích thước nhỏ"),
    ], [2.0, 7.0, 6.6]))

    add_topic_page(doc, "1.1.2 Vai trò của bốn chuỗi MRI", [
        "Mỗi chuỗi MRI cung cấp một tương phản mô khác nhau. T1 native mô tả giải phẫu; T1 có tiêm tương phản nhấn mạnh vùng tăng cường; T2 và T2-FLAIR nhạy với thành phần giàu nước và phù. Ghép bốn chuỗi thành tensor bốn kênh giúp mô hình khai thác thông tin bổ sung thay vì phụ thuộc một ảnh đơn lẻ.",
        "Điều kiện quan trọng là bốn ảnh của cùng ca phải đúng bệnh nhân, cùng không gian và đúng tên modality. Giao diện hiện tại yêu cầu người dùng tải đủ bốn file. Hướng phát triển nên bổ sung kiểm tra header NIfTI, affine, shape và cảnh báo khi các chuỗi không đồng đăng ký.",
    ], table=(["Kênh", "Tên file trong dự án", "Tín hiệu sử dụng"], [
        ("T1n", "t1n", "Giải phẫu nền"), ("T1c", "t1c", "Vùng tăng cường"),
        ("T2w", "t2w", "Dịch và mô bệnh"), ("T2-FLAIR", "t2f", "Phù và tổn thương lan tỏa"),
    ], [2.2, 4.0, 9.4]))

    add_topic_page(doc, "1.1.3 Các bên liên quan và nhu cầu", [
        "Báo cáo dùng cách nhìn theo stakeholder thay vì chỉ mô tả thuật toán. Người sử dụng trực tiếp của bản mẫu là học viên, kỹ sư hoặc nhà nghiên cứu cần kiểm tra nhanh dữ liệu và kết quả. Bác sĩ chẩn đoán hình ảnh là người có chuyên môn để thẩm định mask, nhưng chưa tham gia nghiên cứu người dùng trong phạm vi tài liệu hiện có. Người quản trị dữ liệu quan tâm đến tính hợp lệ, quyền riêng tư và truy vết nguồn dữ liệu.",
        "Nhu cầu chung là nhận đầu vào có kiểm soát, tạo kết quả nhất quán, hiển thị đủ trực quan để phát hiện lỗi và lưu lại đầu ra có thể tái kiểm tra. Không bên liên quan nào nên bị buộc tin vào một con số duy nhất; hệ thống phải để lộ ảnh gốc, overlay và thông tin phiên bản mô hình.",
    ], bullets=[
        "Nhà nghiên cứu: cần benchmark, chỉ số và khả năng tái lập.",
        "Người dùng ứng dụng: cần hướng dẫn upload, thời gian phản hồi và thông báo lỗi rõ ràng.",
        "Chuyên gia y tế: cần mask có thể chỉnh sửa, độ tin cậy và bằng chứng ngoài tập huấn luyện.",
        "Quản trị viên: cần quản lý dữ liệu, logging và kiểm soát phiên bản.",
    ])

    add_topic_page(doc, "1.2 Motivation and Objective", [
        "Động lực của dự án là chuyển một bài toán nghiên cứu phân đoạn 3D thành pipeline có thể quan sát từ dữ liệu đến đầu ra. Thay vì dừng ở notebook huấn luyện, nhóm xây dựng lớp suy luận, trực quan hóa và xuất báo cáo. Giá trị chính nằm ở tính liên kết: người đọc có thể đối chiếu preprocessing, checkpoint, mask, chỉ số và giao diện.",
        "Mục tiêu không phải tạo thiết bị y tế hoàn chỉnh. Dự án tập trung chứng minh tính khả thi kỹ thuật, so sánh hai họ kiến trúc và xác định các bước kiểm chứng còn thiếu. Mọi phát biểu về lợi ích lâm sàng trong báo cáo vì thế được trình bày như tiềm năng hoặc giả thuyết cần xác nhận.",
    ], bullets=[
        "Huấn luyện và lưu checkpoint của một baseline CNN 3D và một mô hình Transformer 3D.",
        "Đánh giá định lượng theo WT, TC, ET; ghi rõ tập dữ liệu và protocol.",
        "Tích hợp suy luận qua giao diện Gradio và tạo đầu ra có thể tải xuống.",
        "Xây dựng kế hoạch follow-up theo vòng kiểm thử và cải tiến của Quy trình SIC.",
    ])

    add_topic_page(doc, "1.2.1 Xác định vấn đề theo Quy trình SIC", [
        "Theo trang Define trong Quy trình SIC, vấn đề nên trả lời Ai, Cái gì, Ở đâu và lý do cần can thiệp. Tuyên bố vấn đề của dự án là: Trong môi trường đào tạo và nghiên cứu phân tích MRI não, người phân tích cần tạo và kiểm tra vùng khối u trên thể tích đa chuỗi; thao tác thủ công và thiếu pipeline thống nhất làm tăng công sức, khó tái lập và khó so sánh mô hình. Hãy xây dựng một công cụ AI tạo mask sơ bộ, trực quan hóa và báo cáo có truy vết, nhưng không thay thế quyết định của chuyên gia.",
        "Cách diễn đạt này cố ý nêu ranh giới sử dụng. Hệ thống giải quyết phần tự động hóa kỹ thuật và hỗ trợ kiểm tra; hệ thống không tự đưa ra chẩn đoán, tiên lượng hoặc chỉ định điều trị. Đây là điều kiện để phạm vi nghiên cứu phù hợp với dữ liệu và mức độ kiểm chứng hiện tại.",
    ], figure=GENERATED / "problem_followup.png", caption="Hình 1. Quy trình xác định vấn đề và vòng follow-up được chuyển hóa từ Quy trình SIC.", source="Nguồn: tổng hợp từ doc/Quy trình.pdf, các phần Define, Analyze, Ideate, Prototype và Test.")

    add_topic_page(doc, "1.2.2 Câu hỏi Who–What–Where–Why", [
        "Who: người phân tích dữ liệu MRI và nhóm nghiên cứu cần một mask tham chiếu nhanh; chuyên gia y tế là người thẩm định cuối nếu nghiên cứu được mở rộng. What: quy trình hiện thiếu cơ chế tự động phân đoạn và so sánh mô hình theo cùng chuẩn. Where: notebook nghiên cứu và ứng dụng cục bộ dùng dữ liệu BraTS chuẩn hóa. Why: giảm thao tác lặp lại, tăng khả năng tái lập và làm rõ sai số theo từng vùng u.",
        "Tuyên bố không sử dụng số liệu thời gian hay tỷ lệ bất đồng chưa được dự án đo trực tiếp. Trong giai đoạn follow-up, các đại lượng như thời gian xử lý thủ công, thời gian hiệu chỉnh mask và mức chấp nhận của người dùng phải được thu thập thực nghiệm thay vì suy đoán.",
    ], table=(["Thành phần", "Trả lời trong dự án", "Bằng chứng hiện có"], [
        ("Who", "Nhà nghiên cứu/người phân tích", "Ứng dụng và quy trình notebook"),
        ("What", "Mask 3D và báo cáo chưa được chuẩn hóa", "Hai notebook dùng protocol khác nhau"),
        ("Where", "Môi trường nghiên cứu cục bộ", "Gradio app và dữ liệu BraTS"),
        ("Why", "Tái lập, kiểm tra, so sánh", "Checkpoint, metric và output trực quan"),
    ], [2.5, 7.0, 6.1]))

    add_topic_page(doc, "1.2.3 Phân tích nguyên nhân gốc", [
        "Nguyên nhân cấp dữ liệu gồm kích thước thể tích lớn, mất cân bằng vùng u, biến thiên scanner và khả năng thiếu hoặc sai modality. Nguyên nhân cấp mô hình gồm giới hạn receptive field, chi phí bộ nhớ, sai lệch do patch sampling và loss bị chi phối bởi lớp lớn. Nguyên nhân cấp đánh giá gồm split khác nhau, chọn checkpoint theo tập validation nhỏ và quy tắc xử lý mask rỗng.",
        "Trường hợp ET ở notebook v3 minh họa rõ nhất: toàn bộ ET volume bằng 0 trên full validation và Dice ET bằng 1 cho mọi ca. Nếu chỉ nhìn Mean Dice 0,8653, người đọc có thể cho rằng mô hình tốt toàn diện. Nguyên nhân gốc không nhất thiết là kiến trúc; có thể nằm ở cohort, ánh xạ nhãn, bộ lọc dữ liệu hoặc cách tính empty–empty. Follow-up phải bắt đầu từ audit nhãn trước khi tối ưu mô hình.",
    ], bullets=[
        "Audit danh sách case và phân bố nhãn trước khi train.",
        "Kiểm tra số ca có ET > 0 ở từng split.",
        "Báo cáo Dice ET có điều kiện chỉ trên ca có ET, bên cạnh quy tắc empty-mask.",
        "Khóa split và metric implementation khi so sánh mô hình.",
    ])

    add_topic_page(doc, "1.2.4 Tác động cá nhân và xã hội", [
        "Ở cấp cá nhân, một mask sơ bộ có thể giúp người phân tích tập trung vào kiểm tra vùng biên thay vì bắt đầu từ ảnh trống. Ở cấp nhóm, pipeline thống nhất giúp giảm khác biệt do thao tác và lưu lại phiên bản mô hình. Ở cấp nghiên cứu, báo cáo rõ giao thức hạn chế nguy cơ so sánh sai giữa các thí nghiệm.",
        "Tác động tiêu cực có thể xuất hiện nếu người dùng mặc định xem mask là đúng, nếu dữ liệu ngoài miền tạo kết quả sai nhưng giao diện không cảnh báo, hoặc nếu báo cáo PDF bị hiểu như tài liệu chẩn đoán. Biện pháp giảm thiểu là gắn nhãn nghiên cứu, hiển thị overlay, cung cấp thông tin phiên bản và yêu cầu review chuyên gia trong các bối cảnh y tế.",
    ], table=(["Cấp độ", "Lợi ích kỳ vọng", "Rủi ro cần kiểm soát"], [
        ("Cá nhân", "Giảm thao tác lặp lại", "Quá tin vào mask"),
        ("Nhóm", "Tái lập và chia sẻ kết quả", "Dùng nhầm checkpoint/protocol"),
        ("Nghiên cứu", "So sánh có cấu trúc", "Thiên lệch dataset"),
        ("Xã hội", "Tiềm năng hỗ trợ tiếp cận công cụ", "Bất bình đẳng do domain shift"),
    ], [2.2, 6.4, 7.0]))

    add_topic_page(doc, "1.2.5 Discover và bằng chứng đã sử dụng", [
        "Dự án sử dụng ba nguồn bằng chứng: tài liệu BraTS và các công trình nền tảng, dữ liệu/nhãn trong kho, và quan sát trực tiếp output notebook cùng ứng dụng. Báo cáo không khẳng định đã phỏng vấn người dùng hoặc bác sĩ vì không có biên bản trong repository. Khoảng trống này được ghi thành hoạt động follow-up.",
        "Việc quan sát output giúp phát hiện các vấn đề mà chỉ đọc mã nguồn khó thấy: đường học, ET rỗng, sự khác nhau giữa fast validation và full validation, cùng khác biệt giữa held-out test và validation. Đây là ví dụ của bằng chứng thực nghiệm nội bộ, nhưng vẫn cần được phân biệt với đánh giá ngoài tập và đánh giá người dùng.",
    ], bullets=[
        "Literature review: BraTS, 3D U-Net, Swin UNETR, MONAI.",
        "Data review: số case, modality, nhãn, phân chia tập.",
        "Artifact review: checkpoint, đường học, CSV metric, ảnh dự đoán, ứng dụng.",
        "Chưa có: phỏng vấn chuyên gia, quan sát workflow bệnh viện, thử nghiệm đa trung tâm.",
    ])

    add_topic_page(doc, "1.2.6 Persona giả định và giả thuyết cần kiểm chứng", [
        "Persona làm việc được mô tả là một nghiên cứu viên trẻ có kiến thức MRI cơ bản, quen notebook nhưng cần một giao diện để tải dữ liệu, xem mask và xuất kết quả. Đây là giả định thiết kế, không phải kết quả nghiên cứu người dùng. Nhu cầu chính là quy trình rõ ràng, thông báo lỗi dễ hiểu và khả năng chọn mô hình.",
        "Một persona thứ hai là chuyên gia thẩm định, quan tâm đến sai số vùng biên, trường hợp thất bại và khả năng so sánh với ground truth. Giao diện hiện tại mới đáp ứng phần xem đa mặt phẳng; chưa có công cụ chỉnh sửa mask, đo độ tin cậy hay lưu phản hồi. Các điểm này trở thành backlog có thể kiểm thử.",
    ], table=(["Persona", "Nhiệm vụ", "Pain point", "Giả thuyết kiểm chứng"], [
        ("Nghiên cứu viên", "Chạy suy luận", "Nhiều bước kỹ thuật", "UI giảm lỗi thao tác"),
        ("Chuyên gia thẩm định", "Kiểm tra mask", "Thiếu công cụ hiệu chỉnh", "Overlay ba mặt phẳng đủ để sàng lọc lỗi"),
        ("Quản trị dữ liệu", "Quản lý ca", "Thiếu metadata/log", "Manifest đầu vào tăng truy vết"),
    ], [3.0, 3.2, 4.5, 5.3]))

    add_topic_page(doc, "1.2.7 Mục tiêu đo lường và câu hỏi nghiên cứu", [
        "Câu hỏi RQ1: hai kiến trúc học được gì trong protocol hiện có và mức độ tổng quát hóa được chứng minh đến đâu? RQ2: sự khác biệt giữa CNN 3D và Transformer 3D có đi kèm chi phí tài nguyên đáng kể không? RQ3: pipeline giao diện có cho phép người dùng hoàn thành luồng từ upload đến tải kết quả một cách có kiểm soát không?",
        "Tiêu chí kỹ thuật bao gồm đọc đủ bốn modality, trả về mask đúng shape, thống kê WT/TC/ET, không dùng test để chọn checkpoint và lưu phiên bản mô hình. Tiêu chí đánh giá mô hình bao gồm Dice theo từng vùng, IoU, HD95 khi có, cùng phân tích ca rỗng. Tiêu chí người dùng cần được đo trong follow-up bằng completion rate, lỗi thao tác và thời gian hoàn thành.",
    ], table=(["Mục tiêu", "Chỉ số", "Trạng thái"], [
        ("Phân đoạn 3 vùng", "Dice WT/TC/ET", "Đã có, khác protocol"),
        ("Đánh giá ngoài chọn mô hình", "Held-out test", "Có ở v2; chưa có ở v3"),
        ("Tích hợp ứng dụng", "Upload–inference–export", "Đã có bản mẫu"),
        ("Usability", "Task success/time/error", "Chưa đo"),
        ("Clinical validation", "Expert agreement/external set", "Ngoài phạm vi hiện tại"),
    ], [5.5, 5.0, 5.1]))

    add_topic_page(doc, "1.3 Members and Role Assignments", [
        "Kho dự án chưa chứa danh sách thành viên, vì vậy báo cáo sử dụng ma trận vai trò thay cho việc gán tên. Khi hoàn thiện hồ sơ nộp, nhóm cần điền tên và bằng chứng đóng góp. Một người có thể đảm nhiệm nhiều vai trò, nhưng mỗi deliverable phải có người chịu trách nhiệm chính và người review.",
        "Phân công được khuyến nghị theo RACI: điều phối chịu trách nhiệm timeline; data lead kiểm tra dataset; model lead phụ trách từng notebook; application lead tích hợp; documentation lead bảo đảm truy vết và tính nhất quán. Quyết định kỹ thuật quan trọng cần được peer review để giảm lỗi chỉ số hoặc nhãn.",
    ], table=(["Deliverable", "Responsible", "Accountable", "Consulted/Reviewer"], [
        ("Dataset manifest", "[BỔ SUNG]", "[BỔ SUNG]", "Model lead"),
        ("3D U-Net v2", "[BỔ SUNG]", "[BỔ SUNG]", "Data lead"),
        ("Swin UNETR v3", "[BỔ SUNG]", "[BỔ SUNG]", "Data lead"),
        ("Gradio app", "[BỔ SUNG]", "[BỔ SUNG]", "Model lead"),
        ("Final report", "[BỔ SUNG]", "[BỔ SUNG]", "Toàn nhóm"),
    ], [5.0, 3.6, 3.6, 3.4]))

    add_topic_page(doc, "1.4 Schedule and Milestones", [
        "Timeline được tái dựng theo các deliverable quan sát thấy trong repository. Vì không có nhật ký ngày hoàn thành chi tiết, báo cáo trình bày milestone theo pha thay vì gán ngày không có bằng chứng. Nhóm cần cập nhật ngày thực tế trong bản nộp cuối nếu giảng viên yêu cầu.",
        "Các cổng quyết định gồm: dữ liệu hợp lệ; baseline chạy được; checkpoint được đánh giá; mô hình thứ hai hoàn thành; ứng dụng kết nối; báo cáo được review. Sau mỗi cổng, rủi ro và vấn đề tồn đọng được đưa vào action plan thay vì bị ẩn dưới chỉ số trung bình.",
    ], table=(["Pha", "Đầu ra", "Tiêu chí hoàn tất"], [
        ("Discover/Define", "Problem statement", "Who–What–Where–Why rõ ràng"),
        ("Data", "Manifest và split", "Không trộn train/test"),
        ("Baseline", "3D U-Net checkpoint", "Có full metrics và test"),
        ("Advanced model", "Swin UNETR checkpoint", "Có full validation và audit ET"),
        ("Prototype", "Gradio app", "Luồng upload–export hoạt động"),
        ("Report", "DOCX > 50 trang", "Đúng mẫu, kiểm tra layout"),
    ], [3.3, 6.0, 6.3]))

    # CHAPTER 2
    next_page(doc)
    add_chapter(doc, "2 Project Execution")
    add_heading(doc, "2.1 Data Acquisition", 2)
    for p in [
        "Dữ liệu sử dụng thuộc hệ sinh thái BraTS 2023, gồm các thể tích MRI đa chuỗi và mask phân đoạn. Notebook v2 phát hiện 1.251 ca hợp lệ; notebook v3 dùng 1.016 ca sau bước đồng bộ/lọc được ghi trong output. Báo cáo không gộp hai con số vì chúng đại diện cho hai pipeline dữ liệu khác nhau.",
        "BraTS cung cấp dữ liệu nghiên cứu đã được chuẩn hóa theo quy trình challenge. Khi triển khai trên dữ liệu mới, không thể giả định ảnh có cùng hướng, spacing, intensity hoặc tiền xử lý. Do đó mọi tuyên bố hiệu năng chỉ áp dụng cho phân bố dữ liệu đã quan sát và protocol mô tả trong notebook.",
    ]:
        add_body(doc, p)
    add_source(doc, "Nguồn: notebook SIC_Capstone_v2.ipynb, SIC_Capstone_v3.ipynb và BraTS 2023 [4].")

    add_topic_page(doc, "2.1.1 Cấu trúc dữ liệu đầu vào", [
        "Mỗi case phải có bốn file modality và một segmentation label trong quá trình huấn luyện. Pipeline ghép bốn volume theo thứ tự cố định thành tensor C×H×W×D. Nếu thứ tự thay đổi, mô hình có thể nhận nhầm tương phản và tạo mask không đáng tin cậy dù shape vẫn hợp lệ.",
        "Một manifest dữ liệu nên lưu case ID, đường dẫn, modality, shape, spacing, affine, trạng thái nhãn và checksum. Notebook hiện chủ yếu dựa trên quy ước tên thư mục/file. Đối với bản mẫu đào tạo điều này khả thi, nhưng production cần validation rõ ràng trước suy luận.",
    ], table=(["Trường manifest", "Mục đích kiểm soát"], [
        ("case_id", "Liên kết bốn modality và đầu ra"),
        ("modality", "Ngăn đảo thứ tự kênh"),
        ("shape/spacing/affine", "Phát hiện không đồng đăng ký"),
        ("label_present", "Phân biệt train và inference"),
        ("checksum/version", "Truy vết dữ liệu"),
    ], [5.5, 10.1]))

    add_topic_page(doc, "2.1.2 Quy ước nhãn và ánh xạ", [
        "Notebook v2 ánh xạ nhãn raw 0, 1, 2, 4 thành chỉ số liên tiếp 0, 1, 2, 3. Việc ánh xạ này thuận tiện cho loss đa lớp nhưng phải được đảo đúng khi xuất NIfTI. Sai mapping có thể làm ET bị biến mất hoặc được gán sang vùng khác.",
        "Đánh giá BraTS thường tạo ba mask nhị phân tổng hợp. WT bao gồm mọi voxel u; TC bao gồm lõi u; ET là vùng tăng cường. Code metric cần công bố cách xử lý trường hợp ground truth rỗng và prediction rỗng, vì quy tắc gán Dice 1 trong empty–empty tác động mạnh đến macro mean.",
    ], table=(["Nhãn raw", "Nhãn nội bộ", "Diễn giải"], [
        ("0", "0", "Background"), ("1", "1", "NCR/NET"),
        ("2", "2", "Edema"), ("4", "3", "Enhancing tumor"),
    ], [3.0, 3.4, 9.2]))

    add_topic_page(doc, "2.1.3 Chất lượng và quản trị dữ liệu", [
        "Kiểm tra chất lượng tối thiểu gồm đủ file, load được NIfTI, finite intensity, shape nhất quán, nhãn thuộc tập giá trị cho phép và thể tích vùng u hợp lý. Các kiểm tra này cần chạy trước khi tạo cache hoặc NPZ; nếu không, lỗi dữ liệu có thể được đóng gói và lặp lại trong mọi lần train.",
        "Dữ liệu MRI là dữ liệu sức khỏe nhạy cảm. Dù BraTS dùng dữ liệu nghiên cứu đã xử lý, ứng dụng tương lai phải có chính sách lưu giữ, kiểm soát truy cập, ẩn định danh và xóa dữ liệu tạm. Bản mẫu hiện chạy cục bộ là một lợi thế giảm truyền dữ liệu, nhưng không tự động bảo đảm tuân thủ.",
    ], bullets=[
        "Không đưa thông tin định danh vào tên file xuất.",
        "Không lưu upload lâu hơn thời gian cần thiết.",
        "Ghi model version và preprocessing version trong báo cáo.",
        "Tách dữ liệu test khỏi quá trình chọn checkpoint.",
    ])

    add_topic_page(doc, "2.1.4 Phân chia dữ liệu notebook v2", [
        "Notebook v2 dùng seed 42 và chia 1.251 ca thành 875 train, 188 validation và 188 held-out test. Validation dùng để theo dõi và chọn mô hình; test chỉ được mở cho báo cáo cuối. Đây là thiết kế bằng chứng phù hợp hơn cho ước lượng tổng quát hóa nội bộ.",
        "Fast validation chỉ lấy 20 ca và chạy mỗi hai epoch để tiết kiệm thời gian; full validation chạy định kỳ dài hơn. Khi diễn giải best epoch, cần xác định nó được chọn theo fast hay full validation. Kết quả test 188 ca vẫn là điểm tham chiếu chính của v2 vì không tham gia tối ưu.",
    ], table=(["Tập", "Số ca", "Tỷ lệ xấp xỉ", "Vai trò"], [
        ("Train", "875", "69,9%", "Tối ưu tham số"),
        ("Validation", "188", "15,0%", "Theo dõi/chọn checkpoint"),
        ("Held-out test", "188", "15,0%", "Báo cáo cuối"),
    ], [3.2, 2.4, 3.2, 6.8]))

    add_topic_page(doc, "2.1.5 Phân chia dữ liệu notebook v3", [
        "Notebook v3 ghi nhận 1.016 ca sau đồng bộ/lọc, với 204 ca full validation và khoảng 812 ca train theo tỷ lệ 80/20. Notebook không tạo held-out test riêng. Fast validation chỉ có hai ca, vì vậy best fast-val Dice 0,9538 có phương sai rất lớn và không đại diện cho toàn bộ dữ liệu.",
        "Sự khác nhau về số ca có thể do đường dẫn, trạng thái cache, tiêu chí hợp lệ hoặc phiên bản dataset. Trước đối chứng, nhóm phải xuất danh sách case ID của từng notebook và tính giao/hiệu tập hợp. Chỉ khi hai mô hình dùng cùng cohort và split, so sánh chênh lệch metric mới phản ánh kiến trúc tốt hơn.",
    ], table=(["Tập", "Số ca", "Vai trò", "Điểm cần lưu ý"], [
        ("Train", "812", "Tối ưu", "Suy ra từ tổng 1.016 và val 204"),
        ("Full validation", "204", "Báo cáo v3", "Không độc lập như test"),
        ("Fast validation", "2", "Theo dõi nhanh", "Không dùng làm kết luận cuối"),
        ("Held-out test", "0", "—", "Khoảng trống đánh giá"),
    ], [3.2, 2.2, 4.3, 5.9]), figure=GENERATED / "split_comparison.png", caption="Hình 2. Hai notebook sử dụng cohort và thiết kế split khác nhau.")

    add_topic_page(doc, "2.2 Training Methodology", [
        "Cả hai notebook dùng MONAI/PyTorch và huấn luyện trên patch 96×96×96 voxel để phù hợp bộ nhớ GPU. Patch sampling làm tăng số mẫu huấn luyện từ mỗi volume và tập trung vào ROI, nhưng cũng có thể thay đổi phân bố lớp. Batch và số sample trên volume khác nhau giữa hai phiên bản.",
        "Hai mô hình được huấn luyện trong các notebook khác nhau: v2 chỉ chạy 3D U-Net dù có cấu hình kiến trúc khác; v3 in rõ thông báo bỏ qua U-Net và chỉ train Swin UNETR. Báo cáo không trình bày hai model như cùng một lần chạy.",
    ], table=(["Thuộc tính", "v2 – 3D U-Net", "v3 – Swin UNETR"], [
        ("Patch", "96³", "96³"), ("Batch size", "1", "2"),
        ("Samples/volume", "2", "2"), ("Epoch", "30 rồi resume tới 50", "20"),
        ("Fast validation", "20 ca", "2 ca"), ("Sliding overlap", "0,5", "0,5"),
    ], [4.7, 5.45, 5.45]))

    add_topic_page(doc, "2.2.1 Tiền xử lý và cache", [
        "Pipeline clip hoặc chuẩn hóa cường độ theo volume, ghép bốn kênh, crop/pad và tạo patch cho huấn luyện. Ở ứng dụng, inference_utils thực hiện clip/normalize trước khi chuyển thành tensor. Sự nhất quán giữa preprocessing train và inference là điều kiện bắt buộc; thay đổi percentile hoặc thứ tự kênh sẽ tạo domain shift ngay trong cùng dữ liệu.",
        "Cache/NPZ giúp giảm thời gian đọc và transform lặp lại. Tuy nhiên cache phải mang version và nguồn tạo. Nếu code preprocessing đổi nhưng cache cũ vẫn được dùng, notebook có thể báo kết quả không tương ứng với mã hiện tại. Khuyến nghị lưu hash cấu hình cùng mỗi artifact.",
    ], bullets=[
        "Xác nhận tất cả intensity hữu hạn sau chuẩn hóa.",
        "Lưu shape trước và sau transform.",
        "Kiểm tra mask không bị nội suy tuyến tính.",
        "Đồng nhất preprocessing giữa notebook và Gradio app.",
    ])

    add_topic_page(doc, "2.2.2 Tăng cường dữ liệu", [
        "Augmentation 3D nhằm tạo biến thiên vị trí và cường độ mà vẫn giữ cấu trúc giải phẫu. Các phép biến đổi thường gồm crop có định hướng nhãn, flip theo trục, biến đổi intensity và nhiễu nhẹ. Với segmentation, ảnh và mask phải dùng cùng biến đổi hình học; mask cần nội suy nearest-neighbor.",
        "Augmentation không thay thế dữ liệu ngoài miền. Hiệu quả cần được kiểm chứng qua ablation trên split cố định. Các phép biến đổi quá mạnh có thể tạo ảnh không thực tế hoặc làm biến dạng vùng ET nhỏ. Vì notebook chưa báo cáo ablation, báo cáo chỉ mô tả vai trò phương pháp, không gán mức cải thiện cụ thể.",
    ], table=(["Biến đổi", "Mục tiêu", "Rủi ro"], [
        ("Random crop", "Tăng patch chứa u", "Lệch phân bố background"),
        ("Flip 3D", "Tăng đa dạng không gian", "Cần áp dụng đồng bộ mask"),
        ("Intensity shift/scale", "Mô phỏng biến thiên scanner", "Làm mất tương phản nếu quá mạnh"),
        ("Noise", "Tăng ổn định", "Giảm tín hiệu vùng nhỏ"),
    ], [4.0, 6.0, 5.6]))

    add_topic_page(doc, "2.2.3 Kiến trúc 3D U-Net", [
        "3D U-Net mở rộng encoder–decoder U-Net sang tích chập ba chiều, dùng skip connection để kết hợp ngữ nghĩa sâu với chi tiết không gian [5]. Cấu hình v2 có các kênh 16, 32, 64, 128, 256; stride giảm kích thước bốn lần; mỗi tầng dùng hai residual unit, instance normalization và dropout 0,2.",
        "Mô hình có 4.811.129 tham số học được. Quy mô tương đối gọn giúp huấn luyện và suy luận dễ hơn trong điều kiện GPU hạn chế. Điểm yếu là receptive field và khả năng mô hình hóa quan hệ xa phụ thuộc độ sâu/tích chập; các vùng u lan tỏa có thể hưởng lợi từ ngữ cảnh rộng hơn.",
    ], table=(["Thành phần", "Cấu hình v2"], [
        ("Input", "4 × 96 × 96 × 96"), ("Channels", "16–32–64–128–256"),
        ("Strides", "2, 2, 2, 2"), ("Residual units", "2"),
        ("Normalization", "InstanceNorm"), ("Dropout", "0,2"),
        ("Trainable parameters", "4.811.129"),
    ], [6.0, 9.6]), source="Tham chiếu kiến trúc: Çiçek et al. [5]; cấu hình thực tế: SIC_Capstone_v2.ipynb [21].")

    add_topic_page(doc, "2.2.4 Kiến trúc Swin UNETR", [
        "Swin UNETR sử dụng encoder Swin Transformer phân cấp để học biểu diễn thể tích, kết hợp decoder kiểu U-Net qua skip connections [6]. Cơ chế shifted-window self-attention cân bằng khả năng nắm ngữ cảnh với chi phí tính toán, phù hợp ảnh y sinh 3D.",
        "Cấu hình v3 dùng feature_size 24 và gradient checkpointing, tổng 15.705.646 tham số học được. So với U-Net, mô hình lớn hơn khoảng 3,26 lần về số tham số. Gradient checkpointing giảm bộ nhớ activation bằng cách tính lại một phần trong backward, đổi lấy thời gian huấn luyện.",
    ], table=(["Thành phần", "Cấu hình v3"], [
        ("Input", "4 × 96 × 96 × 96"), ("Backbone", "Swin Transformer phân cấp"),
        ("Feature size", "24"), ("Checkpointing", "Bật"),
        ("Decoder", "UNETR-style với skip connection"),
        ("Trainable parameters", "15.705.646"),
    ], [6.0, 9.6]), source="Tham chiếu kiến trúc: Hatamizadeh et al. [6]; cấu hình thực tế: SIC_Capstone_v3.ipynb [22].")

    add_architecture_depth(doc)

    add_topic_page(doc, "2.2.5 Loss, optimizer và lịch học", [
        "Notebook v2 sử dụng DiceCELoss với include_background=False, AdamW, learning rate khởi tạo 2×10⁻⁴, weight decay 10⁻⁴; giai đoạn resume sử dụng learning rate 2×10⁻⁵. Notebook v3 cũng dùng AdamW 2×10⁻⁴ và weight decay 10⁻⁴, kết hợp cosine annealing và early stopping patience 5.",
        "Dice component trực tiếp tối ưu overlap, còn cross-entropy hỗ trợ phân biệt lớp voxel. Loại background khỏi Dice giảm sự áp đảo của nền. Dù vậy loss aggregate có thể che mất lớp ET nhỏ; cần theo dõi per-class validation và phân bố patch chứa từng nhãn.",
    ], table=(["Cấu hình", "v2", "v3"], [
        ("Loss", "DiceCE, bỏ background", "DiceCE/segmentation loss trong notebook"),
        ("Optimizer", "AdamW", "AdamW"), ("LR", "2e-4; resume 2e-5", "2e-4"),
        ("Weight decay", "1e-4", "1e-4"), ("Scheduler", "Theo cấu hình notebook", "Cosine annealing"),
        ("Early stopping", "Không phải tiêu chí chính", "Patience 5"),
    ], [4.6, 5.5, 5.5]))

    add_optimization_depth(doc)

    add_topic_page(doc, "2.2.6 Checkpoint và khả năng tái lập", [
        "Notebook v2 đạt best epoch 30 và tiếp tục resume đến epoch 50. Notebook v3 auto-resume checkpoint ở epoch 20. Việc resume giúp tận dụng phiên chạy trước nhưng có thể gây nhầm nếu optimizer state, scheduler hoặc split không khớp. Mỗi checkpoint cần đi kèm epoch, config, seed, code commit, dataset manifest và metric dùng để chọn.",
        "Báo cáo đọc kết quả đã lưu, không chạy lại hoặc thay đổi notebook. Đây là lựa chọn bảo toàn bằng chứng hiện tại. Một lần tái lập chính thức trong tương lai nên bắt đầu từ môi trường khóa phiên bản, xuất log máy/GPU và tái tạo metric từ checkpoint trên một script đánh giá độc lập.",
    ], bullets=[
        "Seed v2: 42; seed của mọi thư viện cần được ghi cùng deterministic flags.",
        "Checkpoint không được đánh giá test nhiều lần để quyết định hyperparameter.",
        "Tệp TorchScript/ONNX phải có input signature và preprocessing contract.",
        "Metric CSV phải chứa case ID để audit outlier và mask rỗng.",
    ])

    add_topic_page(doc, "2.2.7 Sliding-window inference", [
        "Thể tích MRI đầy đủ thường không vừa bộ nhớ GPU. Sliding-window chia volume thành các ROI 96³, suy luận từng cửa sổ với overlap 0,5 và trộn kết quả ở vùng giao. Cả v2 và v3 sử dụng cách này; v3 dùng sw_batch_size 2.",
        "Overlap lớn giảm đường nối giữa patch nhưng tăng số lần suy luận. Kết quả còn phụ thuộc blending mode và padding ở biên. Khi benchmark tốc độ, cần ghi shape volume, thiết bị, precision, warm-up, sw_batch_size và gồm/không gồm thời gian I/O. Con số 0,45 giây trong PDF mẫu chỉ là một output minh họa và không được xem là benchmark chuẩn.",
    ], figure=GENERATED / "end_to_end_workflow.png", caption="Hình 5. Luồng xử lý từ bốn ảnh MRI đến mask và báo cáo.")

    add_topic_page(doc, "2.2.8 Chỉ số đánh giá", [
        "Dice đo mức chồng lấp giữa mask dự đoán và ground truth; IoU cũng đo overlap nhưng nghiêm ngặt hơn ở cùng sai số. HD95 đo khoảng cách biên ở phân vị 95 và có đơn vị theo không gian ảnh, cung cấp góc nhìn về lỗi hình học mà Dice không phản ánh. Với ca rỗng, phải quy định rõ giá trị và số ca đóng góp.",
        "Báo cáo sử dụng WT, TC, ET và Mean Dice như notebook. HD95 chỉ có trong báo cáo v2. Để tránh trung bình sai, future benchmark nên báo cáo mean, median, độ lệch chuẩn, khoảng tin cậy bootstrap, số ca hợp lệ và conditional ET Dice trên ca có ground-truth ET.",
    ], table=(["Metric", "Diễn giải", "Càng tốt", "Điểm yếu"], [
        ("Dice", "2|P∩G|/(|P|+|G|)", "Cao", "Nhạy quy tắc mask rỗng"),
        ("IoU", "|P∩G|/|P∪G|", "Cao", "Không mô tả biên"),
        ("HD95", "Phân vị 95 khoảng cách biên", "Thấp", "Không xác định tự nhiên khi mask rỗng"),
        ("Volume error", "Sai lệch thể tích", "Gần 0", "Có thể triệt tiêu lỗi vị trí"),
    ], [2.7, 5.4, 2.3, 5.2]))

    add_inference_metric_depth(doc)

    add_topic_page(doc, "2.2.9 Thiết kế so sánh công bằng", [
        "Một đối chứng kiến trúc hợp lệ phải giữ cố định case list, split, preprocessing, augmentation, patch sampler, loss, số lần update hoặc compute budget, cách chọn checkpoint và script metric. Chỉ thay mô hình là cách gần nhất để quy khác biệt cho kiến trúc.",
        "Thí nghiệm hiện tại chưa đáp ứng điều kiện này: v2 và v3 khác cohort, batch, epoch, validation subset và test design. Vì vậy phần kết quả trình bày song song theo đúng nguồn, sau đó phân tích độ tin cậy. Không dùng phép trừ 0,8653 – 0,7361 để tuyên bố cải thiện 12,92 điểm phần trăm.",
    ], bullets=[
        "Khóa một split 70/15/15 hoặc cross-validation theo cùng case ID.",
        "Đánh giá cả hai checkpoint trên cùng held-out test.",
        "Stratify theo ET present/absent và kích thước tổn thương.",
        "Lặp ít nhất ba seed hoặc báo khoảng tin cậy theo case.",
        "Báo tài nguyên: VRAM, thời gian train, thời gian inference và kích thước model.",
    ])

    add_topic_page(doc, "2.3 Workflow", [
        "Workflow tổng thể đi từ hiểu bài toán, hiểu dữ liệu, chuẩn bị dữ liệu, modeling, evaluation đến deployment prototype. Đây là cách ánh xạ CRISP-DM với Quy trình SIC: Discover/Define tương ứng business/data understanding; Analyze đi sâu vào nguyên nhân và dữ liệu; Ideate/Prototype tương ứng modeling/deployment; Test tạo vòng feedback.",
        "Mỗi bước tạo một artifact có thể kiểm tra: problem statement, manifest, preprocessing config, checkpoint, prediction, metric table và test log. Artifact là cơ sở follow-up; nếu một chỉ số bất thường, nhóm quay về artifact gần nhất thay vì đoán nguyên nhân.",
    ], table=(["Pha", "Artifact", "Cổng kiểm tra"], [
        ("Understand", "Problem statement", "Phạm vi và người dùng rõ"),
        ("Data", "Manifest/EDA", "Nhãn và modality hợp lệ"),
        ("Prepare", "Cache/config", "Train–inference nhất quán"),
        ("Model", "Checkpoint/log", "Có thể resume và truy vết"),
        ("Evaluate", "Per-case metrics", "Không leakage, xử lý empty rõ"),
        ("Deploy", "Gradio prototype", "Có guardrail và output version"),
    ], [3.2, 6.2, 6.2]))

    add_topic_page(doc, "2.3.1 Quy trình huấn luyện và đánh giá", [
        "Train loader tạo patch tăng cường, mô hình dự đoán logits, loss được backpropagate và optimizer cập nhật tham số. Đến chu kỳ validation, sliding-window inference chạy trên volume hoặc subset; metric được tổng hợp và checkpoint tốt nhất được lưu theo tiêu chí định trước.",
        "Điểm then chốt là tách fast validation khỏi final evaluation. Fast validation có ích để phát hiện run hỏng, nhưng subset nhỏ không đủ cho kết luận. Full validation hoặc test phải chạy qua toàn bộ case list, lưu per-case metric và không bị can thiệp sau khi xem kết quả.",
    ], bullets=[
        "Log train loss theo epoch/step.",
        "Log Dice từng vùng, không chỉ mean.",
        "Lưu best criterion và epoch rõ ràng.",
        "Chạy final evaluation một lần trên tập khóa.",
        "Xuất bảng lỗi và ví dụ định tính đại diện.",
    ])

    add_topic_page(doc, "2.3.2 Quy trình suy luận người dùng", [
        "Người dùng chọn mô hình, tải bốn file MRI, khởi chạy suy luận và xem kết quả theo ba mặt phẳng. Ứng dụng tạo mask, bảng thể tích và cho phép xuất tệp. Tab hướng dẫn giải thích modality và luồng thao tác. Một MRI viewer hỗ trợ kiểm tra ảnh đầu vào ngoài kết quả segmentation.",
        "Luồng cần fail-fast khi thiếu file, đọc sai định dạng hoặc shape không khớp. Thông báo lỗi phải nêu file nào sai và cách khắc phục. Khi suy luận thành công, kết quả cần gắn case ID, model, checkpoint và thời gian tạo để tránh nhầm lẫn giữa lần chạy.",
    ], figure=GENERATED / "system_architecture.png", caption="Hình 6. Kiến trúc logic của ứng dụng Gradio và engine suy luận.")

    add_topic_page(doc, "2.4 System Design", [
        "Hệ thống tách UI, preprocessing, model loading, inference và postprocessing. Cách tách này cho phép thay mô hình mà không thay đổi toàn bộ giao diện, đồng thời tạo điểm kiểm thử đơn vị cho từng lớp. app.py điều phối luồng; inference_utils.py chứa chức năng chuẩn hóa, load model, sliding-window và hậu xử lý.",
        "Hai lựa chọn model 3D U-Net và Swin UNETR chia sẻ contract đầu vào bốn kênh và đầu ra segmentation. Contract cần được chính thức hóa bằng cấu hình thay vì hard-code: label mapping, ROI, intensity normalization, device, precision, checkpoint checksum và version.",
    ], table=(["Module", "Trách nhiệm", "Kiểm thử ưu tiên"], [
        ("UI", "Upload, chọn model, hiển thị", "Validation và thông báo lỗi"),
        ("Preprocess", "Load, normalize, tensor", "Shape, NaN, modality order"),
        ("Model service", "Load checkpoint, inference", "Compatibility và determinism"),
        ("Postprocess", "Mask, volume, overlay", "Label mapping và spacing"),
        ("Export", "NIfTI/PDF", "Metadata và khả năng mở lại"),
    ], [3.3, 6.1, 6.2]))

    add_topic_page(doc, "2.4.1 Thiết kế triển khai mô hình", [
        "Notebook v3 xuất TorchScript khoảng 67,65 MB và ONNX khoảng 62,80 MB. Đây là hướng giảm phụ thuộc vào code huấn luyện khi triển khai. Tuy nhiên xuất file thành công chưa đồng nghĩa parity; cần so logits/mask giữa PyTorch, TorchScript và ONNX trên cùng input với tolerance định trước.",
        "3D U-Net checkpoint quan sát có quy mô khoảng 55 MiB, còn Swin khoảng 67 MiB. Kích thước tệp không chỉ phụ thuộc số tham số mà còn state dict/format. Benchmark deployment phải đo thời gian load, peak VRAM, latency và throughput trên đúng phần cứng mục tiêu.",
    ], figure=GENERATED / "resource_comparison.png", caption="Hình 7. So sánh quy mô tham số và kích thước artifact triển khai.")

    add_topic_page(doc, "2.4.2 An toàn và ranh giới sử dụng", [
        "Giao diện phải hiển thị rõ đây là công cụ nghiên cứu. Báo cáo sinh tự động không được dùng cụm từ khẳng định chẩn đoán; nên dùng 'kết quả phân đoạn của mô hình' và 'thể tích ước tính'. Nếu dữ liệu không phù hợp distribution BraTS, hệ thống cần cảnh báo thay vì trả kết quả có vẻ chắc chắn.",
        "Cơ chế an toàn kỹ thuật gồm validation đầu vào, phát hiện output rỗng/bất thường, log lỗi, timeout, không giữ file quá lâu và kiểm tra checksum checkpoint. Cơ chế con người gồm review mask, tài liệu hướng dẫn và quy trình escalation khi kết quả mâu thuẫn ảnh gốc.",
    ], table=(["Rủi ro", "Guardrail hiện có/đề xuất", "Mức ưu tiên"], [
        ("Sai modality", "Kiểm tra filename/header và preview", "Cao"),
        ("Domain shift", "Cảnh báo dataset ngoài BraTS; external validation", "Cao"),
        ("Mask rỗng", "Thông báo và yêu cầu review", "Cao"),
        ("Nhầm model", "Hiển thị version/checksum", "Trung bình"),
        ("Rò rỉ dữ liệu", "Xử lý cục bộ và xóa tệp tạm", "Cao"),
    ], [4.2, 8.6, 2.8]))

    # CHAPTER 3
    next_page(doc)
    add_chapter(doc, "3 Results")
    add_heading(doc, "3.1 Data Preprocessing", 2)
    add_body(doc, "Các ảnh trực quan từ notebook cho thấy bốn modality được load đồng bộ và mask có thể overlay lên lát cắt. Đây là kiểm tra định tính quan trọng trước huấn luyện: nếu ảnh/mask lệch không gian, mô hình có thể học nhiễu dù loss vẫn giảm.")
    add_body(doc, "Báo cáo không tái chạy preprocessing. Hình và số liệu được trích trực tiếp từ output đã lưu trong notebook, nhằm bảo toàn trạng thái thực nghiệm. Kết luận chỉ dựa trên những gì artifact chứng minh.")
    add_picture(doc, ASSETS / "v2_cell9_output0.png", "Hình 8. Mẫu dữ liệu v2: bốn modality và ground-truth overlay trên năm ca.", max_height=5.7)
    add_source(doc, "Nguồn: output cell 9, SIC_Capstone_v2.ipynb.")

    add_topic_page(doc, "3.1.1 Kiểm tra trực quan dữ liệu v3", [
        "Notebook v3 cũng hiển thị đồng thời T1n, T1c, T2w, T2-FLAIR và mask. Ảnh cho phép xác nhận tương phản khác nhau giữa chuỗi và vị trí tổn thương. Tuy nhiên một vài ảnh mẫu không thay thế kiểm tra tự động toàn bộ cohort.",
        "Audit nên sinh montage ngẫu nhiên theo split và danh sách outlier: mask rỗng, mask quá lớn, intensity gần hằng, shape bất thường. Đặc biệt cần chọn các case có ET > 0 để xác minh mapping nhãn v3.",
    ], figure=ASSETS / "v3_cell13_output0.png", caption="Hình 9. Mẫu bốn modality và mask trong notebook v3.", source="Nguồn: output cell 13, SIC_Capstone_v3.ipynb.")

    add_topic_page(doc, "3.2 Exploratory Data Analysis (EDA)", [
        "EDA v2 biểu diễn phân bố intensity sau xử lý cho các modality. Các đường phân bố khác nhau phản ánh đặc trưng tương phản, đồng thời giúp phát hiện channel bất thường. Vì normalization được thực hiện theo volume, EDA cần đọc cùng thông tin transform để tránh so trực tiếp scale không tương đương.",
        "Đối với segmentation, EDA không nên dừng ở intensity. Cần thống kê thể tích WT/TC/ET, tỷ lệ ca có nhãn, tỷ lệ voxel nền và phân bố kích thước tổn thương. Những thống kê này quyết định sampling, loss và cách diễn giải metric.",
    ], figure=ASSETS / "v2_cell13_output0.png", caption="Hình 10. Phân bố intensity các modality trong notebook v2.", source="Nguồn: output cell 13, SIC_Capstone_v2.ipynb.")

    add_topic_page(doc, "3.2.1 EDA notebook v3", [
        "Đồ thị intensity v3 cho phép kiểm tra định tính pipeline dữ liệu của cohort 1.016 ca. Khi so với v2, khác biệt hình dạng phân bố có thể do cohort hoặc preprocessing; không đủ bằng chứng để quy cho dữ liệu tốt/xấu hơn nếu không xuất cùng thống kê trên cùng case.",
        "Đề xuất tạo bảng EDA versioned gồm mean/std intensity theo modality, percentile, số voxel non-zero, shape và spacing. Với nhãn, báo cáo ET-present count ở train/val/test là kiểm tra bắt buộc sau phát hiện full validation ET rỗng.",
    ], figure=ASSETS / "v3_cell11_output0.png", caption="Hình 11. Phân bố intensity các modality trong notebook v3.", source="Nguồn: output cell 11, SIC_Capstone_v3.ipynb.")

    add_topic_page(doc, "3.2.2 Mất cân bằng lớp và hiện tượng ET rỗng", [
        "Background chiếm phần lớn thể tích não, WT lớn hơn TC và ET thường là vùng nhỏ nhất. Điều này làm accuracy voxel không phù hợp và giải thích việc dùng Dice loss bỏ background. Nhưng ngay cả Dice cũng có thể tạo kết quả gây hiểu lầm khi lớp không xuất hiện.",
        "Trong full validation v3, cột ET volume bằng 0 cho mọi ca và Dice ET bằng 1,0. Theo quy tắc empty–empty, dự đoán rỗng khớp ground truth rỗng nên được điểm tuyệt đối. Về toán học điều này hợp lệ theo quy ước, nhưng về năng lực phân đoạn ET thì tập đó không tạo thử thách. Mean Dice do đó nhận thêm một thành phần 1,0 không đo khả năng phát hiện ET.",
    ], table=(["Cách báo cáo ET", "Ý nghĩa", "Khuyến nghị"], [
        ("Macro gồm empty–empty = 1", "Phản ánh cả khả năng không báo giả", "Báo riêng số ca rỗng"),
        ("Conditional Dice", "Chỉ ca ground truth ET > 0", "Bắt buộc cho năng lực phân đoạn ET"),
        ("False-positive volume", "Sai thể tích ET trên ca âm tính", "Bổ sung"),
        ("Detection sensitivity", "Có/không phát hiện ET", "Bổ sung theo ngưỡng"),
    ], [5.2, 6.0, 4.4]))

    add_topic_page(doc, "3.3 Modeling – Kết quả 3D U-Net", [
        "Đường học v2 cho thấy train loss giảm và validation Dice tăng qua các epoch, với best epoch 30 trong log. Các đường WT, TC và ET có động thái khác nhau, phản ánh độ khó từng vùng. Resume đến epoch 50 cần được phân biệt với checkpoint tốt nhất; epoch cuối không mặc nhiên là mô hình tốt nhất.",
        "Biểu đồ là bằng chứng theo dõi tối ưu nhưng không thay thế test. Kết quả held-out 188 ca được xem là chỉ số tổng kết của 3D U-Net trong báo cáo.",
    ], figure=ASSETS / "v2_cell27_output0.png", caption="Hình 12. Train loss và validation Dice của 3D U-Net v2.", source="Nguồn: output cell 27, SIC_Capstone_v2.ipynb.")

    add_topic_page(doc, "3.3.1 Chỉ số định lượng 3D U-Net", [
        "Trên validation, Mean Dice đạt 0,7555. Trên held-out test, Mean Dice giảm còn 0,7361, tương ứng chênh khoảng 0,0194. Đây là mức giảm hợp lý cần được đọc cùng phân bố per-case; nó cho thấy validation không hoàn toàn đại diện cho test nhưng không có dấu hiệu sụp giảm cực đoan ở mean.",
        "TC đạt Dice test cao nhất 0,7534, tiếp theo WT 0,7489 và ET 0,7061. HD95 khoảng 43,058 mm cho thấy vẫn tồn tại lỗi biên hoặc outlier đáng kể dù overlap tương đối. Không nên dùng Dice duy nhất để nói mask đã chính xác hình học.",
    ], table=(["Metric", "Validation v2", "Held-out test v2"], [
        ("Dice WT", "0,7613", "0,7489"),
        ("Dice TC", "0,7717", "0,7534"),
        ("Dice ET", "0,7334", "0,7061"),
        ("Mean Dice", "0,7555", "0,7361"),
        ("Mean IoU", "0,6490", "0,6310"),
        ("HD95", "43,087 mm", "43,058 mm"),
    ], [5.2, 5.2, 5.2]))

    add_topic_page(doc, "3.3.2 Phân tích định tính 3D U-Net", [
        "Hình dự đoán v2 hiển thị FLAIR, ground truth và mask U-Net trên axial, coronal và sagittal. Sự tương đồng tổng thể cho thấy mô hình học được vị trí tổn thương. Các khác biệt ở vùng biên hoặc cấu trúc nhỏ phù hợp với HD95 còn cao.",
        "Ví dụ định tính chỉ minh họa một ca và không được dùng để đại diện toàn tập. Bộ báo cáo tốt hơn nên chọn ca median Dice, ca tốt, ca xấu, ca ET nhỏ và ca prediction rỗng; mỗi ca gắn metric và thể tích.",
    ], figure=ASSETS / "v2_cell31_output1.png", caption="Hình 13. Ground truth và dự đoán 3D U-Net trên ba mặt phẳng.", source="Nguồn: output cell 31, SIC_Capstone_v2.ipynb; case minh họa BraTS-GLI-01021-000.")

    add_topic_page(doc, "3.3.3 Kết quả huấn luyện Swin UNETR", [
        "Notebook v3 huấn luyện Swin UNETR 20 epoch và ghi best fast-validation Dice 0,9538. Fast validation chỉ sử dụng hai ca, vì vậy con số này chủ yếu dùng theo dõi run. Đường ET giữ ở 1,0 là dấu hiệu cần audit hơn là bằng chứng hoàn hảo.",
        "Đường học cho thấy mô hình hội tụ trên subset theo dõi, nhưng không có held-out test để đo generalization. Kết luận chính phải dựa trên full validation 204 ca và kèm caveat về ET.",
    ], figure=ASSETS / "v3_cell24_output0.png", caption="Hình 14. Đường học Swin UNETR; ET cố định ở 1,0 cần được kiểm tra.", source="Nguồn: output cell 24, SIC_Capstone_v3.ipynb.")

    add_topic_page(doc, "3.3.4 Chỉ số định lượng Swin UNETR", [
        "Trên full validation 204 ca, Dice WT đạt 0,8498, Dice TC 0,7461 và Dice ET 1,0000; Mean Dice 0,8653. Nếu loại thành phần ET rỗng khỏi trung bình và chỉ lấy WT/TC, mean hai vùng xấp xỉ 0,7980. Con số này không thay cho official mean, nhưng giúp thấy mức đóng góp của ET=1,0.",
        "Notebook không báo HD95 và không có test độc lập. Do đó bằng chứng mạnh nhất của Swin ở trạng thái hiện tại là khả năng overlap WT tốt trên full validation. Năng lực ET chưa được kiểm tra thực chất trên tập có ET.",
    ], table=(["Metric", "Full validation v3", "Diễn giải"], [
        ("Dice WT", "0,849789", "Tín hiệu tốt nhất của v3"),
        ("Dice TC", "0,746115", "Gần mức TC của U-Net nhưng khác tập"),
        ("Dice ET", "1,000000", "Empty–empty ở 204/204 ca"),
        ("Mean Dice", "0,865301", "Bị nâng bởi ET=1"),
        ("Mean WT/TC", "≈0,797952", "Phân tích bổ sung, không phải metric notebook"),
        ("HD95", "Không báo cáo", "Khoảng trống"),
    ], [4.2, 4.2, 7.2]))

    add_topic_page(doc, "3.3.5 Phân tích định tính Swin UNETR", [
        "Ảnh ba mặt phẳng của v3 cho thấy mask Swin bám vùng tổn thương tổng thể trên ca minh họa. Kênh FLAIR và overlay giúp quan sát WT, trong khi đánh giá ET cần chọn ca có nhãn ET khác rỗng. Hình hiện tại không đủ để chứng minh năng lực đó.",
        "Visual QA nên bật/tắt từng vùng, điều chỉnh opacity và hiển thị sai khác false-positive/false-negative. Khi có ground truth, một bản đồ lỗi ba màu giúp người review thấy nơi mô hình thiếu hoặc thừa thay vì chỉ nhìn hai overlay riêng.",
    ], figure=ASSETS / "v3_cell26_output1.png", caption="Hình 15. Ground truth và dự đoán Swin UNETR trên ba mặt phẳng.", source="Nguồn: output cell 26, SIC_Capstone_v3.ipynb.")

    add_topic_page(doc, "3.3.6 So sánh kiến trúc", [
        "3D U-Net dùng inductive bias tích chập mạnh, tham số ít và quy trình dễ tối ưu. Swin UNETR có khả năng mô hình hóa ngữ cảnh xa qua attention phân cấp nhưng cần nhiều tham số và tài nguyên hơn. Chọn mô hình phải dựa trên benchmark cùng điều kiện, không chỉ dựa vào độ mới của kiến trúc.",
        "Trong dự án, U-Net là baseline có test độc lập; Swin là mô hình nâng cao có full validation nhưng thiếu test. Swin có WT cao hơn trong số liệu as-reported; TC xấp xỉ; ET không thể so do empty-mask. Vì vậy recommendation hiện tại là giữ cả hai trong prototype, dùng U-Net làm baseline tin cậy và ưu tiên tái đánh giá Swin.",
    ], table=(["Tiêu chí", "3D U-Net v2", "Swin UNETR v3"], [
        ("Họ kiến trúc", "CNN encoder–decoder", "Transformer phân cấp + decoder"),
        ("Tham số", "4.811.129", "15.705.646"),
        ("Ngữ cảnh", "Tích lũy qua convolution", "Window attention phân cấp"),
        ("Bằng chứng cuối", "Held-out test 188", "Full validation 204"),
        ("Điểm mạnh quan sát", "TC/ET ổn định hơn về bằng chứng", "WT as-reported cao"),
        ("Điểm yếu", "HD95 còn cao", "ET chưa được thử thách; không test"),
    ], [4.0, 5.8, 5.8]))

    add_topic_page(doc, "3.3.7 So sánh chỉ số as-reported", [
        "Biểu đồ đặt U-Net test v2 cạnh Swin validation v3 để người đọc nhìn thấy số liệu đã báo cáo, không phải để khẳng định thắng thua. Nhãn tập đánh giá được giữ trong chú giải. Đặc biệt, cột ET v3 được đánh dấu là empty–empty.",
        "Nếu chỉ nhìn Mean Dice, Swin cao hơn U-Net 0,1292. Nhưng chênh lệch này trộn ảnh hưởng kiến trúc với cohort, split, validation/test, batch, epoch và quy tắc ET. Phần chênh do ET=1 cũng đáng kể. Một phép so sánh khoa học phải chạy lại trên cùng held-out test.",
    ], figure=GENERATED / "metric_comparison.png", caption="Hình 16. Chỉ số as-reported của hai mô hình dưới hai giao thức khác nhau.")

    add_topic_page(doc, "3.3.8 So sánh tài nguyên và khả năng triển khai", [
        "Swin có khoảng 3,26 lần số tham số của U-Net, nhưng kích thước artifact quan sát chỉ lớn hơn khoảng 1,22 lần; nguyên nhân là format lưu khác nhau và state đi kèm. Không thể suy latency trực tiếp từ tham số. Attention, sliding window, precision và GPU ảnh hưởng đáng kể.",
        "Đối với ứng dụng cục bộ, U-Net phù hợp khi ưu tiên footprint và tính đơn giản. Swin phù hợp để nghiên cứu ngữ cảnh rộng nếu phần cứng đáp ứng. Quyết định production cần benchmark peak VRAM, P50/P95 latency và failure rate trên máy mục tiêu.",
    ], table=(["Thuộc tính", "3D U-Net", "Swin UNETR", "Nhận xét"], [
        ("Tham số", "4,811 triệu", "15,706 triệu", "Swin ≈3,26×"),
        ("Artifact", "≈55,1 MiB", "≈67,2 MiB", "Khác format"),
        ("ROI", "96³", "96³", "Tương đồng"),
        ("Batch train", "1", "2", "Không công bằng để suy tốc độ"),
        ("Export", "Checkpoint", "TorchScript + ONNX", "Swin có thêm artifact"),
    ], [3.6, 3.5, 3.5, 5.0]), figure=GENERATED / "resource_comparison.png", caption="Hình 17. Quy mô tham số và artifact của hai mô hình.")

    add_topic_page(doc, "3.3.9 Mối đe dọa đến tính hợp lệ", [
        "Internal validity bị ảnh hưởng bởi khác cohort, split và hyperparameter. Construct validity bị ảnh hưởng bởi Mean Dice chứa ET empty–empty và thiếu HD95 v3. External validity bị giới hạn do chỉ sử dụng BraTS, chưa có dữ liệu bệnh viện mới. Conclusion validity bị hạn chế vì chưa có lặp seed hay khoảng tin cậy.",
        "Những hạn chế này không phủ nhận giá trị thực nghiệm; chúng xác định mức độ tuyên bố phù hợp. Dự án chứng minh pipeline và cho thấy tín hiệu hiệu năng, nhưng chưa chứng minh một kiến trúc vượt trội phổ quát hay sẵn sàng lâm sàng.",
    ], table=(["Loại hợp lệ", "Rủi ro", "Biện pháp"], [
        ("Internal", "Protocol khác nhau", "Benchmark chung"),
        ("Construct", "Metric bị empty-mask chi phối", "Conditional metric"),
        ("External", "Một nguồn dữ liệu", "External/multi-site validation"),
        ("Conclusion", "Không CI/replicate", "Bootstrap và nhiều seed"),
    ], [3.3, 6.1, 6.2]))

    add_topic_page(doc, "3.3.10 Kết luận lựa chọn mô hình", [
        "Nếu cần một baseline để tiếp tục phát triển ngay trong phạm vi bằng chứng hiện tại, 3D U-Net là lựa chọn thận trọng vì có held-out test, ít tham số và chỉ số ET không dựa trên tập hoàn toàn rỗng. Nếu mục tiêu là nghiên cứu tiềm năng nâng Dice WT, Swin UNETR đáng để tiếp tục nhưng phải sửa thiết kế đánh giá trước.",
        "Ứng dụng có thể giữ selector hai mô hình để demo, song giao diện nên hiển thị nhãn 'experimental' cho Swin và không trình bày Mean Dice 0,8653 mà thiếu chú thích. Sau benchmark chung, lựa chọn production phải dựa trên Pareto hiệu năng–latency–VRAM–độ ổn định.",
    ], bullets=[
        "Khuyến nghị ngắn hạn: U-Net làm baseline tham chiếu.",
        "Khuyến nghị thực nghiệm: audit ET và đánh giá Swin trên test chung.",
        "Khuyến nghị sản phẩm: hiển thị model version và mức bằng chứng.",
        "Không khuyến nghị: kết luận Swin tốt hơn chỉ từ hai mean Dice hiện tại.",
    ])

    add_topic_page(doc, "3.4 User Interface", [
        "Ứng dụng Gradio cho phép chọn 3D U-Net hoặc Swin UNETR, tải bốn ảnh MRI NIfTI, chạy phân đoạn và xem kết quả. Các tab gồm segmentation, MRI viewer và hướng dẫn. Đầu ra gồm hình ba mặt phẳng, bảng thể tích, mask NIfTI và PDF minh họa.",
        "Thiết kế này hoàn thành cầu nối từ notebook tới người dùng kỹ thuật. Điểm cần cải tiến là validation metadata, progress indicator, lịch sử chạy, biểu diễn uncertainty và cảnh báo domain shift. Mọi output cần gắn thông báo nghiên cứu.",
    ], table=(["Bước", "Thao tác", "Phản hồi hệ thống"], [
        ("1", "Chọn mô hình", "Hiển thị tên/version"),
        ("2", "Tải 4 modality", "Kiểm tra file/shape"),
        ("3", "Chạy segmentation", "Tiến trình và lỗi rõ ràng"),
        ("4", "Review overlay", "3 mặt phẳng và volume"),
        ("5", "Export", "NIfTI/PDF có metadata"),
    ], [2.0, 6.2, 7.4]))

    add_topic_page(doc, "3.4.1 Trực quan hóa đầu ra", [
        "Output axial của Swin đặt FLAIR, ground truth và prediction cạnh nhau. Cách bố trí giúp so nhanh vùng tổn thương nhưng phụ thuộc lát cắt được chọn. Một ca có nhiều tổn thương hoặc vùng nhỏ ngoài lát trung tâm có thể bị bỏ sót.",
        "Giao diện nên cho người dùng di chuyển slice, đồng bộ ba mặt phẳng và bật/tắt từng label. Việc hiển thị legend màu và opacity nhất quán giữa notebook, app và PDF tránh nhầm WT/TC/ET.",
    ], figure=ASSETS / "v3_cell26_output2.png", caption="Hình 18. FLAIR, ground truth và dự đoán Swin trên lát axial.", source="Nguồn: output cell 26, SIC_Capstone_v3.ipynb.")

    add_topic_page(doc, "3.4.2 Báo cáo kết quả minh họa", [
        "PDF mẫu một trang chứa case ID, thời gian suy luận, thể tích vùng và ảnh ba mặt phẳng. Đây là output hữu ích cho demo và truy vết, nhưng nội dung 'diagnostic finding' cần được thay bằng ngôn ngữ trung tính nếu chưa qua thẩm định y tế.",
        "Con số thời gian 0,45 giây và thể tích trong PDF chỉ thuộc một lần chạy mẫu; không được dùng như benchmark hoặc số liệu lâm sàng tổng quát. Báo cáo tương lai nên thêm model/checkpoint hash, preprocessing version, timestamp, device và disclaimer.",
    ], figure=ASSETS / "sample_clinical_report.png", caption="Hình 19. Báo cáo PDF một trang do ứng dụng tạo ở ca minh họa.", source="Nguồn: doc/Report_BraTS-GLI-2026_nzd2b1s9.pdf.")

    add_topic_page(doc, "3.5 Testing and Improvements", [
        "Kiểm thử được tổ chức thành bốn lớp: unit test cho preprocessing/postprocessing, integration test cho model load và UI, model evaluation trên dataset khóa, và usability test với người dùng mục tiêu. Một lỗi ở lớp sớm phải chặn lớp sau; ví dụ shape mismatch không nên đi đến inference.",
        "Repository chứng minh luồng chính qua artifact ứng dụng và output mẫu, nhưng chưa có bộ test người dùng hay clinical validation. Bảng kiểm thử dưới đây phân biệt điều đã quan sát với điều cần bổ sung.",
    ], table=(["Nhóm test", "Ví dụ", "Trạng thái bằng chứng"], [
        ("Input", "Thiếu modality, file hỏng, shape mismatch", "Cần hệ thống hóa"),
        ("Preprocess", "NaN, constant image, affine khác", "Cần unit test"),
        ("Model", "Checkpoint tương thích, output shape", "Có luồng chạy mẫu"),
        ("Metric", "Empty masks, label mapping", "Phát hiện vấn đề ET"),
        ("UI", "Upload–review–download", "Có prototype"),
        ("Usability", "Task completion/error/time", "Chưa thực hiện"),
    ], [3.2, 7.5, 4.9]))

    add_topic_page(doc, "3.5.1 Kiểm thử chức năng đề xuất", [
        "Mỗi test case cần precondition, input, bước, expected output và bằng chứng. Dữ liệu test nên gồm ca bình thường, ca u lớn, ca u nhỏ, ET có/không, file thiếu và file sai shape. Không dùng dữ liệu thật có định danh trong demo.",
        "Test regression cần chạy khi thay preprocessing, label mapping hoặc model export. Một checksum mask cho fixture nhỏ có thể phát hiện thay đổi ngoài ý muốn, nhưng tolerance cần phù hợp với backend/precision.",
    ], table=(["ID", "Tình huống", "Kết quả mong đợi"], [
        ("FT-01", "Đủ 4 modality hợp lệ", "Tạo mask và viewer"),
        ("FT-02", "Thiếu T1c", "Chặn và nêu file thiếu"),
        ("FT-03", "Shape không đồng nhất", "Chặn trước inference"),
        ("FT-04", "Output rỗng", "Cảnh báo và vẫn cho review"),
        ("FT-05", "Đổi model", "Load đúng checkpoint/version"),
        ("FT-06", "Xuất PDF/NIfTI", "File mở lại được, metadata đúng"),
    ], [2.2, 6.0, 7.4]))

    add_topic_page(doc, "3.5.2 Kiểm thử độ bền và lỗi", [
        "Robustness test cần bao phủ intensity cực trị, file lớn, header thiếu, affine khác, orientation đảo và thiết bị thiếu bộ nhớ. Hệ thống không được treo im lặng; cần giải phóng tài nguyên và trả thông báo có thể hành động.",
        "Model robustness cần noise/corruption có kiểm soát và dữ liệu ngoài site. Nếu confidence không đáng tin, giao diện không nên hiển thị phần trăm giả. Một cơ chế đơn giản hơn là cảnh báo out-of-distribution theo metadata/intensity và luôn yêu cầu review.",
    ], bullets=[
        "Giới hạn kích thước upload và thời gian xử lý.",
        "Bắt exception ở bước load NIfTI, model và export.",
        "Log lỗi không chứa dữ liệu ảnh hoặc định danh.",
        "Thu hồi bộ nhớ GPU sau mỗi request hoặc theo session.",
        "Kiểm thử nhiều request liên tiếp để phát hiện leak.",
    ])

    add_topic_page(doc, "3.5.3 Kế hoạch usability test", [
        "Theo bước Prototype/Test của Quy trình SIC, nhóm nên tuyển người dùng đại diện và giao nhiệm vụ cụ thể: tải đúng bốn modality, chọn model, xác định slice có u, đọc bảng thể tích và tải mask. Người quan sát ghi thời gian, lỗi, câu hỏi và điểm người dùng do dự.",
        "Mẫu nhỏ 5–8 người dùng kỹ thuật có thể phát hiện nhiều vấn đề giao diện ban đầu; tuy nhiên không đủ cho tuyên bố thống kê. Nếu có chuyên gia y tế, cần protocol riêng và phê duyệt phù hợp. Mọi phản hồi được chuyển thành issue có severity, owner, deadline và retest criterion.",
    ], table=(["Chỉ số", "Cách đo", "Tiêu chí pilot đề xuất"], [
        ("Task success", "Hoàn thành không trợ giúp", "≥80%"),
        ("Input error", "Số lần chọn sai modality", "Giảm sau hướng dẫn"),
        ("Time on task", "Từ upload đến review", "Thiết lập baseline"),
        ("Comprehension", "Giải thích đúng WT/TC/ET", "Không nhầm label"),
        ("Trust calibration", "Nhận ra mask sai", "Không tin tuyệt đối"),
    ], [4.0, 6.4, 5.2]))

    # CHAPTER 4
    next_page(doc)
    add_chapter(doc, "4 Projected Impact")
    add_heading(doc, "4.1 Accomplishments and Benefits", 2)
    for p in [
        "Dự án hoàn thành một chuỗi artifact từ dữ liệu đến ứng dụng: notebook v2 cho baseline 3D U-Net, notebook v3 cho Swin UNETR, checkpoint và tệp triển khai, công cụ Gradio, mask/ảnh và báo cáo PDF mẫu. Báo cáo này bổ sung lớp truy vết và so sánh có điều kiện.",
        "Lợi ích trực tiếp là đào tạo quy trình AI y sinh 3D, minh họa khác biệt CNN–Transformer và tạo nền cho benchmark tiếp theo. Lợi ích tiềm năng đối với workflow là tạo mask sơ bộ và thống kê nhất quán; lợi ích này chưa được đo bằng nghiên cứu người dùng hoặc thử nghiệm lâm sàng.",
    ]:
        add_body(doc, p)
    add_bullets(doc, [
        "Có baseline với held-out test rõ ràng.",
        "Có mô hình Transformer và artifact TorchScript/ONNX.",
        "Có giao diện đa mô hình và đầu ra trực quan.",
        "Đã phát hiện vấn đề đánh giá ET cần follow-up.",
    ])

    add_topic_page(doc, "4.1.1 Lợi ích kỹ thuật và học thuật", [
        "Về kỹ thuật, dự án cho thấy khả năng xử lý NIfTI đa chuỗi, huấn luyện segmentation 3D dưới giới hạn bộ nhớ và dùng sliding-window inference. Về học thuật, hai notebook tạo trường hợp điển hình để học cách đọc metric theo protocol và nhận diện chỉ số bị thổi phồng bởi mask rỗng.",
        "Điểm đáng giá không chỉ là mức Dice mà còn là năng lực đặt câu hỏi đúng: test có độc lập không, cohort có giống nhau không, ET có hiện diện không, metric xử lý empty thế nào và checkpoint được chọn ra sao. Các câu hỏi này quyết định độ tin cậy của mọi kết luận.",
    ], table=(["Năng lực", "Minh chứng"], [
        ("Data engineering 3D", "Load/chuẩn hóa/caching bốn modality"),
        ("Modeling", "3D U-Net và Swin UNETR"),
        ("Evaluation", "Dice, IoU, HD95, per-region analysis"),
        ("Deployment", "Gradio, NIfTI, PDF, TorchScript/ONNX"),
        ("Responsible AI", "Caveat, threat to validity, guardrail"),
    ], [5.5, 10.1]))

    add_topic_page(doc, "4.1.2 Lợi ích đối với người dùng mục tiêu", [
        "Đối với nghiên cứu viên, một giao diện thống nhất giảm nhu cầu chạy trực tiếp từng cell và giúp chia sẻ demo. Đối với người review, ảnh ba mặt phẳng và bảng thể tích tạo điểm bắt đầu để kiểm tra. Đối với nhóm phát triển, cấu trúc module hỗ trợ thay checkpoint và thêm kiểm thử.",
        "Những lợi ích này là projected impact. Để chuyển thành accomplished impact, nhóm phải đo baseline và after-use: số bước, thời gian, tỷ lệ lỗi, độ chính xác hiệu chỉnh và mức hiểu disclaimer. Chỉ khi có dữ liệu đó mới được phát biểu mức tiết kiệm cụ thể.",
    ], bullets=[
        "Giảm ma sát thao tác cho người dùng kỹ thuật.",
        "Tăng khả năng quan sát output thay vì chỉ nhận file mask.",
        "Hỗ trợ thảo luận mô hình bằng cùng một case.",
        "Tạo đầu ra có thể lưu và xem lại.",
    ])

    add_topic_page(doc, "4.1.3 Đạo đức, thiên lệch và giới hạn", [
        "Mô hình học từ dữ liệu challenge có thể không đại diện mọi quần thể, máy chụp và protocol. Domain shift có thể tạo lỗi có hệ thống trên nhóm ít được đại diện. Không có external validation nên báo cáo không khẳng định công bằng hay tổng quát hóa.",
        "Các vấn đề đạo đức gồm quyền riêng tư, đồng thuận sử dụng dữ liệu, automation bias và trách nhiệm khi mask sai. Cách giảm thiểu là data governance, model card, human-in-the-loop, logging, cảnh báo rõ và không dùng ngôn ngữ chẩn đoán. Khi mở rộng sang dữ liệu bệnh viện, cần quy trình đạo đức và pháp lý phù hợp.",
    ], table=(["Vấn đề", "Biểu hiện", "Kiểm soát"], [
        ("Bias", "Hiệu năng khác site/quần thể", "Phân tầng và external validation"),
        ("Privacy", "MRI/metadata nhạy cảm", "De-identification, access control"),
        ("Automation bias", "Tin mask không kiểm tra", "Overlay, disclaimer, review"),
        ("Accountability", "Không rõ model/version", "Audit log và model card"),
    ], [3.6, 6.2, 5.8]))

    add_topic_page(doc, "4.2 Future Improvements", [
        "Ưu tiên cao nhất là benchmark công bằng. Nhóm tạo một manifest thống nhất, cố định split, xác minh nhãn ET, chạy cả U-Net và Swin với cùng preprocessing/loss/budget, sau đó đánh giá một lần trên held-out test. Per-case CSV và bootstrap confidence interval phải được lưu.",
        "Ưu tiên thứ hai là cải thiện sản phẩm: validation đầu vào, metadata report, uncertainty/quality flag, visual error map và mask editing. Ưu tiên thứ ba là external validation và đánh giá chuyên gia. Thứ tự này tránh đầu tư UI sâu trước khi độ tin cậy mô hình được xác minh.",
    ], table=(["Ưu tiên", "Hạng mục", "Kết quả mong đợi"], [
        ("P0", "Audit nhãn ET và case list", "Giải thích nguyên nhân ET rỗng"),
        ("P0", "Benchmark chung", "So sánh hợp lệ"),
        ("P1", "Input validation và model metadata", "Giảm lỗi thao tác"),
        ("P1", "Test automation", "Regression có kiểm soát"),
        ("P2", "External/expert validation", "Bằng chứng tổng quát hóa"),
    ], [2.0, 7.0, 6.6]))

    add_topic_page(doc, "4.2.1 Quy trình follow-up vấn đề", [
        "Mỗi vấn đề được ghi thành issue với bằng chứng, mức ảnh hưởng, giả thuyết nguyên nhân, owner, hành động, tiêu chí hoàn tất và ngày retest. Issue chỉ đóng khi có bằng chứng kiểm thử, không đóng vì đã sửa code. Nếu retest thất bại, vòng quay lại Analyze để cập nhật nguyên nhân.",
        "Vấn đề ET v3 được xếp P0 vì ảnh hưởng trực tiếp kết luận mô hình. Bước đầu không phải train lại ngay mà là kiểm tra case list, raw labels, mapping và metric. Sau khi xác định nguyên nhân, nhóm mới quyết định tái tạo split, sửa pipeline hoặc chọn dataset phù hợp.",
    ], table=(["Issue", "Bằng chứng", "Hành động", "Definition of Done"], [
        ("ET-EMPTY-V3", "204/204 ET volume=0", "Audit raw label và mapping", "Có ca ET+, metric conditional hợp lệ"),
        ("NO-TEST-V3", "Không held-out test", "Khóa test chung", "Hai model đánh giá cùng test"),
        ("PROTOCOL-DIFF", "Cohort/hyperparameter khác", "Benchmark contract", "Chỉ thay kiến trúc"),
        ("UI-VALIDATION", "Upload phụ thuộc tên/file", "Thêm kiểm tra header", "Test FT-02/03 pass"),
        ("REPORT-METADATA", "PDF thiếu version", "Bổ sung model card fields", "Output truy vết được"),
    ], [3.3, 4.3, 4.2, 4.2]))

    add_topic_page(doc, "4.2.2 Kế hoạch thử nghiệm đối chứng", [
        "Benchmark đề xuất sử dụng cùng 1.251 ca nếu license và artifact cho phép, split cố định 875/188/188 như v2 hoặc một split mới khóa trước khi chạy. Hai mô hình dùng cùng patch 96³, augmentation, loss, optimizer policy, số update và metric script. Hyperparameter đặc thù kiến trúc được ghi trước.",
        "Kết quả báo per-case và tổng hợp theo WT/TC/ET, kèm ET-present subset. Dùng bootstrap theo case để ước lượng khoảng tin cậy chênh lệch. Nếu khoảng tin cậy chứa 0, không tuyên bố superiority. Đồng thời báo latency/VRAM để có kết luận Pareto.",
    ], bullets=[
        "Pre-register protocol và primary metric.",
        "Không xem test trong quá trình tuning.",
        "Kiểm tra parity preprocessing bằng fixture.",
        "Chạy nhiều seed nếu compute cho phép.",
        "Lưu environment, commit và checkpoint checksum.",
    ])

    add_topic_page(doc, "4.2.3 Kế hoạch kiểm chứng với chuyên gia", [
        "Sau khi benchmark kỹ thuật ổn định, có thể thiết kế study nhỏ với chuyên gia. Mẫu ca phải đại diện kích thước, vị trí và trạng thái ET; chuyên gia xem ảnh và mask, chấm mức cần chỉnh sửa, lỗi nguy hiểm và thời gian hiệu chỉnh. Study cần quy trình đạo đức, dữ liệu phù hợp và không ảnh hưởng chăm sóc bệnh nhân.",
        "Kết quả nên gồm tỷ lệ mask chấp nhận được cho nghiên cứu, thời gian chỉnh sửa so baseline, vùng lỗi thường gặp và inter-rater agreement. Feedback được dùng để xác định requirement mới; không dùng vài nhận xét định tính để tuyên bố hiệu quả lâm sàng.",
    ], table=(["Câu hỏi", "Đo lường", "Quyết định"], [
        ("Mask có hữu ích không", "Accept/edit/reject", "Ngưỡng sử dụng nghiên cứu"),
        ("Lỗi ở đâu", "Loại lỗi theo WT/TC/ET", "Ưu tiên cải thiện"),
        ("Có tiết kiệm công sức", "Thời gian hiệu chỉnh", "ROI của công cụ"),
        ("Người dùng hiểu giới hạn", "Câu hỏi comprehension", "Thiết kế cảnh báo"),
    ], [5.0, 5.6, 5.0]))

    add_topic_page(doc, "4.2.4 Lộ trình triển khai", [
        "Lộ trình ba giai đoạn được đề xuất. Giai đoạn 1 làm sạch bằng chứng: audit ET, benchmark chung và test automation. Giai đoạn 2 củng cố prototype: validation, metadata, performance và usability. Giai đoạn 3 nghiên cứu bên ngoài: external data, expert review và governance.",
        "Mỗi giai đoạn có gate rõ. Không tiến sang pilot chuyên gia nếu model evaluation còn mơ hồ; không gọi production nếu chưa có monitoring và rollback. Lộ trình cho phép dự án tiếp tục có kiểm soát thay vì tối ưu một chỉ số đơn lẻ.",
    ], table=(["Giai đoạn", "Thời lượng gợi ý", "Gate"], [
        ("1. Evidence hardening", "4–6 tuần", "Benchmark chung và ET audit pass"),
        ("2. Prototype hardening", "4–8 tuần", "Functional/robustness/usability pass"),
        ("3. External study", "Phụ thuộc phê duyệt", "Protocol và expert evaluation"),
    ], [5.0, 4.5, 6.1]))

    # Reviews
    next_page(doc)
    add_chapter(doc, "5 Team Member Review and Comment")
    add_body(doc, "Bảng dưới đây giữ đúng mục đánh giá thành viên của mẫu. Nhóm điền điểm và nhận xét sau buổi retrospective, dựa trên minh chứng đóng góp thay vì tự đánh giá chung chung.")
    add_table(doc, ["No.", "Team member", "Contribution evidence", "Score", "Comment"], [
        (str(i), "[BỔ SUNG]", "[BỔ SUNG: commit/notebook/test/report]", "[  ]/10", "[BỔ SUNG]") for i in range(1, 6)
    ], [1.2, 3.6, 5.2, 2.2, 3.4], font_size=8.7)
    add_heading(doc, "Retrospective của nhóm", 2)
    add_table(doc, ["Nội dung", "Ghi nhận"], [
        ("Điều làm tốt", "Đã xây dựng hai pipeline mô hình, ứng dụng và artifact đầu ra."),
        ("Điều cần cải thiện", "Chuẩn hóa split/metric và tăng kiểm thử tự động."),
        ("Bài học", "Một chỉ số cao phải được kiểm tra theo dữ liệu và quy tắc metric."),
        ("Cam kết follow-up", "Đóng P0 ET-EMPTY-V3 và NO-TEST-V3 trước tuyên bố so sánh."),
    ], [4.2, 11.4])

    next_page(doc)
    add_chapter(doc, "6 Instructor Review and Comment")
    add_body(doc, "Phần này dành cho giảng viên/hội đồng. Các ô được để trống có chủ đích theo form mẫu.")
    add_table(doc, ["Evaluation item", "Score", "Instructor comment"], [
        ("Problem definition and relevance", "[  ]/10", ""),
        ("Data and methodology", "[  ]/10", ""),
        ("Modeling and comparison", "[  ]/10", ""),
        ("System prototype", "[  ]/10", ""),
        ("Testing and responsible AI", "[  ]/10", ""),
        ("Report quality and presentation", "[  ]/10", ""),
        ("Overall", "[  ]/10", ""),
    ], [6.2, 2.4, 7.0])
    add_heading(doc, "General comment", 2)
    add_table(doc, ["Comment"], [("\n\n\n\n\n\n",)], [15.6])
    add_body(doc, "Instructor name: ____________________    Signature: ____________________    Date: ____/____/________")

    # References
    next_page(doc)
    add_chapter(doc, "REFERENCES")
    add_heading(doc, "Tài liệu tham khảo 1", 2)
    refs1 = [
        "[1] Center for Biomedical Image Computing & Analytics, University of Pennsylvania, “Multimodal Brain Tumor Segmentation Challenge (BraTS),” official challenge portal.",
        "[2] B. H. Menze et al., “The Multimodal Brain Tumor Image Segmentation Benchmark (BRATS),” IEEE Transactions on Medical Imaging, 2015. DOI: 10.1109/TMI.2014.2377694.",
        "[3] S. Bakas et al., “Advancing The Cancer Genome Atlas glioma MRI collections with expert segmentation labels and radiomic features,” Scientific Data, 2017. DOI: 10.1038/sdata.2017.117.",
        "[4] BraTS 2023 Adult Glioma Training Dataset, Zenodo record 7837974, 2023.",
        "[5] Ö. Çiçek et al., “3D U-Net: Learning Dense Volumetric Segmentation from Sparse Annotation,” MICCAI, 2016. DOI: 10.1007/978-3-319-46723-8_49.",
        "[6] A. Hatamizadeh et al., “Swin UNETR: Swin Transformers for Semantic Segmentation of Brain Tumors in MRI Images,” WACV, 2022. DOI: 10.1109/WACV51458.2022.00181.",
        "[7] M. J. Cardoso et al., “MONAI: An open-source framework for deep learning in healthcare,” arXiv:2211.02701, 2022.",
        "[8] CBTRUS, “Statistical Report: Primary Brain and Other Central Nervous System Tumors Diagnosed in the United States in 2018–2022,” Neuro-Oncology, vol. 27, Supplement 4, 2025.",
    ]
    for ref in refs1:
        add_body(doc, ref)

    add_topic_page(doc, "Tài liệu tham khảo 2", [
        "[9] BraTS documentation, “Getting Started and Challenge Documentation,” Read the Docs.",
        "[10] U. Baid et al., “The RSNA-ASNR-MICCAI BraTS 2021 Benchmark on Brain Tumor Segmentation and Radiogenomic Classification,” arXiv/peer-reviewed benchmark documentation.",
        "[11] V. Antonelli et al., “The Medical Segmentation Decathlon,” Nature Communications, 2022.",
        "[12] A. Myronenko, “3D MRI brain tumor segmentation using autoencoder regularization,” BrainLes, 2018.",
        "[13] F. Isensee et al., “nnU-Net: a self-configuring method for deep learning-based biomedical image segmentation,” Nature Methods, 2021.",
        "[14] MONAI Consortium, MONAI documentation: transforms, losses, metrics and sliding-window inference.",
        "[15] PyTorch documentation, AdamW optimizer and reproducibility notes.",
        "[16] ONNX documentation, model interoperability and runtime validation.",
        "[17] Gradio documentation, Blocks application framework.",
        "[18] Quy trình.pdf, tài liệu hướng dẫn xác định vấn đề và triển khai dự án SIC, lưu tại thư mục doc của dự án.",
        "[19] SIC_AI_Capstone Project_Action Plan.docx, biểu mẫu kế hoạch dự án.",
        "[20] SIC_AI_Capstone Project_Final Report.docx, biểu mẫu báo cáo cuối kỳ.",
        "[21] SIC_Capstone_v2.ipynb, notebook thực nghiệm 3D U-Net và đánh giá held-out test.",
        "[22] SIC_Capstone_v3.ipynb, notebook thực nghiệm Swin UNETR và full validation.",
        "[23] BraTS_Model/app.py và BraTS_Model/inference_utils.py, mã nguồn ứng dụng và suy luận.",
        "[24] Report_BraTS-GLI-2026_nzd2b1s9.pdf, báo cáo kết quả minh họa do ứng dụng sinh.",
        "[25] O. Ronneberger, P. Fischer, and T. Brox, ‘U-Net: Convolutional Networks for Biomedical Image Segmentation,’ MICCAI, 2015. DOI: 10.1007/978-3-319-24574-4_28.",
        "[26] Z. Liu et al., ‘Swin Transformer: Hierarchical Vision Transformer using Shifted Windows,’ ICCV, 2021. DOI: 10.1109/ICCV48922.2021.00986.",
        "[27] I. Loshchilov and F. Hutter, ‘Decoupled Weight Decay Regularization,’ ICLR, 2019.",
        "[28] D. P. Kingma and J. Ba, ‘Adam: A Method for Stochastic Optimization,’ ICLR, 2015.",
    ])

    # Appendices
    next_page(doc)
    add_chapter(doc, "APPENDICES")
    add_heading(doc, "Appendix A – Cấu hình tái lập notebook v2", 2)
    add_table(doc, ["Nhóm", "Giá trị"], [
        ("Dataset", "BraTS 2023; 1.251 ca hợp lệ"),
        ("Split", "875 train / 188 val / 188 held-out test; seed 42"),
        ("Modalities", "t1n, t1c, t2w, t2f"),
        ("ROI", "96×96×96; batch 1; num_samples 2"),
        ("Training", "150 steps/epoch; 30 epoch ban đầu; resume tới 50"),
        ("Optimizer", "AdamW; LR 2e-4, resume 2e-5; weight decay 1e-4"),
        ("Validation", "Fast 20 ca mỗi 2 epoch; full mỗi 10 epoch"),
        ("Inference", "Sliding window overlap 0,5"),
        ("Model", "3D U-Net; 4.811.129 tham số"),
    ], [5.0, 10.6])
    add_body(doc, "Tái lập phải sử dụng đúng manifest và checkpoint thay vì chỉ dựa vào tên notebook. Held-out test không được dùng để quyết định learning rate hoặc epoch.")

    add_topic_page(doc, "Appendix B – Cấu hình tái lập notebook v3", [
        "Bảng ghi lại cấu hình quan sát từ output notebook v3. Số train 812 được suy ra từ tổng 1.016 và full validation 204. Khi chạy lại, cần xuất case list để xác nhận con số và nguồn lọc.",
    ], table=(["Nhóm", "Giá trị"], [
        ("Dataset", "1.016 ca sau đồng bộ/lọc"),
        ("Split", "Khoảng 812 train / 204 full validation; không held-out test"),
        ("ROI", "96×96×96; batch 2; num_samples 2"),
        ("Training", "20 epoch; early stopping patience 5"),
        ("Optimizer", "AdamW; LR 2e-4; weight decay 1e-4; cosine annealing"),
        ("Fast validation", "2 ca; interval 2"),
        ("Inference", "sw_batch_size 2; overlap 0,5"),
        ("Model", "Swin UNETR feature_size 24; checkpointing; 15.705.646 tham số"),
        ("Export", "TorchScript 67,65 MB; ONNX 62,80 MB"),
    ], [5.0, 10.6]))

    add_topic_page(doc, "Appendix C – Định nghĩa chỉ số và empty-mask policy", [
        "Metric implementation phải được xem như một phần của protocol. Cùng prediction nhưng các thư viện có thể xử lý empty mask khác nhau. Báo cáo chính cần ghi số case được loại, gán 1, gán 0 hoặc NaN cho từng vùng.",
    ], table=(["Trường hợp", "Dice đề xuất", "Cách báo cáo"], [
        ("GT>0, Pred>0", "Tính Dice chuẩn", "Đưa vào conditional và macro"),
        ("GT>0, Pred=0", "0", "False negative hoàn toàn"),
        ("GT=0, Pred>0", "0", "False positive; báo volume"),
        ("GT=0, Pred=0", "1 hoặc NaN theo protocol", "Luôn báo số ca; không dùng để nói năng lực phân đoạn"),
    ], [4.0, 5.0, 6.6]))

    add_topic_page(doc, "Appendix D – Ma trận truy vết yêu cầu", [
        "Ma trận nối nhu cầu với thiết kế, bằng chứng và khoảng trống. Đây là artifact follow-up: khi requirement thay đổi, test và bằng chứng liên quan phải được cập nhật.",
    ], table=(["Req", "Nhu cầu", "Thành phần", "Bằng chứng", "Khoảng trống"], [
        ("R1", "Nhận 4 modality", "UI/preprocess", "app.py", "Header/shape validation"),
        ("R2", "Phân đoạn WT/TC/ET", "Model/postprocess", "Notebook/output", "ET v3"),
        ("R3", "Xem 3 mặt phẳng", "Viewer", "Ảnh/PDF mẫu", "Interactive QA"),
        ("R4", "Xuất mask", "Export", "NIfTI output", "Metadata/checksum"),
        ("R5", "So sánh model", "Evaluation", "Bảng báo cáo", "Benchmark chung"),
        ("R6", "Sử dụng an toàn", "UX/governance", "Disclaimer", "External/expert validation"),
    ], [1.4, 3.8, 3.2, 3.4, 3.8], 8.2))

    add_topic_page(doc, "Appendix E – Danh sách test case chi tiết", [
        "Các test case ưu tiên có thể được chuyển thẳng thành unit/integration test. Dữ liệu fixture cần được de-identify và nhỏ đủ để chạy trong CI; test GPU nặng có thể tách nightly pipeline.",
    ], table=(["ID", "Input", "Expected", "Priority"], [
        ("T01", "4 NIfTI hợp lệ", "Mask đúng shape/label", "P0"),
        ("T02", "Thiếu một modality", "Error rõ, không inference", "P0"),
        ("T03", "Affine khác", "Cảnh báo/chặn", "P0"),
        ("T04", "Intensity toàn 0", "Cảnh báo input bất thường", "P1"),
        ("T05", "Ca GT ET>0", "Conditional ET metric", "P0"),
        ("T06", "Ca GT ET=0", "False-positive check", "P0"),
        ("T07", "Checkpoint sai kiến trúc", "Load fail có kiểm soát", "P1"),
        ("T08", "PyTorch/ONNX cùng input", "Parity trong tolerance", "P1"),
        ("T09", "10 request liên tiếp", "Không memory leak", "P1"),
        ("T10", "Export PDF/NIfTI", "Mở lại và metadata đúng", "P1"),
    ], [2.0, 5.0, 6.2, 2.4]))

    add_topic_page(doc, "Appendix F – Model card tóm tắt 3D U-Net", [
        "Intended use: nghiên cứu và đào tạo phân đoạn glioma trên dữ liệu có preprocessing tương thích BraTS. Không intended use: chẩn đoán, điều trị hoặc dữ liệu chưa kiểm chứng. Input: bốn modality; output: mask các vùng u. Bằng chứng: test nội bộ 188 ca.",
        "Known limitations: Mean Dice 0,7361, HD95 khoảng 43 mm; lỗi ở vùng nhỏ/biên; chưa external validation; không có uncertainty. Cần ghi checkpoint hash và environment trong bản model card triển khai.",
    ], table=(["Mục", "Thông tin"], [
        ("Model", "3D U-Net, 4.811.129 tham số"),
        ("Data", "BraTS 2023, split v2"),
        ("Primary evidence", "Held-out test 188 ca"),
        ("Mean Dice", "0,7361"),
        ("Limit", "Không external/clinical validation"),
        ("Status", "Research baseline"),
    ], [4.7, 10.9]))

    add_topic_page(doc, "Appendix G – Model card tóm tắt Swin UNETR", [
        "Intended use: nghiên cứu kiến trúc Transformer 3D trên dữ liệu tương thích pipeline v3. Bằng chứng hiện tại là full validation 204 ca; không có held-out test. Mean Dice 0,8653 phải đi kèm lưu ý ET empty–empty.",
        "Known limitations: cohort khác v2; fast validation chỉ hai ca; 204/204 ca không có ET theo bảng output; chưa HD95; chưa external validation. Status nên là experimental cho đến khi audit nhãn và benchmark chung hoàn tất.",
    ], table=(["Mục", "Thông tin"], [
        ("Model", "Swin UNETR, 15.705.646 tham số"),
        ("Data", "1.016 ca pipeline v3"),
        ("Primary evidence", "Full validation 204 ca"),
        ("Mean Dice", "0,8653, bị ET=1 ảnh hưởng"),
        ("Limit", "Không test; ET chưa được thử thách"),
        ("Status", "Experimental"),
    ], [4.7, 10.9]))

    add_topic_page(doc, "Appendix H – Hướng dẫn sử dụng nhanh", [
        "Bản mẫu được chạy trong môi trường cục bộ theo hướng dẫn của repository. Người dùng chuẩn bị bốn file NIfTI của cùng một ca, xác minh không chứa thông tin định danh, chọn model và tải từng modality vào đúng ô. Sau khi chạy, người dùng xem overlay ba mặt phẳng trước khi tải mask hoặc PDF.",
    ], bullets=[
        "Bước 1: chọn 3D U-Net hoặc Swin UNETR; ưu tiên baseline theo mục tiêu thử nghiệm.",
        "Bước 2: tải đúng T1n, T1c, T2w và T2-FLAIR của cùng case.",
        "Bước 3: kiểm tra preview/shape nếu giao diện hiển thị.",
        "Bước 4: chạy phân đoạn và đợi thông báo hoàn tất.",
        "Bước 5: review mask ở axial, coronal và sagittal; không chỉ xem lát trung tâm.",
        "Bước 6: tải NIfTI/PDF và lưu cùng model version.",
        "Bước 7: nếu mask bất thường, không dùng kết quả; ghi issue và chuyển cho người có chuyên môn.",
    ])

    add_topic_page(doc, "Appendix I – Bản đồ artifact và khả năng tái lập", [
        "Bảng liệt kê artifact chính trong repository. Đường dẫn tương đối được dùng để tài liệu có thể di chuyển cùng dự án. Trước khi nộp, nhóm cần xác nhận tất cả file còn tồn tại và checkpoint không bị thay đổi.",
    ], table=(["Artifact", "Vai trò"], [
        ("SIC_Capstone_v2.ipynb", "Huấn luyện/đánh giá 3D U-Net"),
        ("SIC_Capstone_v3.ipynb", "Huấn luyện/đánh giá Swin UNETR"),
        ("best_unet3d.pth", "Checkpoint baseline"),
        ("best_swin_unetr.pth", "Checkpoint Transformer"),
        ("BraTS_Model/app.py", "Giao diện Gradio"),
        ("BraTS_Model/inference_utils.py", "Tiền xử lý và suy luận"),
        ("doc/Quy trình.pdf", "Hướng dẫn xác định/follow-up vấn đề"),
        ("doc/Report_BraTS-GLI-2026_nzd2b1s9.pdf", "Output minh họa"),
    ], [7.2, 8.4]))

    # Word requires a paragraph after the final table. Keep it at one point so
    # the mandatory marker does not spill onto an otherwise blank last page.
    if doc.paragraphs and not doc.paragraphs[-1].text.strip():
        final_p = doc.paragraphs[-1]
        final_p.paragraph_format.space_before = Pt(0)
        final_p.paragraph_format.space_after = Pt(0)
        final_p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        final_p.paragraph_format.line_spacing = Pt(1)
        final_run = final_p.add_run(" ")
        set_run_font(final_run, "Arial", 1)

    force_times_new_roman(doc)
    doc.save(OUTPUT)
    normalize_package_font_names(OUTPUT)
    print(f"Saved: {OUTPUT}")
    print(f"Size: {OUTPUT.stat().st_size} bytes")


if __name__ == "__main__":
    build_report()
